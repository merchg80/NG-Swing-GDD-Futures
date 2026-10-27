import os
from pathlib import Path
import pandas as pd
import streamlit as st
from pathlib import Path
import pandas as pd
import numpy as np

REQUIRED = ['trade_date','hub','product','strip','contract','settlement_price']

def normalize_prices(df):
    df = df.copy()
    df.columns = [str(c).strip().lower().replace(' ','_') for c in df.columns]
    missing = set(REQUIRED)-set(df.columns)
    if missing: raise ValueError('Missing pricing columns: '+', '.join(sorted(missing)))
    df = df[REQUIRED].copy()
    for col in ['trade_date','strip']:
        df[col] = pd.to_datetime(df[col].astype(str), errors='coerce', format='mixed').dt.normalize()
    df['settlement_price'] = pd.to_numeric(df['settlement_price'], errors='coerce')
    df = df.dropna(subset=REQUIRED)
    for col in ['contract','hub','product']: df[col] = df[col].astype(str).str.strip()
    if df.empty: raise ValueError('No valid pricing records found.')
    return df

def read_prices(source):
    name = str(getattr(source,'name',source)).lower().split('?')[0]
    if hasattr(source,'seek'): source.seek(0)
    if name.endswith('.parquet'): df = pd.read_parquet(source)
    elif name.endswith('.csv'): df = pd.read_csv(source)
    else: df = pd.read_excel(source, sheet_name=0, usecols='A:K')
    return normalize_prices(df)

def workbook_defaults(source):
    import openpyxl
    if hasattr(source,'seek'): source.seek(0)
    s = openpyxl.load_workbook(source,data_only=True,read_only=True)['Pricer']
    usage = {int(s.cell(r,32).value):float(s.cell(r,43).value) for r in range(17,29)}
    return dict(customer=s['AN1'].value,start=str(s['N3'].value)[:10],transport=float(s['AH1'].value),broker=float(s['AJ1'].value),margin=float(s['AG40'].value),usage=usage)

def curve(prices, asof, selections, months):
    snapshot = prices[prices.trade_date.eq(pd.Timestamp(asof))]
    out = pd.DataFrame(index=pd.DatetimeIndex(months,name='month'))
    for i,(hub,product,contract) in enumerate(selections):
        x = snapshot[(snapshot.hub==hub)&(snapshot['product']==product)&(snapshot.contract==contract)]
        if x.strip.duplicated().any(): raise ValueError('Duplicate settlements for selected product/date; clean the source before pricing.')
        out[f'component_{i+1}'] = x.set_index('strip').settlement_price.reindex(out.index)
    return out

def calculate(pricing, usage, transport, broker, margin, fuel=0., fixed_quote=None, capital_hurdle=0.08, financing_days=30., pricing_mode='Fixed price', customer_adder=.38, supplier_premium=0., transport_pass_through=False):
    x=pricing.copy()
    components=[c for c in x if c.startswith('component_')]
    if x[components].isna().any().any(): raise ValueError('Missing settlement months; incomplete terms cannot be quoted.')
    x['usage_mmbtu']=np.asarray(usage,dtype=float)
    if not np.isfinite(x.usage_mmbtu).all() or (x.usage_mmbtu<0).any() or x.usage_mmbtu.sum()<=0: raise ValueError('Usage must be finite, nonnegative and total above zero.')
    if not all(np.isfinite(v) for v in [transport,broker,margin,fuel,capital_hurdle,financing_days,customer_adder,supplier_premium]): raise ValueError('Deal assumptions must be finite.')
    if fixed_quote is not None and not np.isfinite(fixed_quote): raise ValueError('Offer must be finite.')
    if not 0<=fuel<1: raise ValueError('Fuel must be between zero and 100%.')
    x['market_rate']=x[components].sum(axis=1)
    x['supply_rate']=(x.market_rate+supplier_premium)/(1-fuel)
    x['transport_rate']=transport
    x['broker_rate']=broker
    x['landed_rate']=x.supply_rate+transport+broker
    if capital_hurdle<0 or financing_days<0: raise ValueError('Capital hurdle and financing days must be nonnegative.')
    x['capital_base_rate']=x.landed_rate.clip(lower=0.)
    x['capital_rate']=x.capital_base_rate*capital_hurdle*financing_days/365.
    x['all_in_cost_rate']=x.landed_rate+x.capital_rate
    indexed=pricing_mode=='Index + adder'
    passback=transport if transport_pass_through else 0.
    x['transport_recovery_rate']=passback
    residual=x.all_in_cost_rate-passback-(x.market_rate if indexed else 0.)
    target=float(np.average(residual,weights=x.usage_mmbtu)+margin)
    quote=(customer_adder if indexed else target) if fixed_quote is None else float(fixed_quote)
    x['offer_rate']=quote
    x['customer_rate']=(x.market_rate if indexed else 0.)+quote+passback
    x['transport_recovery']=passback*x.usage_mmbtu
    x['revenue']=x.customer_rate*x.usage_mmbtu
    for name,col in [('supply_cost','supply_rate'),('transport_cost','transport_rate'),('broker_cost','broker_rate')]: x[name]=x[col]*x.usage_mmbtu
    x['capital_cost']=x.capital_rate*x.usage_mmbtu
    x['gross_profit_before_capital']=x.revenue-x.supply_cost-x.transport_cost-x.broker_cost
    x['gross_profit']=x.gross_profit_before_capital-x.capital_cost
    x['margin_per_mmbtu']=x.customer_rate-x.all_in_cost_rate
    return x,target

def required_credit(result, ft_rate):
    volumes=result.usage_mmbtu.to_numpy(dtype=float)
    supply=result.supply_rate.to_numpy(dtype=float)
    if len(volumes)<1 or not np.isfinite(volumes).all() or not np.isfinite(supply).all():raise ValueError('Valid monthly usage and supply prices are required for credit sizing.')
    top_two=sorted(volumes,reverse=True)[:2]
    ft_volume=float(np.mean(top_two))
    highest_volume=float(np.max(volumes))
    highest_supply=max(0.,float(np.max(supply)))
    ft_credit=float(ft_rate)*ft_volume
    supply_credit=highest_volume*highest_supply
    return dict(ft_credit=ft_credit,supply_credit=supply_credit,total_credit=ft_credit+supply_credit,top_two_average_volume=ft_volume,highest_volume=highest_volume,highest_supply_price=highest_supply)

def investment_cashflows(result, initial_investment=0., terminal_recovery=0.):
    """One initial outlay and at most one recovery, in the actual final month."""
    n=len(result)
    if n<1: raise ValueError('At least one operating month is required.')
    operating=np.r_[0.,result.gross_profit.to_numpy(dtype=float)]
    collateral=np.zeros(n+1,dtype=float)
    collateral[0]=-float(initial_investment)
    collateral[n]=float(terminal_recovery)
    return pd.DataFrame({'month_number':np.arange(n+1),'operating_profit':operating,'collateral_cashflow':collateral,'net_cashflow':operating+collateral})

def deal_summary(result, annual_discount=0.10, initial_investment=0., terminal_recovery=0.):
    """Monthly cash proxy: GP at month-end; investment at delivery start."""
    if annual_discount<=-1 or initial_investment<0 or terminal_recovery<0:
        raise ValueError('Invalid discount rate or investment inputs.')
    n=len(result)
    schedule=investment_cashflows(result,initial_investment,terminal_recovery)
    flows=schedule.net_cashflow.to_numpy()[1:]
    monthly_discount=(1+annual_discount)**(1/12)-1
    npv=-initial_investment+sum(cf/(1+monthly_discount)**(i+1) for i,cf in enumerate(flows))
    cash=np.r_[-initial_investment,flows]
    irr=None
    status='Requires an initial investment and a cash-flow sign change'
    if initial_investment>0 and (cash>0).any():
        roots=np.polynomial.polynomial.polyroots(cash)
        candidates=[float(r.real) for r in roots if abs(r.imag)<1e-7 and r.real>0]
        if len(candidates)==1:
            annual=(1/candidates[0])**12-1
            if np.isfinite(annual): irr=annual;status='Annual effective IRR'
        elif len(candidates)>1: status='Multiple IRRs; no unique return'
        else: status='No valid IRR'
    revenue=float(result.revenue.sum());volume=float(result.usage_mmbtu.sum());gp=float(result.gross_profit.sum())
    return dict(annual_revenue=revenue*12/n, volume=volume, annual_volume=volume*12/n,
                contract_price=float(np.average(result.customer_rate,weights=result.usage_mmbtu)),collateral_recovery_month=n if terminal_recovery>0 else None,npv=float(npv),irr=irr,irr_status=status,
                spread=gp/volume,gross_margin=gp/revenue if revenue>0 else None)

def standard_offers(prices, asof, selections, start, usage, fallback_profile, transport, broker, margin, fuel, capital_hurdle, financing_days, discount=0.10, investment=0., recovery=0., overrides=None, automatic_credit=False, full_credit_recovery=True, **deal_options):
    """Exact entered months first, then repeat first-year calendar-month estimates."""
    schedule=usage.copy()
    schedule['month']=pd.to_datetime(schedule['month'])
    exact=dict(zip(schedule['month'],schedule['usage_mmbtu']))
    profile=dict(fallback_profile)
    for row in schedule.iloc[:12].itertuples(index=False): profile[row.month.month]=row.usage_mmbtu
    rows=[]
    for n in [12,18,24,36,48]:
        months=pd.date_range(start,periods=n,freq='MS')
        volumes=[exact[m] if m in exact else profile[m.month] for m in months]
        projected=sum(m not in exact for m in months)
        pricing=curve(prices,asof,selections,months)
        missing=int(pricing.isna().any(axis=1).sum())
        if missing:
            rows.append(dict(term_months=n,offer_price=None,break_even=None,contract_volume=None,contract_profit=None,npv=None,irr=None,discount_rate=discount,collateral_recovery_month=None,irr_status="Unavailable",projected_usage_months=projected,status=f'Unavailable: {missing} months missing settlements'))
            continue
        try:
            result,offer=calculate(pricing,volumes,transport,broker,margin,fuel,fixed_quote=(overrides or {}).get(n),capital_hurdle=capital_hurdle,financing_days=financing_days,**deal_options)
            credit=required_credit(result,transport)
            term_investment=credit['total_credit'] if automatic_credit else investment
            term_recovery=term_investment if full_credit_recovery and automatic_credit else recovery
            returns=deal_summary(result,discount,term_investment,term_recovery)
            rows.append(dict(term_months=n,upfront_credit=term_investment,ft_credit=credit['ft_credit'],supply_credit=credit['supply_credit'],credit_recovery=term_recovery,target_price=offer,offer_price=float(result.offer_rate.iloc[0]),projected_all_in_price=returns['contract_price'],break_even=offer-margin,spread=result.gross_profit.sum()/result.usage_mmbtu.sum(),gross_margin=returns['gross_margin'],annual_revenue=returns['annual_revenue'],contract_volume=result.usage_mmbtu.sum(),contract_profit=result.gross_profit.sum(),npv=returns['npv'],irr=returns['irr'],discount_rate=discount,collateral_recovery_month=n if term_recovery>0 else None,irr_status=returns['irr_status'],projected_usage_months=projected,status='Available'))
        except ValueError as e:
            rows.append(dict(term_months=n,offer_price=None,break_even=None,contract_volume=None,contract_profit=None,npv=None,irr=None,discount_rate=discount,collateral_recovery_month=None,irr_status="Unavailable",projected_usage_months=projected,status=str(e)))
    return pd.DataFrame(rows)

def build_summary_pdf(meta, summary, result, scenarios):
    from io import BytesIO
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
    from reportlab.lib.enums import TA_LEFT, TA_CENTER
    buffer=BytesIO()
    navy=colors.HexColor('#074F42');teal=colors.HexColor('#862131');pale=colors.HexColor('#F1F5F3')
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Brand',fontName='Helvetica-Bold',fontSize=17,leading=21,textColor=navy))
    styles.add(ParagraphStyle(name='SmallREM',fontName='Helvetica',fontSize=8,leading=11,textColor=navy))
    styles.add(ParagraphStyle(name='SectionREM',fontName='Helvetica-Bold',fontSize=10,leading=14,textColor=teal,spaceBefore=9,spaceAfter=6))
    styles.add(ParagraphStyle(name='ProfitREM',fontName='Helvetica-Bold',fontSize=9,leading=12,textColor=navy,alignment=TA_CENTER))
    def para(text):return Paragraph(escape(str(text)),styles['SmallREM'])
    def usd(v):return f'${v:,.0f}'
    def pct(v):return f'{v:.2%}' if v is not None else 'N/A'
    logo=Image(BytesIO(get_logo_bytes()),width=300,height=300*402/2096)
    logo.hAlign='LEFT'
    story=[logo,Paragraph('DEAL ECONOMICS | INTERNAL REVIEW',styles['SectionREM']),para(meta['customer']),para(f"Utility: {meta.get('utility') or 'Not specified'}   |   Facility: {meta.get('facility_type') or 'Not specified'}"),para(f"Location: {meta.get('client_location') or 'Not specified'}"),para(f"Broker: {meta.get('broker_name','Not specified')}   |   Offer due: {meta.get('due_date','Not specified')}"),para(f"Delivery start: {meta['start']}   |   Term: {meta['term']} months   |   Settlements: {meta['asof']}"),Spacer(1,8)]
    metrics=[('Annualized revenue',usd(summary['annual_revenue'])),('Contract volume',f"{summary['volume']:,.0f} MMBtu"),('Projected weighted all-in price',f"${summary['contract_price']:.4f}/MMBtu"),(f"NPV @ {meta['discount']:.2%}",usd(summary['npv'])),('Break-even adder' if meta.get('pricing_mode')=='Index + adder' else 'Break-even offer',f"${meta['break_even']:.4f}/MMBtu"),('Contract profit',usd(result.gross_profit.sum())),('Spread after capital',f"${summary['spread']:.4f}/MMBtu"),('Margin after capital',pct(summary['gross_margin']))]
    cards=[]
    for i in [0,4]:
        cards.append([Paragraph(f'<font size="8">{escape(label)}</font><br/><br/><b>{escape(value)}</b>',styles['SmallREM']) for label,value in metrics[i:i+4]])
    table=Table(cards,colWidths=[132]*4,rowHeights=[54,54])
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),pale),('VALIGN',(0,0),(-1,-1),'TOP'),('BOX',(0,0),(-1,-1),.4,colors.white),('INNERGRID',(0,0),(-1,-1),4,colors.white),('LEFTPADDING',(0,0),(-1,-1),10),('TOPPADDING',(0,0),(-1,-1),10)]));story.append(table)
    story.append(Paragraph('Offer comparison',styles['SectionREM']))
    rows=[['Scenario',('Adder\n$/MMBtu' if meta.get('pricing_mode')=='Index + adder' else 'Offer\n$/MMBtu'),'Spread\n$/MMBtu','Annualized\nprofit','Contract\nprofit','NPV']]
    for row in scenarios:
        npv=next(v for k,v in row.items() if k.startswith('NPV @'))
        rows.append([row['Scenario'],f"${row['Offer $/MMBtu']:.4f}",f"${row['Spread $/MMBtu']:.4f}",usd(row['Annualized GP $']),usd(row['Contract GP $']),usd(npv)])
    def styled_table(rows,widths,highlight_profit=False):
        t=Table([[Paragraph(escape(str(v)).replace('\n','<br/>'),styles['ProfitREM']) if highlight_profit and col in (3,4) else para(v) for col,v in enumerate(row)] for row in rows],colWidths=widths,hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),pale),('LINEBELOW',(0,0),(-1,0),1,teal),('LINEBELOW',(0,1),(-1,-1),.3,colors.HexColor('#DBE5E9')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        if highlight_profit:
            t.setStyle(TableStyle([('BACKGROUND',(3,0),(4,-1),pale)]+[('BOX',(col,row),(col,row),1,teal) for col in (3,4) for row in range(len(rows))]))
        return t
    story.append(styled_table(rows,[116,78,78,88,88,80],highlight_profit=True))
    story.append(Paragraph('Pricing and credit assumptions',styles['SectionREM']))
    assumptions=[['Upstream transport',f"${meta['transport']:.4f}/MMBtu",'Brokerage',f"${meta['broker']:.4f}/MMBtu"],['Target spread',f"${meta['margin']:.4f}/MMBtu",'Supply fuel',pct(meta['fuel'])],['Capital hurdle',pct(meta['capital_hurdle']),'Financing period',f"{meta['days']:g} days"],['Capital charge',f"${result.capital_cost.sum()/summary['volume']:.4f}/MMBtu",'Discount rate',pct(meta['discount'])],['Upfront credit cash',usd(meta['investment']),'Terminal credit recovery',usd(meta['recovery'])]]
    assumptions.insert(0,['Pricing structure',meta.get('pricing_mode','Fixed price'),'Transport pass-through','Full cost' if meta.get('transport_pass_through') else 'None'])
    if meta.get('pricing_mode')=='Index + adder': assumptions.insert(1,['Customer index adder',f"${meta['customer_offer_rate']:.4f}/MMBtu",'Supplier premium',f"${meta['supplier_premium']:.4f}/MMBtu"])
    if meta.get('credit_mode')=='Automatic':assumptions.append(['Automatic FT credit',usd(meta['ft_credit']),'Automatic supply credit',usd(meta['supply_credit'])])
    story.append(styled_table(assumptions,[132,132,132,132]))
    story.append(Paragraph('Source and valuation basis',styles['SectionREM']))
    story.append(para('Source: '+meta['source']))
    story.append(para('Market components: '+meta['components']))
    story.append(Spacer(1,6))
    story.append(para('NPV is valued at contract start using the selected effective annual rate, month-end profit after capital charge, upfront credit and terminal recovery. Annualized figures equal contract totals x 12 / term months.'))
    story.append(Spacer(1,5))
    story.append(para('Indexed bills = signed market curve + adder + transport pass-through when enabled. Projected all-in price is volume-weighted. Transport recovery equals its expense; financing remains. Break-even shows the adder or fixed offer required after recovery. Overhead and taxes are excluded.'))
    def footer(c,doc):
        c.setStrokeColor(teal);c.line(42,39,570,39);c.setFont('Helvetica',7);c.setFillColor(navy);c.drawString(42,27,'CONFIDENTIAL | Riverbend Gas Marketing');c.drawRightString(570,27,f'Page {doc.page}')
    SimpleDocTemplate(buffer,pagesize=(612,792),leftMargin=42,rightMargin=42,topMargin=36,bottomMargin=50,title='Riverbend Deal Economics',author='Riverbend Gas Marketing').build(story,onFirstPage=footer,onLaterPages=footer)
    return buffer.getvalue()

"""Durable named workstreams in a private GitHub repository."""
import base64
import json
import re
import uuid
from datetime import datetime, timezone
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from urllib.parse import quote

class WorkstreamError(ValueError): pass

def github_request(repo,token,path='',method='GET',body=None):
    if not re.fullmatch(r'[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+',repo): raise WorkstreamError('Set WORKSTREAM_REPO to owner/private-repository.')
    url=f'https://api.github.com/repos/{repo}'+path
    headers={'Accept':'application/vnd.github+json','Authorization':f'Bearer {token}','X-GitHub-Api-Version':'2022-11-28','User-Agent':'REM-Deal-Economics'}
    data=None if body is None else json.dumps(body,allow_nan=False).encode()
    try:
        with urlopen(Request(url,data=data,headers=headers,method=method),timeout=20) as response:return json.load(response)
    except HTTPError as e:
        if e.code==404: raise WorkstreamError('Not found. Check the private repository, branch and token access.') from None
        if e.code in [409,422]: raise WorkstreamError('Save conflict: reopen the current saved version or save a new workstream.') from None
        raise WorkstreamError(f'GitHub request failed (HTTP {e.code}). Check token permissions or expiry.') from None
    except (URLError,TimeoutError):raise WorkstreamError('Unable to reach saved-workstream storage. Your current inputs remain on screen.') from None

def verify_private(repo,token):
    info=github_request(repo,token)
    if not info.get('private'): raise WorkstreamError('Use a private repository for saved customer economics. WORKSTREAM_REPO currently points to a public repository.')

def list_workstreams(repo,token,branch='main'):
    verify_private(repo,token)
    try:items=github_request(repo,token,'/contents/workstreams?ref='+quote(branch,safe=''))
    except WorkstreamError as e:
        if str(e).startswith('Not found'):return []
        raise
    return sorted([x for x in items if x.get('type')=='file' and x['name'].endswith('.json')],key=lambda x:x['name'].lower())

def load_workstream(repo,token,path,branch='main'):
    if not re.fullmatch(r'workstreams/[A-Za-z0-9_.-]+\.json',path):raise WorkstreamError('Invalid saved workstream path.')
    item=github_request(repo,token,'/contents/'+path+'?ref='+quote(branch,safe=''))
    payload=json.loads(base64.b64decode(item['content']))
    if payload.get('schema_version')!=1: raise WorkstreamError('Unsupported workstream version.')
    return payload,item['sha']

def save_workstream(repo,token,name,snapshot,branch='main',path=None,sha=None):
    verify_private(repo,token)
    if not name.strip():raise WorkstreamError('Enter a workstream name.')
    if path is None:
        slug=re.sub('[^a-z0-9]+','-',name.lower()).strip('-')[:60] or 'deal'
        path=f'workstreams/{slug}-{uuid.uuid4().hex[:10]}.json'
    if not re.fullmatch(r'workstreams/[A-Za-z0-9_.-]+\.json',path):raise WorkstreamError('Invalid saved workstream path.')
    if sha is None and snapshot.get('_existing'):raise WorkstreamError('Reopen the saved workstream before updating it.')
    payload=dict(snapshot,schema_version=1,name=name.strip(),saved_at=datetime.now(timezone.utc).isoformat())
    body=dict(message='Save deal workstream: '+name.strip(),branch=branch,content=base64.b64encode(json.dumps(payload,indent=2,allow_nan=False).encode()).decode())
    if sha:body['sha']=sha
    response=github_request(repo,token,'/contents/'+path,'PUT',body)
    return path,response['content']['sha'],payload

import base64

LOGO_BASE64 = 'iVBORw0KGgoAAAANSUhEUgAACDAAAAGSCAIAAACQADTEAAC8JElEQVR4nOzdf1xUVf4/8LvOtItiCaMwbGiglIpBmKOmiLiA5S/6ksouSiRJiRomJhu6mdMGWYErqSur4ophRrJByoa/SmFFRVMxCRK1QEhpgVF+mOjsxrTfP24fluXnnZl7zzn3zuv52Mdjbbwz5y3MzL33vM95v3/xn//8hwMAAAAAAAAAAAAAAJBSH9oBAAAAAAAAAAAAAACA8iEhAQAAAAAAAAAAAAAAkkNCAgAAAAAAAAAAAAAAJIeEBAAAAAAAAAAAAAAASA4JCQAAAAAAAAAAAAAAkBwSEgAAAAAAAAAAAAAAIDkkJAAAAAAAAAAAAAAAQHJISAAAAAAAAAAAAAAAgOSQkAAAAAAAAAAAAAAAAMmpe/7rphs1N4q/JBMKAIBEvEKCaYcAAAAAAAAAAABg63pJSNwo/nLf8ngyoQAASAQJCQAAAAAAAAAAAOpQsgkAAAAAAAAAAAAAACSHhAQAAAAAAAAAAAAAAEgOCQkAAAAAAAAAAAAAAJAcEhIAAAAAAAAAAAAAACA5JCQAAAAAAAAAAAAAAEBySEgAAAAAAAAAAAAAAIDkkJAAAAAAAAAAAAAAAADJISEBAAAAAAAAAAAAAACSQ0ICAAAAAAAAAAAAAAAkh4QEAAAAAAAAAAAAAABIDgkJAAAAAAAAAAAAAACQHBISAAAAAAAAAAAAAAAgOSQkAAAAAAAAAAAAAABAckhIAAAAAAAAAAAAAACA5JCQAAAAAAAAAAAAAAAAySEhAQAAAAAAAAAAAAAAkkNCAgAAAAAAAAAAAAAAJIeEBAAAAAAAAAAAAAAASA4JCQAAAAAAAAAAAAAAkBwSEgAAAAAAAAAAAAAAIDkkJAAAAAAAAAAAAAAAQHJISAAAAAAAAAAAAAAAgOSQkAAAAAAAAAAAAAAAAMkhIQEAAAAAAAAAAAAAAJJDQgIAAAAAAAAAAAAAACSHhAQAAAAAAAAAAAAAAEhOTTsAAAAAAIXbW3j0ISet60CnDo+7DnRSq1RUQrJxLUbjp2dPerkNu79vvw5/5ebsQiWkNtX1tfwfyq9XTX50tL2dncAn8v8oyeKC3j3kpPX19KYdBQBA74rKS78z1NGOwnbN859KbCwGLw9CJwXY1AUw+x83XMAAeUhIAADYHENzU9xfN7f9Z1bhMYrB9EDrqPmN9+NSjzJisJvHr13b/tPBvr/nEHeO4/r9ys5pgIPUoysJ/76aOc63/YNdTvhyHDfoAQfh06ySajEab95u6vz4D/fullVXtn/k4LmiDS8ut+xdEbkh0dynSPr+H/SAw/gRo8x6ysSRXh0e6fAjktedzM3bTeb+UqT7jfyj9Mu6xoYu/+rqjix7O6EJEgv+USCuMP8gYp+CrQc+OX25jP/zpe+qSqsqyIwrHIGTeIevsrYzODunGHZU19cOXxQW6KNrfxbr4VzQ3embvPLrVU0td3o+5uC5oh7+trv3iY1f6W07uI/ZWwBbQDIhweDlwcFzRbvj9LSjIIf9j5tYFzALNiR0uBnsfAfB2fzXL/CQkAAAsDl3/2Vk/JKIV9fYwEKcYf5BHMdNHOk11OVBzyHuWNLeHf59xcKvTDqJz0UTG0vq939qXo64L0hyKpYKRr6RAHinL5cx/oYk85Hp+auMP4PPHOfL71Gjvv+JuvyS4g6PiH4uYFOv7xM+aeHlNkzroME0GYDUsgqPTRzptXTWHNqBgMgUfzMoFv76hD/18JcoNjjJgIQEAABAT/iLqvaXVt7uHv5eo8ePGDVxpBdmNwAAANjU+Qwe6KPze9Qn4LExPkMfwS4K4PHvkLakhdZREzop4Kkx43UPj0RyAkAiK9I2TfYa7eU2jHYgABS0XZm0z5d7u3uMesh95jhfL7dhQ7UPKv4qBQkJAAAA85RWVZRWVbRdPcQEz31qzHizSr0DAAAAefklxfklxQmZ6RzHebt7zPadEjJhMmbEoL26xobUvBz+Mk/rqFk845no6SHITACIbuYbccWbduHDBcDjJxnachX8VYqCl1AgIQEAAGCVtrvWQB9dbMjvkJkAkAuto6a75g1y1O9XdvwGcK7HvhQAwOPv/BMy0/lJ54iAadj1CB3UNTYkZKYnZKZ7u3u88FTwgqCZirnGayvsbmhu6lzLC4CAusaGBRsSPn1jva1VqgEQou0qheO4QB/d/3vCL9QvUEkJvF/85z//6eGvy3Lz9i2PJxYNAIAU1lZfoh0CW1qMxk/PnuQbAKLIoxRigudGPRVsaysu+abWSp0G5buzWtzUutVkqrllOH257OyVS4op2M3PzvgMe0QBZU8NzU3ffH+dqZaDonyNtBiNJ76++PruNAY7HitSmH8QsS6dReWlBV9duHKjms2O1jLl7e7x1oLoqaPHyfoLrUv8td/ZK5du3m7CBLQ19OFRSs1dlVVXpn+Wp5irFMb9K/c4yeEMzU3ZJ/N3fpbH4PkiJnhuyqLltKOQFps3AmH+QRNHevFX8mJ9py3YkIDLEknx91/KyEwgIQEAyoeERM+q62vLr1dl/uMzdmbilMHb3WPzklfGDx+lvHkNIarra3+4d7esuvLslUvZpwpklKVo6xHi5Tbs/r79pJh0KKuuzD1zYvuh/TL6sbQX5h8UHxqhyJRbi9G4MTeLX4tEhdZRkxwV8/R4P3EX4Rqam97522527oGVimRCor0Wo7Hk2jcFX12Q77cKU/gNEytCwhSzFr5LLUbjtbrvZX0yoijMPyjxuWhFpiVaTaajF89F/zkJ7wpJEU5ItCmrrkzO3sPaTV+uPmm6bgLtKAjh08N/ysmkMmvPT2fPHOdL4Our7Wbw4LkipS5Zo8jb3eP3c8NFv2UgCQkJAFA+JCQEajEadx87yObaGfnipxfn+U+lHQhl7K/UJl9xi7/nZ/ln0pm3u0fOmrcVOQXTnqG5aYZ+JeHfC/9dETopQLoUZlF56bwkPW4IpUMrIdGm1WQ6e/VS4ke7sP5dFPrwqFWhEbawqoBfPb0ibRPtQGQmzD9oa0y8fCeDetBiNIa+/Rq+SaRDKyHBo778orOrO7IUf3nZwd7Co/HpqWSuylhoh8OfaN75+ANciIpLvrUZkJAAAOVDQsIsrSZT9qmCyA2JtANRFG93j/dXvi7HCwXR7S08yuC7i+7KrLLqyld3bmH8tl/rqEl7eZXtrF9rNZmefvNVYr+UjdGxi6aHEJj3NDQ3LdiQwPibTb6oJyTaIPkkFpv66sP3g2WUmrhqNZmSsvcwNWetJHQTErzq+tolW5IZ+chrHTXl2z5SZHqvB60m047DuVIng5n6jmo1maI2rmNtj44CyLHmJBISAKB8SEhYADelUlDwSjqzVNfXzl33GiPbArzdPQ4lpLBQhZPkOilzaR01xZt2sfBTIqnVZJqwcpHUb9RAH93uOD3Jny2Zf5dtYichwXFcq8kUn56KOl2iIP85pYjNdQOM0zpq9q5K8PX0ph2I+NZlZSAnIQUWEhK8w8VnGKnQFeijO5SQQjsKCqS7M9I6ak4kb2Vw6wlONBIhsN9aREhIAIDyISFhMWZvQtqq/HMcN3GkF8dxgx5waJvobzEab95u4v98+nIZx3EHzxUx0sLRptZa9qDFaNSETaMdBad11FTuzGbnio3NCgm2mY3gGZqbBi8Ike71M+LWUqnnJvW/q1dh/kEzx/mK+IIHzxW1/ZnisjumEhK8ovLSgNXLaEfRUaCPzu9RH49fuzrY9/cc4s5xnOtAp7av4ur6Wv4PNbcM3xnq+D7MLKympPWBJW/rgU9QvskCSl13MkO/kuKVibjni4p/1ly5Uc3/mW7rXXYSEhxL15/68Kg1YZG0o6BAijUEgT667NfeZvYbCTkJ6cglLYGEBAAoHxIS1mAnJxHoo4ucOnPiSK/20xZm4YtrF3x1YV/Rcbqrg8P8g9JXrGH8EkFqLFyGMji7RLhSkBAM/pRIkuiNSn3BNd15agJvKj4zzXdTPHvlUmHZRQJf+wwmJDhmJpf5WcWJI70sXqpZXV978FwR3U5XjM/viKXVZBr2QigLK6Y5jvN29xj1kPugBxz4ZSgPOWldBzrxf8WnrDiOa/zh9unLZSwsPVHkupMWo9FzyXxa7wcC5wtDc9PdfxnLr1ddq/3+9OUyMulPphISPEZu+gre3aLIzUZCiHjNydqiq86YOtEoEvudLJGQAADlQ0LCSnQXRnEc5+3usXnJK+Jemxqam46VnKdYIceWV53zqF+GMnulzlROwtvd40zKDgZ/SsRI8UYN9NF9+sZ66j/VBRsSaK06p5LlajEaS659U/DVhe2H9kv0zcNmQoKjPc0kehelFqPxxNcXX9+dRiUzYSOn78PFZ0ISVlEZWuuoCZ0UMH7EKMvWoPBvj88unM0+VUDrGkN5604o5rCpnC+q62tPXy7LOHpQuosxBhMSHBsJbBv5ju2OWJ81WTQJZ2F1Ws/4jHh3f0t3l5VALHeyREICAJQPCQkr0Z0elXQ+lO/g/aecTFoXE7a8CIijfRnK8sJ/dnISNv4W5Yl7f85INoKzvQmm9orKSxM/2iX6R4zZhARHL/8k6Ru+qLx0+bb3qJzBbeG7cWxsFPmf7cbo2EXTQ8R6w/C7at75+APymQlmS7dbptVksp8TSGVouueLFqNxY26WFGlsNhMSHO2Lc56Nr4ax/vKM+lWWQNRXp/HC/IP4TXhtO/DMzYW3mkw1twz8rly+LhwLdSbbY7OiIBISAKB8SEhYr7q+dviiMCpDk7nnLyovnZekp3I9lKtPUtjWfuHoFrK/sTuX5eVXhuYmXexCutfoWkfNd+/voxgAI8qqK3XLF4ryUqwtm33o+dlU3mOM3CpX19eu/SBNxJtGlhMStDr3EDiJ00pLKD4nQX6ttETnZb5iJ5U3iZKu8WhttGLhfMEvYBJ3XzWzCQmO9qY6HsvnUwKsyUkwuwW8S7RWS2gdNYtnPBMRME26tHF1fW359arib69ItzHXXCx8nbbXh3YAAAAgA27OLvrwKPLjBvroyNzt+3p6V+7MjgmeS2CsDkISVq3LyiA/LgucBjhoHTVUhtY6aljORnAc5zTA4UTyVroxLJ7xDN0AGDFysJsor8P3jWDqFtHGf8Vuzi674/QF727xdvegHYvk7O3sNkbHEh7U292DwEnc19P7TMoO8mfwgNXLispLCQ9K0mSv0SSH04dHSXReVqtUvp7e5zelF7y7JdBHJ8UQ3VHSNV7IhMm0Q6BGrVLN859avu0jKndD5K0JiwzzD6IbQ1bhsb2FR+nGQJGvp3euPsmy5y6e8QxTl5o9E7FlvUBh/kEF72757v19a8IiJd3E5ubsMl03YU1Y5Hfv77u6I2tjdCz1q83IDYkz9CsNzU10w2iDhAQAAAgSPZ3CSvbIqTOJjaVWqVIWLbf44s8aCZnpirlfNddvvB+3qXHNQisR2MaWZx/aU6tU1s9haR012a+9LUo8IsKvmOM4fqYyI24t7UAktyCI3CmV9/u54WQGonUGV3ZOYqj2QZLDrQqNkHoIX0/vT99YTzgzp5hrPDZLkJNkb2e3Jizy6o4swmktKtJXrKH+z4zckFhWXUk3Boqm6yZYlhaics9usYkjvQiPmPhcNPndjW7OLktnzTm/Kf3G7ly6N3f5JcWDF4QcLj5DMYY2SEgAAIAgTgMcqF+YEjBdN+HG7lzy6xcSMtMVPKnRgxEirT2Xy7jmWhUaQWsTCYfZh3asX7d78M0NrFVu5Tju/r79aIfAinn+U2/szlX2ac7ezo76oldJ8Wdwwt+ZCs5JkPzK8nb3ILOkV61SLZ015+qOLJLvE8XkJChekLDDzdnlUEIK+Q1nhKlVKhZWUcx8I46d1dzkpa9YY+6Hztvdg/Et4DbOaYDDmrDIlk/yM+LWUtwwEZKwasGGhFaTiVYAPCQkAABAqNiQ3xEekfyiCY7jnAY4HE/6C/mbLgVPavTA49euNjWuudQq1cE3N1AZWtkzs4RtjI5lM7sz6AEH2iEwxGmAw6GEFGVvlVgyczbJ4cifxJ0GOBRv2kU+J6HUKTNi0yWjHnInMxDPzdmlcmc2ydOcMnISsthdSgaf1qJegEVS9nZ21E+IdY0NLEyb0qJWqZKjYsx6ygtPBUsUjEQkLZrELL4KHL89l1aiN6vw2LAXQqvra6mMzkNCAgAAhPIc4k47BELs7ezIz2hwip7U6A6tWVo2Z4e75OU2jMqiZiywEkugj27prDm0o+gag5s2qJvnP5Xw6mmSHnlwCO0QJOc0wIF8HlcXu1CRU2bE8gTkty2qVapP31hPOCfBSJUMEIWbswuV7jUkzfOfSj3pkl9SnJS9h24MFD093s+s4wn3/gErzfOfWrkzm1bmr66xYfiiMIoLIpGQAAAAoWxqCQOVGQ2O42boVypyUqM7tCrGyKtSTeJz0eQHpbI/SZG2LYunHQKYx83ZpXzbR4rcJGQjiUYvt2EF724hOWJdY0PUxnUkR1QYTf/7yQ/K5yRIzreGJKyytXUnysZ3r1F2+aacNfQLN9lyMs/ezs6sFRKEe/+IgnrSiy5+t0RD1hFaRTUDVi+jtYEPCQkAADCDIidoukN+RoPjuNKqivj0VMKDAuPcnF3IX6SevlxGeESW/aP0S8ueqA+PsqlUrmLY29l9+sZ6RS59VXYbiTa+nt6E+0ZmFR5T3pTZzHG+ZAaidcZRq1SHElJI7oiytXUntmDprDnk7xeIcXN2oduDlxeSsIpubRmKzCqVJsedr4RL9rHJ3s5ud5w+V59EZYduQmY6lXMTEhIAAGAGG1lc2cbX05v83E1qXk5ZdSXhQYFx8aERhEe0eApekeoaGyx4ltZRs4r4Lw7Ewi99ZWEiBiyzKjSC8LrLkIRVLUYjyRHBek4DHE4kbyU2nKzXnRBLUMmOr6f3jd25Sq31FxEwjXYIHMdxk+OX2mYyT3hFOxtZcKBg03UTaO3QzS8pfvrNVwl/xJCQAAAA6MmGF5eTH3TmG3E2cs1Na/G47Bate7kNIzyzVtfYYCNvwl5ZvCgvOSpGrVKJGwwQtiYsEjkJmVKrVJuXvEJ40KWpyYRHBOu5ObuQ3A6FdSeK5DTAgUr/OQIYuWaua2x4+s1XaUdBgcevXQUeSb4ZD4jO3s7uUEIKlStPPidBcl0FEhIAAGAGGywr7zTAgfw1QV1jw47DuYQHBcb9fm444REv36gmPCKbyq9XWfAsraMmdFKA2LGIz8ZL9wqhsJyETa1x9vX0JrzSMKvwGOaa5SjxuWiSU8nPp7xFbCwgRsE5CUZq9uaXFNMqdi8LwlMXwLg1YZG5+iTy4+aXFHsumU+s1xESEgAAYAbH+x+gHQIFq0IjyN9drEjbhMoP0N7T4/0Ij5h75gThEdlU/O0VC54ll+0RKN0rhMJyEjaFfFd5zDXLkb2dXXJUDLHhSqsqispLiQ0HxBCuAEYMOzV7EzLTbe2z85CTlnYIQMF03YSrO7LIT0HUNTboYheSmYVAQgIAAKAXapVq8YxnyI+7MTeL/KDALHs7O8Itdrcf2k9yOGZZ8HOQy/YIEG5NWKQie1xLyrLdReJyc3YhvA0Ic80yRfhLe/m290gOB8S4ObsouMc1C+Yl6Ymt4GaB60AngUfaYCUDZXNzdqGy6aqusSH07dcIVO5FQgIAAKB30dNDyA+akJmOTRLQXtRTwSSHq2tswLSaobnJgo7WctkeAWZJjopBhSuzNLXcoR0Cx3HcWwuiCY+IuWY5UqtUJDdClVZVHC4+Q2w4IMnX0xub6qRT19gwQ78Sfc7AFtAqBEemxzUSEgAAAL1zGuBApXwqNklAe+RbWyd+tIvkcAw6VnLe3Kdge4RSqVWqQwkpiqwPrmxTR48jPGJpVQU6ScgR4dUnm3L/RnI4IGlVaAQjfRcUqbSqIj49lXYUzOn3KzvaIYD46OYkJB0CCQkAAABB1s5fSH7Q7Yf2YwUQtPcC2U0S+SXFNrUvvrOMowfNfQq2RyiY0wCHvasSaEcB5lGrVOTLbaV/lkd4RLAe4dUn+SXF2AirVGqVanecHgls6aTm5ewtPEo7Craw0+oDxEUxJyFpG3kkJAAAgF2nL5fRDuG/fIY+Qn7QusaGs1cvkR8XmLUgaCbhEdMO5xIekR0tRmN+SbG5zyLffhxI8vX0DvMPoh0FmCfUj/SmpdS8HKwnkKP/9wTRL/BPz54kORyQhAS21CI3JGIvGtgIPidBftyEzHTpMn9ISAAAAAhib2dHZfM1auZAe/Z2doQnQ215m86Jry+a+xR9eJS9HbbMK9zWmHjaIYB5qCwpOHrxHPlBwUqhfoEkh7NgEx7IiK+nNzoPSWrmG3HYZgQ2wmmAQ8G7W8iPK13mDwkJAAAAoQivm+PllxTb7HQwdGnJzNkkh7PlbToWFPiOCJgmRSTAFHs7OzQslRd7Ozvy04KfXThLeESwntMAB5JlMSzYhAfysnnJK7RDULK6xobQt1/DjRLYCF9P74y4teTHlSjzh4QEAACwa+JIL9oh/I/JXqOpjHv5RjWVcYFNvp7ehKuI2uY2HQvqNXm7e7g5u0gUDzBlRUgYioPLC+EGPBzHZZ8qIDwiiCJ0EtECX9X1tSSHA8JQ5U9q+SXFSdl7aEcBQMg8/6nkv1L4zJ/oL4uEBAAAgFAjB7tRGTf3zAkq4wKz/vDb50gOZ5utrS0o7f3WgmgpIgEG2dvZJUfF0I4CzOAzjHTVprrGBsw1y9H4EaNIDsdUvzSQQuJzsr82+Efpl7RD6ElCZvrh4jO0o6AMxcFsR/qKNVQaXG898Im4r4mEBAAAgFBqlYrK1d6+ouPkBwWWES5yzdlka+s/5WSa+5Spo8dJEQmwKchnLO0QwAyuA53ID4q5ZjkivD234p81JIcD8tycXeQ+WVzX2EA7hF6EJKyywaUz7Y16yJ12CECIWqU6kbyV/Lgr0jaJu8wCCQkAAAAzULnaK62qID8osMxpgAPhFuvbD+0nORx11fW15n7u9OFRapVKoniAQU4DHOQ+xyQ1L7dhtEP4Lyrl1M5esdEGPLI26AEHksNdQVlOGzDbdwrtEJRPF7sQzSTARrg5u1BpJjF3nZgtW5CQAAAAMMPMcb5UxkXZB+hg7fyFJIera2woKi8lOSJdewqOmPuU6OkhUkQCLCPflkBe7u/bj3YI/4NwHpfjuPLrVYRHBOvZ29mRHO7Sd1UkhwMqQiZMph2C8tU1Njz95qu0owAgJHRSAPllMaVVFTvE2zSPhAQAAIAZHnLSUhkXkxrQwfjhowjXD912cB/J4ShqNZnM3RES5h/kNMBBkmiAYZO9RtMOAcxA/kOaX1JMeEQQBcncFXbB2gKmtospmBRl7gHYpFapcta8TX5cEQs3qUV5FQAAABtBpQg1x3FNLXeojAvMUqtUi2c8k5CZTmzErMJjW2PiCS8dpeLs1UvmlkteMnO2RMEAy7zchrXfMk8rYw0CzRznm1V4jPCghuYmZCtlB78yEF3Bu1u+M9TRjkL5VqRt8hn2iK+nN+1AACTn5uyiD48ieTPIW7Il+VBCivWvgx0SAAAAZuj3KzqzsQfPFVEZF1hGvkbQ7mMHCY9IReJHu8w63tvdA7e+Nmue/9S2/+FtAJ3d/ZeRdghgNsJ9rcEW+Hp6tz9f0A5HyeYl6W28wTXYjlWhEYR3zHMcl19SLEohXyQkAAAAzEBr0RwurKEz8j113/n4A5LDUWFobjK3ysrv54ZLFAwAiIhK1ZTTl8vIDwpWcrz/AdohAICF6hobZuhXosE12AK1SpUcFUN+3HlJeus/YkhIAAAAyADqUEOX3loQTXK4usaGsupKkiOSl2Z+r7bQSQFSRAIA4mKtyTYAD4tOAMRVWlURn55KOwoAEkInBZDfJFHX2JB9qsDKF0FCAgAAAECupo4eR3jE9M/yCI9ImLntrPXhUWqVSppYAED2Gn+4TTsEYB3qegGILjUvZ2/hUdpRiADb7KBnapVq76oE8uNGbkhsMVp18kJCAgAAwDxh/kFUxsXWY+hMrVLpw6NIjpial2Pl1SfLispLzW1nTb6TBwDICOaSAACoiNyQqPh9vQAcx/l6ehOu4svbmJtlzdORkAAAAJCHmlsG2iEAiyICphEe8dOzJwmPSMy2g/vMOj7MP4hWXxkAMNegBxxohwAAoHAxwXNph/BfM9+IU/AyGoA2hKv48hIy0635fCEhAQAAACBjbs4uhBfF/Cknk+RwxLQYjVmFx8x6ypKZsyUKBgBEh8I4wCakykBJkqNiAn10tKP4WV1jQ+jbr2GXOSje1NHjyHeS4Dhu97GDFj8XCQkAAAAAefv93HCSw5VWVShyC7y5l9Te7h6+nt4SBQMAokNCAthkb2dHOwQA0ahVqt1xeipzo13KLylOyt5DOwoAaalVquSoGPLjvvPxBxYn/JCQAAAAAJC3p8f7ER5Rka2t3/n4A7OOJ5wHAgAAAGCf0wCHg29uoB3FfyVkpheVl9KOAkBaoZMCyA9a19iQfarAsuciIQEAAAAgb/Z2doQr9qbm5Shs/7sF7aypXPcDAAAAMM7LbViuPol2FP8VsHqZobmJdhQAElKrVPrwKPLjxqenWvZEJCQAAADMc+m7KtohAHQU9VQw4RGPXjxHeERJmdvOWh8epVapJAoGAAAoqvhnDbGxwvyDiI0FQNJ03QQq06Pd0cUuVNhiGoAOIgKmkR+0rrHBsh1ISEgAAACYp7SqgnYIAB15uQ0j3Nr69d1pJIeTlAXtrKOnh0gUDAAA0HXlRjXtEACUYFVoBFMNrqM2rqMdBYCE3JxdCN8P8sxd18VDQgIAAEAe+v0KPQ+hJy+Q3SRRWlVRXV9LckTpmNvOOsw/yGmAgzSxAAAAZSRLu8wc50tsLADC1CpV9mtvs9PgOqvw2NYDn9COAkBCby2IJj9oVuGxFqPR3GchIQEAAGAGilt9MQEKPVsQNJPwiHsKjhAeUSLmtrNeMnO2RJEAgHROXy4jP+igBxzIDwpWyi8pJjaWg31/YmMBkGdvZ3cieSvtKP5rRdomNLgGBZs6ehyVcc1d3cUhIQEAAGCWmlsG2iEAdM3ezo5wKeqEzHQFVOM1t521t7uHr6e3dPEAgJKMHzGKdghgHguWeVpD9/BIksMBkOfm7MJUg+t5SXo0uAalUqtUMcFzyY9r7uouDgkJAAAAs/xw7y6VcamUgwTZIb9y/+zVS4RHFJ25ZU9/PzdcokgAQFIHzxXRDgFk4Frd98TG0jpqsP8VbAFTDa7rGhsWbEhQwJIagC6F+gWQH7SusaGsutKspyAhAQAAYAZzT7RiGfWQO5VxQV58Pb0JF+pN/GgXyeFEZ0E769BJFK7yAeSIVgqfKRNHetEOAcxD8koPJxSwHUw1uM4vKY5PT6UdBYAkxg8fRaVxS+6ZE2Ydj4QEAACAGSr+WUNlXMxogEB/+O1zJIfLLymW9bZ3cwue6sOj1CqVRMEAKAytFH53zM0+gm06e4Xczr+nxownNhYAXaw1uE7Ny9lbeJR2FADiU6tUVLLd2w/tN+t4JCQAAADMcPLrEirjDnV5kMq4IDvkW1unHc4lPKJYWk0mcwueRgRMkygYAJAU4cYAbdycXaiMCxbLPlVAbKzJj44mNhYAdaw1uI7ckFhdX0s7CgDx0araZNYHCgkJAAAAM+SXFFMZ13OIO5VxQXbs7ewI74jffmi/TOvwnr16yax21oE+OswtAsgUycYAbdipTwICGZqbzDovWCPQR2dvZ0dmLABGsNbgenL8UlrpagDpjB8+isq4ZjXrQkICAABAKIqlaTANCsKtnb+Q5HB1jQ0ybW1tbgMMwj9YABARlfpRfo/6kB8UrFH87WViY+GcAraJtQbXoW+/JtOFNQDdUatUVJZE7PwsT/jBSEgAAAAIRfI2tb0w/yAq44JMkW9ltu3gPpLDicLQ3GTWhieto8bX01u6eABAUiQbA7QJeGwM+UHBGp9dOEtsLForWAGoY63BdVL2HtpRdItW/0KQu8ippKv4chxXWlUhfMsREhIAAABCkbxNbW/mOF8q44JMqVWqxTOeITliVuEx2W14N7f1RXJUjESRAAABJBsDtPEZ+gj5QcFirSZTal4OmbH04VFqlYrMWACsYa3BdUJmelF5Ke0ounblRjXtEECWJo70ojJuybVvBB6pljQOAAAAJaEyncFxXJDPWCrjgnxFTw9JyEwnOeLuYweXzppDckRrtJpM2w/tN+spT4/3kyYWoKaovPQ7Q12Xf/X0eD/UdlcSko0B2qBDgOyQLD+4IiSM2FhgpRaj8dOzJ7v8q4ectNg9aRm+wfXwRax8EAJWL7uxO9dpgAPtQADEQavgc/bJAoHfikhIAAAACEJlOoPjOG93D1wcg7mcBjgE+uhI9mB/5+MPZJSQMLedtT48ChOLyrPt4L6swmNd/tXVHVn2dujcoxxUKi7+vyeQxZSZ7JOE1p3gnCIvN283RW5I7PKvwvyDkJCwmJuzS8G7WwJWL6MdyM9m6FeeSdmBrUugGGH+Qd1d6Eon+1RByqLlQo5EySYAAABBaDWQeOGpYCrjgtzFhvyO5HB1jQ3M7nbvbPm298w6PiJgmkSRAEWXvquiHQIQQqXiYqhfIPlBwWIk6zVhe4S8lF+voh2CYvl6erPT4Lq0qiJq4zraUQCIhkrZ57rGBoGFfJGQAAAAEGRT7t+ojIsZDbDM1NHjCI8ol9bW1fW1pVUVwo8P9NHR2vUM0mk1mcx6G4B8kZxoboPdjbJz9OI5MgNlxK3F9gh5Kf72Cu0QlGxNWCQ7Da6zCo9tPfAJ7SgAxMF4GwkkJAAAAHrXYjSSrH7TJtBHhxkNsIxapSK86Ewura33FBwx6/i18xdKFAlQdBldIm0GycYAbd5aEE1+ULDG67vTCIwS6KOb5z+VwEAgopNfl9AOQeE+fWM9Ow2uV6RtKquupB0FgAgGPeBAZdySSiQkAAAARCIwzy86zISCNcgXGuqu6yM7Wk0ms9p9ax01KA+tSCfKLtIOAQgp+OoC4RG1jhrye9TAGkXlpQS2TGkdNdmvvS31KCCuVpOJypokm6JWqYo37aIdxX/NfCPO0NxEOwoAa9nb2VFJ9f39C0H3g0hIAAAA9C7xIwpXyd7uHpgJBWu4Obt4u3uQHPFPOZkkh7OAuUU5kqNiJIoE6Nr5WR7tEIAEc3OQolg84xm0RZWRVpPJ3MZCltm7KgHFmmSHyhYrG+Q0wKHg3S20o/hZXWPDgg0JrSYT7UAArPUb78fJDyowiYuEBAAAQC8MzU1U1kZtXvIK+UFBYQiXDSmtqmB8n7u5RTmeHu8nUSRAUVl1JRpI2AhijQHaQ8tieck+VUDgC0EfHoVVJnJEZU2SbWKqwXV+SXF8eirtKACsRaWvNcdxQvYYISEBAADQi7TDueQHxfYIEMXkR0cTHjGd4YXn5raz1odHYTWrIrH8LgVxkWkM0N7G6Fh8b8iIobkpckOi1KPow6PWhEVKPQqIjtaaJJvFVIPr1Lycw8VnaEcBIEvffH+912OQkAAAAOhJi9FIvtoDx3E5a1BlGERgb2cXEzyX5IipeTnMtrY2t501+SYcQECL0Zial0M7CiCBTGOA9rSOmkXTQ0iOCNZoNZlm6FdKPQqyEfJFZU2SjWOqwXVIwqrq+lraUQBYbuJILyrjfmeo6/UYJCQAAAB6sjE3i/ygMcFz3ZxdyI8LihT1VDDhEU98fZHwiEKYm1wM9NHhY6hIS1OTaYcAhJBpDNBe2sur0D1CLlpNpqfffFXqlBWyEfJVXV9LZU2SjWOtwfXk+KXMrrMBYNbBc0W9HoOEBAAAsOv05TK6AVC5FdE6atBHF0Tk5TaMcGtr8jVShPj07Emzjl87f6FEkQBFZdWVWYXHaEcBJJDfHhHmHzRdN4HkiIpBfgknn42QuhrPxuhYZCPka8kWZK/pYK3Bdejbr9GOAsBCrgOdqIyLHhIAAACWazWZ5q6jcAF68M0NWF8J4vr93HCSw5VWVTC4w/1POZnCD9Y6atDERXlaTaaZb8TRjgJIaDEa5yXpSY6oddRsjYknOaKSON7/AMnhCGQjtI6a4s27ls6aI90QIKnDxWfQPYIi1hpcr8vKoB0FgCVozSoI+f5EQgIAANj1kJOW4uhJ2XsIL67kOG5jdKyX2zDCg4LiPT3ej/CI5nZrkFpZdaVZH2fsUlKkpOw9dY0NtKMAEpamJhP+Xe9dlYBe1rJQXV877IVQSeeaA3105ds+wuWcfBmam6L/nEQ7ClvHVIPrhMz0ovJS2lEAyEmvtc6QkAAAAHbR2mPIcVxReSn5Yk0xwXOxmA6kYG9nF+YfRHLEhMz0VpOJ5Ig9S/8sz6zjyadwQGqHi8+gGjhhtFop7i08Srgw18boWOypYl+rybQuK2P4ojDpklVaR02uPulQQgqyU/LVajIt2JCA7LVlxO1HzVSD64DVy4RUoQEA3s3bTT0fgIQEAABAR0XlpQGrlxEeNNBHh0XZIJ0lM2cTHvHoxXOER+xOi9GYmpcj/Hh9eBTmkhSmqLw0JGEV7SiAhKLy0sgNiSRH1IdHYTGBlRp/uC3p67eaTHsLjw57IVTSrGRM8NzybR+hj4iskWkuomC/8X5cxFdjrcH1DP1KplbbALCs5pah5wOQkAAAAPgftLIRn76xHq0jQDq+nt6EV5ltyv0byeF6YG4764iAaRJFAlRQ+VYHKsj/rsP8g9C12Hp//8K8b2nhDM1N67Iy7OcERm5IlG7Ne5h/0NUdWSmLliOZLWvIRjCIqQbXpVUVURvX0Y4CwDy0Sp99Z6jr+QAkJAAAAP5r64FPkI0ApfrDb58jOVx+STEje9vj01OFHxzoo3NzdpEuGCCJr9CCbISNIH8GD/TRpa9YQ3JEpcovKV6wIWFv4dHq+lrrFyC3mkxl1ZXrsjLGxkYNXhAi6a6IQB9d8eZdu+P0OHHInaG5acLKRchGMMjX03tjdCztKH6WVXhsb+FRkiP+o/RLksOB8jgNcKAdQtfUtAMAAABgQovRuDQ1mXDVaY7j9OFRq0IjkI0AAhYEzVyRtonkiGmHc6mvHS4qLzVrVeza+QulCwZIajEaQ99+DbNLtoDKGVwfHkX9+01JsgqPtf0GtY6a33g/PnGkl+P9D3i5Dbu/b79BDzj0sPmgur72h3t3y6orK/5Zc/LrEjKf+jD/oMTnopGHUIbDxWdQ1o9lS2fNOX25jPxtWpciNyR6uQ0j1rUe7UxAqZCQAAAAW9dqMmWfKiBccpqXEbd2nv9U8uOCbbK3swv00ZGcn91+aD/1fNu2g/uEH6x11KAzrQK0mkxJ2XvQwtpG7C08Gp+eSnjKBqdvSdU1NrTPTzBF66hZPOOZ6OkhzK45BbNU19cu2ZKM1DX70lesufRdVWlVBe1AOI7jZr4RV7xpF74EAKyBhAQAALDrh3t3JX19PhXxp5xM8le3WkfNwTc3EFtcA8BbO38hybvuusaGs1cvUZzibzEazZrSQmN5uWsxGncfO/jOxx9gRSELJD2Jt5pMZ69eWr7tPcJncJy+bVZM8NxQv4Dxw0dhV6sylFVXJmfvYTPvBZ2pVapDCSm62IUsnN/rGhsWbEhAxV0AayAhAQAA7Mo9c0Kie/6y6srcMye2H9pP5aI2JnhuclQMLmGBvPHDR2kdNSTf9okf7TqUkEJsuA52Hzto1vFPj/eTKBKQTnV9bc0tQ0nlN3//4iRWuTJFopO4obkp+2Q+lbRTmH9Q+oo1OH3bDm93jxeeCp7sNXrkYDf83uWu1WSquWUov15V/O2VfUXHGVlrz77q+lraIfzMaYDD3lUJjPSFyi8pTsreg8J9ABZDQgIAANiVkJm+r+i4v9fo8SNGTRzp5TrQyZq7wRaj8Vrd9xTzEBxWVgJtapVq8YxnSJay4Vtb09rV/s7HHwg/WB8e1UONchBR5IbEyA2JfJl4i18Ey1rZx5/EZ/tO0T08wnOIuzXV9vmZxIPninZ+lkdlGlHrqEl7edV03QTyQwMtWkfNbN8pjvc/cPtuS80tg5VXoWAZ688XhuYm5KoVg29wTbgjWncSMtN1D4/AeQHAMkhIAAAA00qrKkqrKlLzcvj/5G9IRgx28/i1q4N9f88h7vzjHaY5DM1Nd/9l5Diu/HpVU8udg+eKWKg6ujE6dtH0ENzNAl3R00MI19bPPpm/dNYckiPyzG1nHREwTbpgoDO+TDztKEBa/Em87T/bNyt+yEnrOtCJ47gO/Yr53APHceTbFHcH+xptU11jQ4fTpbe7B79KJshnLMrHk4TzBbRhqsF1SMKqqzuy0NwewAJISAAAgJzI9IZEHx61IiQMi6+BBU4DHAi3tn7n4w+oJCQSP9ol/OBAHx3uJwGkxnKz4i6F+QclPheNLwfgtV8lw3e3DpkwGdWcAAhjqsH15PillTuz8SUAYK4+tAMAAABQspjguVd3ZK0Ji0Q2AtgRG/I7ksPVNTYUlZeSHJEzv0TD2vkLpQsGAGQn0EdXvHnX7jg9shHQJX7/hG75Qvs5geuyMsqqK2lHBGAr+AbXWkcN7UA4juPqGhuefvNV2lEAyA8SEgAAAJLQh0c1ZB1JWbQccxnAmqmjxxEecdvBfYRHTDucK/xgraPG19NbumAAQEbC/IOu7sg6lJCChk8gEJ+ZGBsbdbj4TKvJRDscAOXjG1zTjuJn+SXF67IyaEcBIDNISAAAAIjJ290jV5/U8kk+dkUAs9QqlT48iuSIWYXHWoxGYsO1mkzbD+0XfnxyVIxksQCAPGgdNfrwqBu7c7ErAixTWlURkrBq2AuhewuPIi0BIDW+wTXtKH6WkJlOfjcwgKwhIQEAACACfiLj6o6s85vSp+smoJAoMI58A+fdxw4SG+vs1UtmtbN+eryfdMEAAOMCfXQF726p3Jm9JiwSnYrBSnWNDZEbEvm0BO1YABRu6aw5Yf5BtKP42bwkvaG5iXYUALKBptYAAACW83b3mO07JWTCZBR2AHlxc3bxdvcg2Q9w52d5xFpbm9XOWh8ehc1MADYo0EcXG/K7yY+OxjcA4wJ9dJ5D3MePGMVx3MSRXvyDrgOd2i/+aDEab95u4jjuh3t3y6orG3+4ffpyGcW2t3xa4k85mTlr3saGGwDpsNPguq6xYYZ+5ZmUHViXBiAEEhIAAABmC/MPmjnON8hnLJZSgny9tSA6JGEVseFKqyrKqisJpO6q62vNamdNfrMIANDCLyMIeGzM+OGjMGfELK2jZvGMZwIeG/PIg0MEXmjZ29nZ2/0878+faPgUeKvJdPlGdVl1ZcbRg2adGkRRWlUxfFFYTPDc5KgYvN8ApMA3uNbFLjRra6xESqsqojau2x2npx0IgAwgIQEAACBImH9Q+G+e0j08EkkIUIbJj44mPGL6Z3kpi5ZLPcqegiPCDw700WHtKoCyaR01oZMCQv0CfIY+gs0QjNM6apKjYkInBYg1fa9Wqbzchnm5DZvnP7XVZDp79VL2yYLUvBxRXlyg1Lyc7FMFB9/cgN20AFJwGuBw8M0NuuULaQfCcRyXVXhs5jjfef5TaQcCwDokJAAAAHqBpW2gSPZ2djHBc0nOy6Tm5SQ+Fy3phGCryZSQmS78+LXzmbh9tTX8JrO2/2yrwdJezS3Dd4a6Lp9+8FwR/4d/lH7JwopIYNnG6NhF00NwBpcFqS+31CqVr6e3r6d3clTM0YvnXt+dRqzMS11jg275wo3RscRKFyqGkPNF+fWqppY7nR/nK3fxf84qPCZRhMACL7dhGXFrIzck0g6E4zguckMinwelHQgA05CQAAAA6AWWtoFSRT0VTHih6KdnT0q6auzoxXPCD9Y6anw9vaULBrojZPGgm7NLd7+dDs/la8eXX6/67MLZwrKLLBSSBnasSNv09y9O7o7TY3cj4/i9EWRSR2qVarpuwnTdhKLy0uXb3iP2pcG/Gz99Yz0yZMIJPF9091dtGSC+io6huamuqaGsuvLguSKkKBRmnv/Us1cuEb6s7c7MN+KKN+3CeQegB31oBwAAACAD/NK2dVkZrSYT7VgAROPlNszb3YPkiH/KyZT09V/fnSb84OSoGOkiAWLs7ezcnF2m6yakLFp+flN6Q9aRjdGxhN/YwLL8kuLBC0KKyktpBwI9obIV1dfT+/ym9Fx9ktZRQ2bE/JLiCSsXGZqbyAwHHTgNcODrd+2O0/8r93jx5l0xwXNpBwWiSY6KCfTR0Y6C4ziurrFhwYYE3DYC9AA7JAAAgGmBPjqBq0sMzU1StytMyEw/+XVJ9mtvowg1KMbv54aT3OFeWlVRXV8rUduG6vpasxa6Pj3eT4owgC57O7uls+YsnTWnrLoyOXsP1sDSJfwkLnUBroDVy/ThUWvCIqUbAqwROimA1tDTdRMqd2bHp6eSWVtdWlWhi12I1dMs8HIblrJoeXJUTPapgvj0VNQAlDu1SrU7Ts9Ig+v8kuKk7D1WnnSQ0gAFQ0ICAACYtm1ZvMC5y1aTadgLoVJfgOaXFHsumY/bSFAM8pPyewqOSDQnaFY7a314FDKLyublNmx3nD7xueglW5KlTldDd4SfxIvKSwNWL5M0GH5VAQrmMEjrqKH7S1GrVCmLlj81ZnxIwioCw9U1NiAnwQ61SjXPf2ropACkJRSAqQbXCZnpuodHTNdNsPgVam4ZRIwHgCko2QQAAAqhVqn2rkogMFBdYwOKP4Bi2NvZhfkHkRwxITNdigVfLUajWe2sIwKmiR4DMMjN2eVQQkpG3FragUAvfD29CZTayC8pfvrNV7HmlDW/8X6cdggcx3HTdRNu7M4lU76Jv5hE7SZ28GmJ4k27GKn5AxbjG1zTjuJnIQmr8DEH6BISEgAAoBy+nt7EplYDVi9DTgKUYcnM2YRHNKv1tECfnj0p/OBAH51EZaOATfP8p17dkUWsTDxYZtuyeAKj5JcUD3shFDNE0CWnAQ7Fm3YR+67QxS7EW5EpTgMcPn1jvT48inYgYJV5/lPZ6Q6ii12ILDhAZ0hIAACAomx4cTmxsQJWL1uXlUFsOACJ+Hp6E56oNav1tEBmtcteO5+JvfxAkpuzS/GmXbSjgJ64ObuQmQfkC+ZgIhi6RDInwb8VW4xGAmOBQGqVak1YJDvT2WAZphpcP/3mq7SjAGAOEhIAAKAoTgMcNkbHEhsuITMdOQlQgD/89jmSw/GtrUV8wbLqSuHtrLWOGl9PbxFHB7lwGuBQ8O4W2lFAT1aEhJGcCEZOArrkNMDhRPJWMmPVNTaEvv0ambFAOHams8EyfINrRnZG5pcUbz3wCe0oAIjychvW8wFISAAAgNIsmh5C8uoTOQlQgAVBMwmPaFYD6l6lf5Yn/ODkqBgRhwZ58fX0xrpXltnb2RH7hGJxOvTAzdklV59EZqz8kmJcSbKGn86mHQVYhW9wTTuKn61I24Riv2BT7u/br+cDkJAAAAClIdbduk1CZjouMUHW7O3sCK8E3H5ov1gVdVuMxtS8HOHHPz3eT5RxQaYSn4tmZMkkdGme/1Rvdw8yY/GL01HdG7o0XTeBWGeyhMz0supKMmOBQIR3XYMUmGpwPS9Jj215AG2QkAAAAAUi2d2ahx7XIHeE2yrUNTacvXpJlJcyq521PjzK3s5OlHFBpuzt7AjXKANz5ax5m9hY+SXFT7/5KnIS0KWtMSQarfNmvhGH9yFrFk0PoR0CWIudBtd1jQ0z9CvxMQcb4TrQqecDkJAAAABlItndmhewehlWt4F8jR8+ivCy8cSPxOkwbFY764iAaaIMCrI2c5wv7RCgJ27OLiRXFeSXFMenpxIbDmTE3s6O2PLqusaGqI3ryIwFAqlVKnSSUAB2OoKUVlXgdAM2Qq1S9XwAEhIAAKBMVPZZz3wjDltxQabUKtXiGc+QHDG/pNj6z0tReanwdtaBPjo3ZxcrRwQFcHN2QdUmxpFcmc5xXGpeDjqOQpdCJwUQqyGWVXgMS1tYExvyO9ohgLWYanCdmpezt/Ao7SgApCUkBagmEAfIUX+nQe6+T1jwxEeCprT9+Ztjx816blXRF3cMNy0YFACgS4umh7zz8Qd1jQ3ERuQ7ZFbuzO51RQAAg6KnhyRkppMcMe1w7pqwSGteYdvBfcIPJlyWCli2eMYzhN/tYBZ+ZXrkhkRiI65I2+Qz7BFfT29iI4IsqFWqtxZEhySsIjPc8ylvnUnZgctIdkx+dDTtEEAETgMcTiRvHb4ojHYgHMdxkRsSvdyGebkNox0I2ISswmPkB3Ua4NDrMUhIKFxbXqGfxtH18cf4BwfrHm874AEXlz5qqS53vEKCrX+Rn1pNt2tr2z9y85sK4+3b/J/vNTXfKL7I/9lw9du68ivWjwgAisF3tw5YvYzkoHWNDU+/+eqhhBSSgwKIwmmAQ6CPLr+kmNiI2w/ttyYh0WI0Cr/I1jpqMNUIbUImTEZCgnGhkwLi01NJriqYl6Qv3rRLyF002JSpo8dpHTVk3oqlVRXZpwrm+U8lMBYIYW9nR/jSCCTi5uySq08illzs2cw34sq3fYSuZqBUE0d69XoMEhKyp/Uc4TT84bZ8g/OI4b/sb//Lvn37DWRiP5r1+qhVDoNd2z/S4T/HRT7b+VlNN2r4P7RlL2q+/OpuQ2PLrYZrJ09LFiwAMIfvbk14XUB+SfG6rAwr130DUBEb8juSd911jQ1F5aUW5wl2Hzso/ODkqBjLRgFFwspE9pFfVcB3HMX6dOiAL2lILIUZn54aOikAb0J2RE6diYSEMkzXTdCHR7GwHKGusSH07dc+fWM9PukgqRajkcq4Q10e7PUYJCRkwytkFvd/BZH4LQ6Sbm6Qu7akRdsfOmzXuHur4d/37v37Tkv9latt2yzKcg+QDRMASNjw4nLyGxUTMtMDHhuD5dggO1NHjyM84raD+yz+pLzz8QfCD356vJ9lowAALeRXFfAdR1MWLSc2IsjCipAwYpOYdY0N2CQBIJFVoREnvy5hIcOUX1KclL0HK9hAUjdvN1EZ13OIe6/HICHBHH7Hw2Dd6L4OA/jEQ4cNASCKfgM1/TiO4zjnkcO5/9tmMXvzeo7jmm7U8ImKhqrvblVcQ2cLALnju1uvSNtEeFxUfgA5UqtUhNeOZRUe2xoTb8Gm9aLyUuEFNPThUdgXDx2gCocsJD4XTXhVQWpezlNjxk/XTSA5KDCOcN0ebJJgykNOWtohgGjUKlX2a297LplPsh5gdxIy03UPj8DpBqTzw727VMZ1HejU6zFISNDEN3hoyz0oqc6SrPEZID5R0abpRs0PtXXNNd/XfPlV9ZlzaFYBIC/ku1tzHFfX2LBgQwK24oLsRARMI7yZffexg0tnzTH3WWa1s44ImGbu64PiIWEsC27OLuQrbIQkrLqxOxfvEGiPZN0ebJJgipCZNZARezs7dhpc43QDkiqrriQ/qLe7h5AJECQkyBnqN9HpEQ/Xxx8b4Prg/S5a7HuQF4fBrg6DXYeMHdNW+olPUdR+XX6j+CJqPQEwjkp3aw5bcUGe3JxdvN09SqsqiI34zscfmJuQMDQ3CV80Heijc3N2MT8uAGDCqtCI7Yf2E15VgCUF0EGQz1iSw/0pJxMJCQCJMNXgWhe7sHJnNk43IIXGH26TH9Tfa7SQw5CQkApfeemRoCnOI4b3dxqErQ/K05aiGBf57OzN6+/earh1rQr5CQBmUeluzXFcQmZ6yITJ6J4K8vLWgmiSN2l1jQ1l1ZVmfUzSDucKP3jt/IXmBwUArFCrVGkvryI8c5RfUrzjcK4Fm7dAqZwGOJDM1pdWVZh7ZgQA4ZhqcB21cd3uOD3tQECBTl8uIz/o+BGjhByGhIRovEJmDfQY+uBjXoMe8cDuBxvUb6Cm30BNW36i6UbNjeIvvzl2HC0oANhBpbs1x3Ez34jDsheQl8mPjiY8YvpnecK7yLaaTNsP7Rd4sNZRg/byAHI3XTeB8M4tjuNWpG2aOc4X+6ugzWzfKSTfhGadGQHAXOw0uM4qPDZxpFfnFHj59Soa4YBy/KP0S/KDThzpJeQwJCQs5xUya7ButMujngOHumMDBHTA75/g6zvdvdXw/Vdlpfs+xc4JALpodbfGsheQHXs7u5jgual5OcRGTM3LSXwuWmDf6bNXLwkv3pIcFWNFXKBkG15cnvhcNP/nfr+yQwVnxuWseZt8ye+56147k7IDSwqAp3t4BMnhzDozgnRcBzpd3ZHV/j/xnaAMTDW4XpG2yWfYIx3W0DS13KEVDygDlfe2wJUcfaSOQ0m8QmZNT1jzfM6euAsn11Zfmr15/bjIZ4eMHYNsBPSs30DNwwH+szevX1t9afGR/dMT1mg9iV7LAkCbRdNDtI4UvrSzCo8dLj5DflwAi0U9FUx4xE/PnhR4ZOJHu4S/7NPj/SwKB5TPaYCDm7ML/z9kI9jn5uwSEzyX8KClVRU7zCkQB8rmOcSd8Ignvr5IeEToTK1StZ0s3JxdkI1QEr7BNe0ofjYvSW9obqIdBShHdX0t+UHD/IMEHomERE+G+k2csnLZ/Pe3IQMBYnEeOXxc5LPRh/fFXTg5//1tQ/0m0o4IwLaoVaqDb26gMnT0n5NajEYqQwNYwMttmLe7B8kR/5STKeQwQ3OT8M31+vAorC0FUIy2HS0krUjbROWWHhhEvn7XZxfOEh4RwNbwDa5pR8FxHFfX2LBgQ0KryUQ7EFAIKiW/Zo7zFXgkEhL/o7/TIK+QWbM3J7986vM1FaURH+70j33p4QB/ZCBAdPy2iYgPd64uL0ZmAoAkL7dhwvP2IqprbFiamkx+XACL/X5uOMnh+AaevR5mVjvriIBpVkQEAGyxt7PLiFtLfty5614jPyiwKdBHR3K41LwczE4CSI1vcE07Co7juPyS4vj0VNpRgEIUf3uF/KACG0hwSEhwHKf1HDF+YQS/DeKV84WzN6/3Cgl2GOzaR42NeEDCff36tmUmUM0JgIytMfFUxs0qPFZUXkplaAALkC92lP5ZXs8HtJpMCZnpAl8t0EeHbrQAChM6KYDw5i2O40qrKvYWHiU8KLCJfHm3yzeqCY8IYINWhUYQTjd2JzUvB2ccEMXJr0sIj6h11Ai/+bLRhITWc8SUlcuez9mzurw4+vC+aX98DdsggLr7+vXlqzktPrLfK2QW7XAAlIzWEkuO4+Yl6bHSDeTC3s6O8HaiXpeCHr14TvirrZ2/0OqIAIAtapVq85JXyI8buSERdReBM6cYhVhOlF0kPCKADeIbXFPpNdhZ5IZElAoE6wkvciuW0EkBwg+2oYREhySEf+xLQ8aOua9fX9pxAXTkPHL47M3rV5cXT1m5jHYsAIpFZYklx3F1jQ3YhwsysmTmbMIj9pxyeH13msDX0TpqfD29xYgIANji6+lNpfQi6i4CFX//4iTtEABsgr2dXfGmXbSj+Nnk+KXIgoM1qOS0Qv2QkPg//Z0GjV8YgSQEyNF9/fr6x760pqJ09ubk/k6DaIcDoDRqler9la9TGTo1LwdrXkAufD29CS8W6yHlUF1fW1pVIfB1kqNiRIoIAJiz4cXl5AfNKjwmpM8NKJvw6thiIb/EFcBmOQ1wKHh3C+0oOI7j6hobQt9+rfGH27QDAbk6fbmM/KDjh48SfrAyExJeIbPaekJM++NrSEKAfPVRq7xCgmPPFCAtASA6Wt2tObTHBFn5w2+fIzlcaVVFdxm7PQVHhL8O+QYYAECM0wCHjdGx5Md9PuUt1F0E8rCQBYAYX09vdhpcv/PxB7SjALk6eK6I8IgxwXPVKjOaMSsnIaH1HDE9Yc3iI/vXVJTO3rwePSFASfi0xLLCIyjiBCAuWt2t0R4TZGRB0EzCI3aZeDCrnbU+PMrezk7UoACALYumh5Av9l1aVZF9qoDwoAA1twy0QwCwIWvCIhlpcF3X2EA7BJCrrMJjhEc0q14Tp4CEBL8Zgq/INC7yWeeRw/uozUjIAMgIX8Tp5VOfaz1H0I4FQCEodrdGe0yQC3s7O8J3ZQmZ6Z3XIJs1CRgRME3UiACAOWqVKu3lVeTHxenbxrkOdCI/6HeGOvKDAtiyT99Yz0iDawALkN9Xp3XUmFWviZNpQqK/06ApK5e13wyBikxgOxwGu76Ylz09YQ3tQAAUglZ3aw7tMUE+1s5fSHjEs1cvdXjkTzmZAp8b6KNzc3YROyIAYM503QQqi1hx+rZlZtWjEMvZKx3PiQAgKbVKxU6DawBzka/XtHjGM+aeH9UShSIFreeIx+eHPho8A7WYlKFzyq78elVTyx1zX8fBvr/nEPf2j7gOdKJypUhMH7VqXOSzbk+M+zDixTuGm7TDAZA3vru1bjnp+VaO47IKj8WHRni5DSM/NIBZxg8fpXXUkNw2nvjRrkMJKW3/WVZdKbydNfn0CQDQsm1Z/PBFYYQHzSo8lvhcNBKfQMzN2020QwCwOXyD64DVKJoN8rPzszzCI1qwPV0GCYmhfhMnvBj54GNeyEOwr8Vo5K+Wam4Z+I2ljT/cbuvtfum7KuGzCWJp61g7c5wvx3EPOWldBzr1+5Wd0wAHwpGIznnk8MVH9u959oW68iu0YwGQN767NfkyixzHPZ/y1pmUHcrOoYICqFWqxTOeEd7CwXr5JcWG5qa2k3W64KtqraPG19NbqrAAgDFuzi4xwXNT83IIjzt33WvnN5H7SgQbZ2huoh0CgC3iG1yTvAAGsF6L0Uh46tXb3cOCVRrsJiT4PITbE+NQjokprSYT31OLTzNU/LPmyo1qjka/FIHaAuscodZR8xvvx0cMdvP4tauX2zCtg0Z2WYp+AzUv5mVnRkZfO3madiwA8rY1Jp7K9xjfHnOe/1TyQwOYJXp6COH7sbTDuWvCIjmOazEahc82JkfFSBkUADAn8bno7FMFhDt/llZV7C08itM3kJFfUkw7BAAbtSYs8uTXJfgMgoyc+Poi4RHfWhBtwbOYS0ggD8EIPvHAb3Tgsw6G5iaFfQvXNTZ0nn8M9NH5Peqje3iE7uGRsshP9FGrwjPSNk0IQO0mAGvw3a0jNySSHzo+PfXp8X72dnbkhwYQzmmAQ6CPjuSVwPZD+1eFRqhVqk/PnhT+rKfH+0kXEgAwyN7OLjkqhvwZPD49NXRSAPY4AoCylV+voh0CZZ++sX7YC6GE094AFtuU+zeSw2kdNVNHj7PgiawkJJCHoIivs8T3bzh4rkh5iQez5JcUt/3ztY6axTOeWRESxvhEYR+1KuLTrBeS3+z7b9Pge6a2x9Xf17dWVPF/7tf60y9/ohMegFyETgr4U04m+cpydY0Naz9IS1m0nPC4AOaKDfkdySuEusaGs1cv+Xp6C29nrQ+PYvyUDQBSoHIGr2tsSMrew2/kAgBQKgvafCoM3+B68IIQ2oEA9I78jO4ffvucZYszKCck0KeavOr62h/u3S2rrjx75dLN203MllpiQV1jQ0JmekJmuj48il+kSTuibjn9+tczpj61Im3T/zz6K44bNbDzwdOGeXIcZ/r3v0fe4ziOa62tV39TxXHcA//+qY/0oQIwi2J369S8nNiQ36E9JjDOssUv1th2cB/HccInGS1opwYACkDrDJ6QmR49PUQWm6pBLNX1tbRDAGDXP0q/pB2CJNDgGuQi+2Q+4REXBM207Il0EhL9nQbpnpv32NwQh8GuVAKwHYbmpm++v/6doe7guSIqPaWVISEzffuh/XtXJbDcJ3PR9JCdn+UJ+RUfqSzn/3CU/7/7/idvMbBf/7EuQ+7/RR+XJuO/79z55ZeXOI5z+De2V4BNoNjdGu0xgX1qlYpwZ7+swmPCO3kG+uiQ1QOwWbTO4HF/3bw7Tk94ULA1Yf5BtEMAEETBdY3Q4Bpk4Z2PPyA5nDXb00knJLxCZo1dMN91tE8fNbuLzeWLL750+nLZ2SuXyq9X2XLlJdHVNTYErF5W8O4WZnMSapXq93PDra+fe+vunbaMBcf9T65i2jBPvipUa/m36psN2FEBikSxu/Xh4jPTdRPIDw0gXETANMJ3YsIvZtbOp7C9CQDYQeUMnlV4LD40wsttGOFxAQCAMDS4BsaVVVcSTgquCAmz+LmEEhJazxG+S18Y8WQQWkSIy9DcVNfUUFZdefBc0T9Kv1RwOpoRAauX3didy+y+bKk7ef43UeH8C855IMdxg+zv12kHD7z3Y/+GH3755SV0qgAFoNjdOvrPSZU7s1muDgfg5uzi7e7B4IZLraOG2RUDAECGvZ3dxujYjiVMpfd8ylvY42g70OAXwJZ9+sb6CSsXMXglDMBxXHL2HpLDWdm9T/KExJSVy1CaSUR8CaaSym9OXy5D+wcqir+9zOwSZvKdPG+2/PDfLMX/7aWYNszz/lu3B928rf6mCrsoQI4odrdGe0xg31sLokMSVtGOoqPkqBjaIQAAfYumh7zz8QeEF2lhj6NNodLgd8RgN/KDAkBnapXqUEKKLnYhVgMDawzNpJsEW7M9gpMuIaH1HBG46pVhkyehNJOVWozGa3XfYw8EOz67cBb3Gz37OUXxf60pBva1f/RXDww13P5VxXWHf5uwhQLYR7G7NdpjAvsmPzqadghdkHqPIADIglql2rsqgXzfUexxtB0HzxWRH9Tj11jfCcAKpwEOVE40AD1LO5xLcjgrt0dwUiQkxi+MeOLFSGyJsEZ1fS3fB6Kw7CL2grEm1C+Adggyc+teS+G9lsL7OG6khvu//MSD175/oK5xkNGE/RPAJordrdEeExhnb2cXEzw3NS+HdiD/Zf0FMYDoispLvzPUtX/kISctCosR4OvpHeijI1zju66xYcfh3KWz5pAcFGzHxJFetEMAqZRVV5ZVV3Z4cJ7/VCrBgEC+nt5UKgQCdKfFaCTZ50/rqFkVGmHli4iWkOjvNOjJtfHoEmGZFqOx5No3JZXf/P2Lk+iQwzKto2b88FG0o5A3Pj/BOf6Cc9RwHOcxYKDX3Z8cLlVqjCZ7039oRwfwX7S6W2cVHlsyczYmrYBlUU8FM5WQiAiYRjsEgP/RYjR2XjsZEzwX3+1kbFsWP3yRVWUELLAibVOoXyD2OCoelYvDQQ84kB8UyJj5RlyHMhhh/kFISLBv6aw5qKMO7NiYm0VyuLSXV1m/K1SEhMRQv4lTXokZMnaM9S9lUwzNTcXfXv7swllsg5CRvasSWN6LXVReSjsEs1U036rgOO4RR47jNHb9Hm+9z+nKd7+++yOSE0Adxe7W85L0qPwALPNyG8ZOa+tAH52bswvtKAD+x6dnT3Z+cPwILGohxM3ZRR8eRXKpIA97HBXP0NxEZVzsAlSqovLSzkW50TJELtJXrLn0XRUj18Ngywhvj/B29xCljr1VCQk0rDZXWy2m7FMF6AYhOwXvbmF8Xdvybe/RDsEqDca7xziO8xjAcZzHgIHDb7U4X/un9l4r2k4ALRS7W2efKsDaKGDZ7+eGU0nXdbZ2PoV2LwA9aDWZ4tNTOz/u5TaMfDA2a0VI2PZD+wnfcGUVHosPjcAvWsG++f46+UHD/IPIDwpkbDu4r/ODuodHkI8ELIAG18AIwtsjcta8LcrrWJKQ6O80SPfcvImLFqI6kxB8EuLguSJs5pIvraPm4JsbGL+7KCovVVJyvqL5VoX6550TPv01bjduPfR9g8O/kZoAoih2t47ckPj0eD8siANmMdJEWuuoYXytANigs1cvdTk3MVT7IPlgbJa9nV1yVAz5vOnzKW+d30R6ZwYQU1L5DflB0UBCqarra7ucI/Ic4k48FrAQGlwDddX1tSS3R2yMjhVrb7p5DWX7Ow2avTk59kyBf+xLyEb0wNDcdLj4zIINCb8KmTJ8UVjkhkRkI+QrJnhu5c5sxrMRrSaT3LdH9KDkTsPfHX6xZdTAv0xwK/Yact1ejbwEEMN3t6Yy9NLUZCrjAghhb2fHwprN5KgY2iEAdNTdJRlyzITN85/q7e5BeNDSqoq9hUcJDwrE/P2LLqqxSW3mOF/ygwIBaz9I6/Jx14FOhCMBa/ANrmlHAbZryRZykwbe7h6LpoeI9WpCd0hoPUfMeGut62ifPmqUtO5aq8l0+UZ17pkT+4qOK2mhui3zdvfYvOQVWSy9jNq4zhbedQ3Gu4f+b9tEkHrA4PIq15ZW89KqAObbGhP/j9IvyW/FReUHYNySmbOpr7dgZKMGQJu9hUe7vCRjIYFng3LWvE2+u3V8emropAA0glKeVpMpv6SY8KBaRw36JClSd9sjtI4afHvIDhpcAy2Hi8+QPDHlrHlbxC+o3hMSfCoCPau702I0nvj64mcXzqbm5dCOBUSjddQkR8XIpYD7uqwMGzz5HWttRmYCyKBV9oHjuOdT3jqTsgO3JcAmX09vraOGYtlcfXgUlpwDU7rrHsGhSSklbs4uYf5BhK+T6xobkrL3rAmLJDkoEHD26iXyg4ZOCiA/KBDQ3faI33g/TjgSEAUaXAN5LUZj9J+TiA0nYrEmXi8zeCOeDIo+vA/ZiM74okwz9Cs1YdNCElYhG6EYWkdNRtzayp3ZMspGkCwYx6Bjrc0Zjzhu9x1a7DWk6ZfISoAkqJR94DiutKoi+1QB+XEBBPrDb5+jOHpEwDSKowN0tuNwbncpOo9fuxIOBnhbY+LJD5qQmV5dX0t+XJBUwVcXyA8a6oeEhAIVlZd2lyhFyxCZ4htcax01tAMBG7I0NZnYyrBAH93SWXPEfc1eJu/QKKIDQ3PT1gOfjI2NGrwgJCRhFfk9myAdb3ePgne38KkIWaxHbjWZVu7YbOPZiDa37t45pDZuGTXwoykjv3Ub9G8kJkBsOWvepjJu5IZEQ3MTlaEBerUgaCatoQN9dKhiAUwxNDetSNvU3d9ijokWezu7jLi15MclWdMZyNhXdJzwiFpHzfjhowgPClJrNZnmJem7+9uhLg+SDAZExDe4ph0F2IrDxWeIbQDVOmp2x3X7rWUxTNoJ0j4PsSJtE/ZhKUxM8NzizbvOb0r39fSWRSqC4zhDc9OElYuwNaeziuZbex1/kfyY0+Upo7FhAkTk5uwSEzyXytBxf91MZVyAXtnb2QX66KgMvXb+QirjAnRnwYaepiHQpJSi0EkB5Fet5pcUHy4+Q3hQkI6huYn8JMDiGc/I5eYUhOthLx3HcZ5D3AnGAiJDg2sgw9DcRLJY095VCU4DHER/WczW9aTFaOTrMiEPoUje7h4ZcWsbso6kLFour7axh4vP6GIX4g3Zs+zmmi2jBh5+cvTNBwfRjgUUIvG5aCr7cLMKj2FSA5hFJTGgddT4enqTHxegO3sLj/awcxpNSulSq1RUVq1G/zmp1WQiPy5IIftkPvlBV4SQbskOUquur+1hLx3HcQzu/jx7hUL3FPlaOmtOmH8Q7ShAyVpNpgUbEogVa9oYHSvRbRcSEl1oNZmKykvb+kOgLpPCaB01+vAofkvEPP+p8uqH2WI0LtiQEJKwimITUXk5b6jZ5vwL2lGAQvDdrakMHf3npBajkcrQAD0bP3wU+UQdrU8iQJcMzU2RGxJ7OABNSqnz9fQmP0NU19jQXZNzkJ2dn+URHjHMP0heN6rQq1aTaXL80h4OoNKyrlc3bzfRDkFm0lesobWBGGxBUvYeYtPUMcFzRW8d0QYJif9RVl25LivDfk5gwOplyEMoT0zwXL5LxJqwSHltieA4rtVk2lt4VBM2jVidOADojFZ367rGhrUfpJEfF6BXapVq8YxnCA/69Hg/wiOCuA6eK6IdgmhaTaYZ+pU9H4MGEizY8OJy8oOm5uWUVVeSHxfEVV1fS35venxoBOERQWrx6ak9Lysc9ZA7qVhAQmqVanecHg2uQQpbD3xCrI9soI9O0kVgSEhwHMe1GI18iwjd8oVoEaw8fB6i5ZP8lEXLZdQlor2i8tIJKxf1vPgOAMig1d0akxrArOjpISSH04dHYdEosCNq47peZyrRpJQFTgMc9OFR5Md9PuUtFG6Su025fyM8Yph/kOzWz0HP9hYe7bUB5MxxvmSCMYuhucms4/GNx3Gc0wCHg29uoB0FKE1ReWnPNd9EpHXUfPrGekmnT206IcGXZlqwIUETNg0tIhSGr8sk9zwEx3HV9bVjY6MCVi/D+xOAERS7W898Iw6Fm4BBTgMcSO5MjwiYRmwsgJ6ty8oQsnUVTUoZsSo0gvyS1dKqih2HcwkPCiJqMRp7nUcW3daYeMIjMqvinzW0QxBBUXmpkMWFDzlpCQRjLnNrh9TcMkgUibx4uQ3LiFtLOwpQjqLy0oDVy8iMpXXUFG/aJfUkqlrSV2eWobkp7XDu9kP7UYhfYQJ9dP/vCb+Z43wZbAZlrur62rUfpKFAEwCDEp+Lzj5VQP4MUtfYsDQ1eXecnvC4AL2KDfkdmVqXgT46BZzimfKP0i9phyBXReWlArdW403LCLVKlfbyqpCEVYTHXZG2abLXaCx4l6mNuVmER8yIW4uNgG2u3KimHYK1DM1N85IEXb27DnSSOhggaZ7/1LNXLpHPaILyVNfXEs5GOA1wkHogm0tIFJWXbju4D5O8SqJ11IROCgj1C/AZ+ogyLt2KyksTP9qFLiYAzOK7W1OpopZVeGzmON95/lPJDw3Qg6mjx5EZaO38hWQGsh1YnWMZ4evU2GxSarOm6yZ4u3uQ33k88424yp3ZMt2xbctajEbCJZ293T1CJwWQHBEkZWhu0sUuFHiqZTB7bW69JuggOSqm/HoV5nbAGkXlpQKTmtYjlo3gbCch0WI07j528J2PP8BNlzJ4u3uMesh95jjfIJ+xZD4qBLSaTEcvnnt9dxqqMwGwb57/1D/lZFL5tEZuSFTSVx8og1ql0odHST1ro3XU+Hp6SzoEkCH3CQ6zds2jSSlrcta8PXxRGOFB6xob4tNTUxZRaKwN1lj7QRrhEd9f+TqziauzVy7RDkFmzMpGkKx+Kdzdf6FarFX4BtfC3wYAHZCv1ERsnkH5CYnq+tpNuX/DJim5C/MPmjjSa6jLg55D3F0HOjF7lWYZJMwA5IjKjAZvhn7lmZQdCvsmBLmLCJgmdULiD799TtLXB2JkvVTQ3DtDNpuU2jK+FxT528PUvJynxoyfrptAeFywWFl1JeH3SUbcWpZLe9283UR+0EvfVZEfVBRmZSM4VrsNlV+voh2C7PENrnXLsccXzHa4+AyxOpOEsxGcshMSReWly7e9h8XmctS2AcLLbdhQ7YPKKMTUJdQQA5AvWjMaHMeVVlVEbVynjGYS1fW1tMZlcF+8rLk5u0hdCGVB0EzpXhxAiHVZGeYm3thsUmrjEp+LpnL6DklYdWN3LvY4ykKryTTzjTiSI8YEz0VNzs5kOp/Dl1gxa7nh+BGjpIvHYtdqvzf3Kacvl+EauwO+wTWVer8gXxZcc1os0EeX/drbhKdeFZiQaDWZsk8VxKenYrG5jAT66DyHuI8fMWriSC/lbYDoDFsiAJSBVndrjuOyCo+NGOy2JiyS/NAA3XlrQbR0q3higucqeIGCDWoxGuX1C201meLTUy2YxX7kwSFSxAPWsLezozU3pItdKKNmEo0/3KYdAjVRG9eRvMAL9NElR8UQGw4kZVmJFTY3x5y+XGbuUyr+WSNFJHKHBtcgXIvRGPr2a8T2Ewf66D59Yz35KxNFJSRajMaNuVmEu06BZQJ9dH6P+ugeHuE5xN128ud8l4hNuX+TdaUCAGhDsbs1x3EJmekBj41BSX1gx+RHR0v34rEhv5PuxW0WxV4ON2832dvJ5grQ0Ny0YEOCZddvWA7PptBJAVR6QdU1Njz95quHElIIj2sZC+YiLZNVeCx9xRp28jTrsjJIbmHXOmp2x+nZ+eezxtDcJJcv0laTKSl7j2VTUloHjejxWO8fpV+a+5SEzPRVoRF4P3eGBtddsuA9pmxl1ZUz34gjlhHXh0fRWuOokIREdX3tnoIjSEWwzNvdw99rtO3sgWiv1WQ6e/VS9skC5MMBlIdid2uO4wJWL7u6I0vWad0f7t2lMm7NLYOsf25ssrezk6iOmbe7B35fUqDYrFJGJR0OF5+J/nOSZXeGbDYpBY7j1CrV+ytfp1LUO7+keF1WBvY4dnD26iVG1liQrJLB0SjbbTFa84bffH9dFj8fa1LXHJPZa0Nzk2XnvuxTBag/1hkaXHeJ8E9jT8ERZs+/Fm/GtVjBu1sonnn70BpYLNX1tQs2JAxfFIZsBIPC/IMy4tYWb97VkHXk/Kb0lEXL5/lPdXN2sZFsRKvJVFReunLHZvs5gQGrlyEbAWCZmlsG2iH0ImfN2xRHnxy/lOIaZ+uVVVdSGfc7Qx2VcRUv6qlgKV5285JXpHhZoPgFe/bKJVpDC9diNK7csTkkYZXFd8tsNiklieWTuJfbsDD/ICpDJ2Smr8vKoDK0WUj2E84+WUBsrB4gG9EDWrOoJZXfUBnXLHsLjw5eEGJxNoLN7HX2yXzLnhifntpipLbigWV8g2vaUdi0hMx0Nu+dDxefGfZCKLFpQ62j5uqOLLrrAGSckGhLRaAhMDu83T304VG5+qQbu3P/lXt8d5x+nv9UL7dh8qoRbCXkIUDZMo4eJDzi8m3vER7RXHx3a1qj1zU26GIXsnldJQStOrOobysRL7dh3u4e4r6m1lEzfjiLnR4VgGJmLjUvh/HZir2FRz2XzLfyQo7BJqWEz+OMn8Q3vLic1tCyyEmQ3ACampdTXV9LbLjOWk0mZCN6QPG3887HH7B8vqiur52hX2llBVe/R33EikdEf//ipGVPrGtsWPtBmrjBKAbf4Jp2FKyg8sWyYENCq8lEftzu8JPb1ix/MVdM8NzKndnUdyrLMiGBVARTAn10G6Nj27ZBrAmLnK6bIJfrKhG1GI2Hi88gDwHKVl1fS77qZWlVBd0bVCESn4vWOlIr/CrrnMSVG9U2Na4t+P3ccHFf8A+/fc5G9laSR3ebArOzFUXlpWNjoyI3JFp/Z8hak1JDcxPh8zjjJ3GnAQ4bo2NpjZ6Qmb71wCe0Ru8V+euKyfFLaU0SGZqbJqxcRDIbEeijq9yZLaO7Zoq7neoaGzbmZtEavQctRuO6rIzhi8Ks/171+LWrKCGJyMr7vtS8nKLyUhHjUZJ5/lMprmZjCpXivfklxU+/+SoLOQm+zhvJyW2toyZXn5SyaDkL91YyS0ggFcEIvhbT1R1ZLZ/kH0pIWTprjq1tg2hjaG7aW3h0hn6lJmxaSMIq5CFA2WhNHjE7adWG725NMQA+J0Gr/JE1aJ3QcSEhnafH+4n7gguCZor7gsBrMRrpXrdQXxDdWVF56Qz9yoDVy8RaGD5U+6AoryOWuL9uJj/onoIj5AcVbtH0EIpLClakbWI2J/HN99cJj8h3/Ca8Fr7VZNp64JPBC0JIbgfZGB17KCGFhfkg4bYd3EdxdNaqrPCpCE3YNLGSWBNHeonyOiKy/v5rXpKeqd8aU5KjYtis00UYrbvX/JLiYS+EUrwKra6vXZeVMXhBCMl70jD/oOJNu6brJhAbsWeySUggFUFdWxKirRaT7XSD6IAvyrQuK+Oh52cPXhASuSGR/JpxAPIOF5+hOHfM/lT7PP+poleqMUtdY4Nu+cK9hUcpxmAuuncpuEeSiL2dnYhl2WOC59rmigcCPj1rYSkGETHSBYe/tBsbGxWwepm4F3VMvXuLykupnMdZm0nsQK1Spb28imIAK9I2rdyxmYWlmh0UfHWB/KD5JcWeS+YfLj5DYKxWk2lv4dFhL4SuSNtEYDie1lFT8O6WpbPmEBtRFC1GI/V5GEZ2A/PTiCKmIniDHnAQ8dWsV11fa/1vvK6xgbXaOOxQq1TZr71NMR3OiIPnimgNXdfYMHxR2LqsDMJZcH7tC+FGyPzGiN1xeqa25f3iP//5D+0YelFdX7v2gzTq5z/bFOiji5w6c+JIL+rFxVhQXV97+nJZxtGDSD/Izr9yj9MOQfaq62uHLwqjGIAsyuxS/ynx9OFRq0IjZJEw3lt41MqSu9bIiFs7z38qrdGVrai8NGD1MlFe6uqOLFyESGRsbBTJFcHdofv1XlZdmXvmhET3hGH+Qbvj9FK8sgUMzU2DF4TQGl3rqCnf9hFT6ZkOqH8cAn102a+9zc6PqNVkmrByEcWfSZh/UOJz0RJ9/7cYjZ+ePRmfnkq4S3Ogj461+SCBth74hGTapjsUzxctRuOJry++vjtNog8FU7er4n78A310n76xXhY3JuSJePPI1CWHQC1GoyZsGu0oOK2jZvGMZ0ImTJa0zCZ/wbn90H7C5x2O42KC5yZHxTD4GWQ6IdFiNC5NTUYqgjBvd48Xngqe7DV65GA3Bt+yhBmam4q/vfzZhbOoxSRrTF3hyVFReem8JD35c2cHWkdN5c5sxr+XVu7YzMLXBWvzGl1qMRo9l8yn+L6SxTtKvh56frb1v1xvd4/zm8itHrIpZdWVuuULaUfxXxlxa0MnBZD5PLYYjSXXvin46oLUt4X68Kg1YZHSvb5whuYmXexCuudxxuekWFhSoHXUnEjeykgKlu6KgTZh/kHxoRFi3Zbyn/1tB/dRmWGQ7zKIVpNp2Auh1G8EeFpHzR9++9yCoJlkrnLJLElkbSpZ9NuZMP+g9BVrmP3+p+tw8ZmQBBF26bH2LhJiXVYGyV0CveK/XiZ7jR6qfVCUb5i26cTsUwVUvkIlTe1bj9GERIvRuDE3i6m3prJpHTWhkwJC/QJ8hj7C+AQWAdS/NUB0SEhYQ6yLJFGwv7KM+orC9khO8FmAhWtQdqYLlUeU1ZQF727x9fQWJR5oj3o6sEttK9TEug9s02I03rzdVH69ivDVHSPzjyLuWLIS4zkJRtaAs7DNkcGviEAf3f97ws/x/gcectK6DnTiH+z3K7sergn5D37NLcN3hrqKf9bsKzpO6/KM/cvXnjGy2qYDcZNVbarra3+4d/dE2cW/f3GSWF0Edi5HW02m+PRUKX7djH//0yXKPZHsEhKMbI/oAV+BdsRgN3N7zp+9cunm7Sa6a+u93T02L3mF8dso5hISrSZT9qkC8tsnbRN/YTdznC+zGTNiqutryd+mAjFISFiGzW1qWkdN2sur2OnF1BndshgdeLt75Kx5m8EveXauQRuyjiATLwXrf8XYwiKRVpPp6TdfZb/+pPBOJDPH+XZ+kK9KbGhuovgvpV5wrNVkSsreQz31257WUXPwzQ2SVkWwxgz9ShY+GtR/SiysGFAMRhKTFmNkr0zPBJ4vupxYbPzh9unLZRzH/aP0S1rzALn6JBZubVqMxtC3X5PuO1AWG7ipEOXCTHYJCZxoJKJ11OxdlcB4KoLHVkKCkaogihcTPBebIVpNpss3qsuqKw+eK2JtvhVEh4SEuVpNph2Hc1lYJ9gdxrcfsrMclRcTPDfqqWB2JoDKqitnvhHHyOmeqRIZCmPl1N7G6FjZtf2UBTbXuipSyyf5FDNqh4vPRP85iZFv2g5igucmPhfN4J0IC7Wt2vALwAmfu9m/ApSRMP+grTHxDL7PhWPtglapWMhek/ngs7+wjBbr96XJKyHByJZEheF3RYwfPkouy7lYSUhU19cu2ZLMwoIUpfJ295jtO0XqPi2M47dBFH97heKOXaACCQnhWozG3ccOvvPxB4zcjfcs0Ee3dv5CNk+6TJW64vHlUAIeG0PxJ8bsTIfcFzCyycqJDGxeEZ10pRigM62j5rv395Efl+/W+6ecTPavdWOC58aG/I61fDBTOQmuXSkzAu39qutr5657jf13Dvu83T3eX/m63O+7WU5qKgzF7HV1fe2egiOEG+0ym5Omy8pWRnJJSEi9Ecc28bMistgV0R79hESL0bj2gzTcGkmE71Ad6hco35qV1jA0N33z/fWCry6c/LoEX3m2DAmJXrUYjSe+vrgp928y/aSE+QdNHOk12Wu01kHDztcdyxtRw/yDZo7znTjSq20myNDc5Nj/fuluh6rraw+eK9r5WR6zMx02fsaUgjVtMGOC56YsWi56SLYMU0uEEZ4aaDWZzl69lH2yQHZ3VXwru6fGjPcc4j7oAQcWpqhYaHDdpZjgueNHjJo40st1oBN/vjY0N1l/zpLvm4dBWkdNclQMyx28hDA0N8X9dTOqCJBBPnvN14o4UXaR7mW5PjxqRUgYC9/57LBmQRv7CYmy6srcMycIZ78UTx8eFREwjbW1HQJRTkjsLTyKdhFSCPTRxYb8bvKjo23w+726vvb05bKD54ooVoEE1iAh0VmryVRzy6DsPUMsXJbJovCut7sHx3GlVRUibhg3NDfd/Zfx9OUyvjKv7O5p+YSNg31/zyHuPffMhJ5ZnJajXr5AMVqMxmt13ydn75Hdx1DuCNQcMzQ31TU1KLUAKd0vAaZKC3bJ292jvrnxN96PW3apwzd8Pn25LOPoQZkuRmGNMlIR/K00+9euSkLghoW/7zt9uezslUvl16uY+si3rSojsAlMFiy+cmbhzrc9/maw/HrVtdrvSbaItxH8lO/U0eNk/alR0xoYG0KlYJt5iLYMP77mAARidumf8szzn+rlNozxSQ2xzsWslbmwRlbhsc6ze1pHTfGmXUhOmCV6eogFt1Xe7h7IRliAr9XT9p9YnEGXz7BHJH39BRsSlJeEYIeX27DiTbsWbEhg9ubC3HM3v+ad47hL31XhHlxc3u4ev58bLrtURFF56XeGOv7PFf+sQUUBWiaO9JJ6iJpbBmZv/douubEShbcqNELWH8axsVE4xUhHYRv6KSQkUL5WdN7uHm8tiLapPASSEAAgF15uw8q3fWQLtTLv/suo7KnPusaGu/8y0o5CZpwGOAT66Mx9829e8opE8SjbzdtNWNbKDteBTrRDAKs4DXD49I31irlvvfsvIzJYoosJnhv1VLBMe0VsO7gPbwkWDHV5kHYIwBC1SpX92ttWNrimCNkIKSgsD9GGdEIC5WtFpNQ3ZQ/4PaTYWQwA8mJvZ3coIaWovHRekh5nQLA1sSG/M+usrXXUjB8+Srp4AMjASk8FUKtUKYuWh/oFLN/2HiZZoE2Yf1D4b56Se60MYITnEHfaIQBb7O3sTiRvZXZTCxATEzz3qTHjFbz0nFxCosVoXJqajCS89bSOmsUzngmZMFmmazHM1WI0llz7JvtkQfapAkzkAYB8+Xp6V+7Mzj5VgOZJYFOmjh5n1vF/+O1zmOIBueN784Ay+Hp6n9+UXlReirSELePXAqLSPYgO2WvozM3ZJVefZHGDa5AvvpHhxJFetvDNQCghgY0RogjzD1oyc/b44aNs4RrI0NyUfTIfFZkAQEnUKtU8/6mhkwLOXr2UfbJAGVUgpBPmH2T9ixiam3AeoUutUunDo4R3klgQNFPSeAAIGPWQO+0QQGS+nt5nUnacvXqp4KsL+4qOIzOhbFpHzW+8Hx8x2M3j164TR3q5DnSyhRtwII9i9rrzZfagBxzGjxBti+rBc0VdPo42NgJN100w6/pZYcy9DZw40svx/gf4P5+9cunm7SYLBqXSei3QR+f3qI/u4RG6h0faTvEb3i/+85//SDqAobmJ5W5gssAvx1gQNFOp+3TaK6uuzD1zYvuh/UhfgYj+lXucdghs6dD4VKkectL6enrTjqInrSZTzS3D6ctljT/cPn25rP1ftb+oIubp8X7WnGg6vK+83Ibd37dfh2NYvqXnfx0dHqy5ZWhr+chZ/SOyWYbmpmMl53s4oO2udeJIr6Wz5hAJSoHamtYCdUtmzpb6BNS+Ia0iMf5922I0nvj6YvG3V67cqG7/+MxxviTDEH6p0/77ofM1xkNO2p67ngx6wIHAr6O6vrbXYzqcl9tU/LOmw+9CCD7lwP+Zv27p9ys7G5kS2nrgkw4Xn0AemcuetktcBtdctxiN/Nwxy/cIFK3csbn93Hr7ryzuf2feGbmE3lt4tO3PnW8GlfEFa2hu6rmt4A/37pZVV/b8Il5uw7QOGgX8NKwhbUJib+FR9NazhqzbZJkFeQiQFBISAAAAAAAAAAAA1ElVsgkdI6yhddT84bfP2cKWCOQhAAAAAAAAAAAAAGyEJAmJovLSeUl6TDFbINBHt3b+QsZrjFiP7w+x87M8VA8EAAAAAAAAAAAAsBEiJyRaTab49FR06bRATPDc2JDfMVjUT0StJtPRi+c25f4NPUUAAAAAAAAAAAAAbI2YCQlDc9MM/UqseTeL1lGzeMYzK0LClF2dqbq+dk/BkYTMdNqBAAAAAAAAAAAAAAAdoiUkDhefCUlYJdar2QKtoyY5KiZ0UoBapaIdi1T4LRGv705DmgoAAAAAAAAAAADAxomQkGg1maI2rkP/auFsIRXRYjRuzM1Ct2oAAAAAAAAAAAAA4FmbkECZJrPYQioC1ZkAAAAAAAAAAAAAoDOrEhJF5aUBq5eJFYqy2UgqYu0HadgrAwAAAAAAAAAAAACdWZ6QWJeVgVXwQiAVAQAAAAAAAAAAAABgSUKi1WR6+s1X80uKRY9GefThUatCI5CKAAAAAAAAAAAAAAAbZ3ZCAk0jBArzD9oaE29vZ0c7EKnwbauxSwYAAAAAAAAAAAAAhDAvIVFWXTnzjbi6xgaJolEGb3ePnDVvuzm70A5EKq0mU/apgvj0VLwTAAAAAAAAAAAAAEAgMxISaGHdK75dxDz/qbQDkVB1fe3cda9hiwwAAAAAAAAAAAAAmEVoQgItrHul+BpNrSZTfHpqal4O7UAAAAAAAAAAAAAAQH4EJSSQjeiZ1lFz8M0NXm7DaAciIVTrAgAAAAAAAAAAAABr9JKQaDWZnn7z1fySYjLRyFFM8NzkqBi1SkU7EKm0mkxJ2XuQkQIAAAAAAAAAAAAAa/SSkMg+VYBsRHdsYWOEoblpwYYEvAcAAAAAAAAAAAAAwEpmNLWG9hTfMYJDmSYAAAAAAAAAAAAAEA8SEpbIiFs7z38q7SikVVReGrB6Ge0oAAAAAAAAAAAAAEAhkJAwj9ZRcyJ5q5uzC+1ApIU25gAAAAAAAAAAAAAgrj60A5CTQB9d+baPkI0AAAAAAAAAAAAAADAXdkgIFRM8NzkqRq1S0Q5EWshGAAAAAAAAAAAAAIAUkJAQRB8etSYsknYUkkM2AgAAAAAAAAAAAAAkgoRE7wre3eLr6U07CskVlZciGwEAAAAAAAAAAAAAEkEPiV7YSDbC0NwUsHoZ7SgAAAAAAAAAAAAAQLGQkOiJjWQjWk2mGfqVtKMAAAAAAAAAAAAAACVDQqJbNpKN4Dhux+Hc0qoK2lEAAAAAAAAAAAAAgJIhIdG1XH2SjWQjWozGFWmbaEcBAAAAAAAAAAAAAAqHhEQX9OFR03UTaEdByMbcLNohAAAAAAAAAAAAAIDyISHRUaCPblVoBO0oCGk1mbYf2k87CgAAAAAAAAAAAABQPiQk/ofWUZP92ttqlYp2IIScvXqprrGBdhQAAAAAAAAAAAAAoHxISPyPg29usLezox0FOQVfXaAdAgAAAAAAAAAAAADYBCQk/ismeK6X2zDaURB18usS2iEAAAAAAAAAAAAAgE1AQuJnWkdN4nPRtKMgLb+kmHYIAAAAAAAAAAAAAGATkJD4WXJUjE0Va+I4rrq+lnYIAAAAAAAAAAAAAGArkJDgOI7TOmpCJwXQjoK0mlsG2iEAAAAAAAAAAAAAgK1AQoLjOG7xjGfUKhXtKAAAAAAAAAAAAAAAFAsJCY7juJAJk2mHQIHP0EdohwAAAAAAAAAAAAAAtgIJCY7jOC+3YbRDoMDezs7b3YN2FAAAAAAAAAAAAABgE5CQ4GKC59IOgZrZvlNohwAAAAAAAAAAAAAANgEJCe6pMeNph0BNwGNjaIcAAAAAAAAAAAAAADYBCQlO9/BI2iFQ4+vpjapNAAAAAAAAAAAAAEAAEhKc0wAH2iHQlLPmbdohAAAAAAAAAAAAAIDy2XpCQuuooR0CZW7OLhujY2lHAQAAAAAAAAAAAAAKZ+sJid94P047BPqWzpoT5h9EOwoAAAAAAAAAAAAAUDJbT0gAL33FmkAfHe0oAAAAAAAAAAAAAECxkJAAjuM4tUr16RvrkZMAAAAAAAAAAAAAAIkgIQE/43MS+vAo2oEAAAAAAAAAAAAAgALZekLC0NxEOwSGqFWqNWGRGXFraQcCAAAAAAAAAAAAAEpj6wmJ/JJi2iEwZ57/1Ks7srzdPWgHAgAAAAAAAAAAAADKYesJCY7jWk0m2iEwx83Z5UzKjpjgubQDAQAAAAAAAAAAAACFQEKCO3v1Eu0QWKRWqVIWLcdWCQAAAAAAAAAAAAAQBRISXMFXF2iHwC43Z5fzm9Iz4tZqHTW0YwEAAAAAAAAAAAAAGUNCgttXdJx2CKyb5z+1cme2PjyKdiAAAAAAAAAAAAAAIFdISHClVRWG5ibaUbBOrVKtCYtsyDqCtAQAAAAAAAAAAAAAWAAJCY7juLTDubRDkAd7OzukJQAAAAAAAAAAAADAAkhIcBzHJWSmt5pMtKOQjba0xMboWPSWAAAAAAAAAAAAAAAhkJD4WfapAtohyIy9nd3SWXMqd2bn6pO83T1ohwMAAAAAAAAAAAAATENC4mfx6anYJGEBtUo1XTfh/Kb04s27YoLn0g4HAAAAAAAAAAAAABiFhMTP6hobdqCThBW83IalLFrekHUEGyYAAAAAAAAAAAAAoDMkJP5rRdomQ3MT7Sjkzd7Ojt8wcXVHlj48Ch0mAAAAAAAAAAAAAICHhMT/iPvrZtohKISbs8uasMjv3t+HUk4AAAAAAAAAAAAAwCEh0UFW4bG9hUdpR6EofCmnlk/yC97dgswEAAAAAAAAAAAAgM36xX/+858e/npv4dHIDYnEomHE1R1Zbs4utKNQplaT6ezVSwVfXdh+aH9dYwPtcMBW/Cv3OO0QAAAAAAAAAAAAbB0SEl3QOmoqd2arVSragShcWXXlibKLOz/LK62qoB0LKBwSEgAAAAAAAAAAANShZFMX6hobnn7zVdpRKJ+X27Cls+ac35R+Y3durj4pzD+IdkQAAAAAAAAAAAAAIBUkJLqWX1K8LiuDdhS2wmmAw3TdhN1x+pZP8os379KHR3m7e9AOCgAAAAAAAAAAAADEpKYdALsSMtM5jlsTFkk7EBuiVqm83IZ5uQ1bExbZYjSWXPsm+2RB9qkCdJsAAAAAAAAAAAAAkDskJHqSkJmu6X//0llzaAdii+zt7Hw9vX09vVMWLTc0NxV/e/mzC2eRnAAAgM5+ajXdrq39952W+itXuzxggOuD97toH3Bx6aNGgygAAAAAAKCs6UYNx3E3ir/s8m/tHnhg0CMe9hrNff36ko0LgAQ0te7dxuhY5CTYYWhu+ub76wVfXdhXdBzdsEEgNLUGUJ6mGzU3ir/85tjxqqIv7hhuCnxWf6dBo4Kne0zxe/Axr34DNZJGyA7+bqcDpGcAbBOfwe38OL4TAAAAJPXj3XvVX5z7/quy784WXzt5WvgTh/pNHPFUoNsT4wY97IGTNSgDEhKC6MOjULuJQa0m0+Ub1SfKLp6+XJZVeIx2OMAuJCQAFKP+8tXLR44Wf7BXeBKiO/2dBumemzdy2lTnkcNFiY1NTTdq/jzpyc6PT1m5zD/2JfLxkHf3VsP3X5UZb9++19R8o/hi2+ODdaP7OgwYrHscS88sc/38hfO7P7Lsuf00jtP++Jq48VigcNNfblVcs+y5szevFzcYYspy8/Ytj+/8+PM5e4aMHUM+HlEc+ePbdxsaJR3ikaAp/B/4LXe/7NtXRlltaz6tXRqsGz0u8lmSQ49dMJ/Y+/NcxoftTxbsIPm1w+wPoUvT3vgDgc8j/2Ye6DFU4/6Q9a/GX5MI/yhZiV/HY8ETh/n5SvezLcvNs+bpI54Mksv1291bDV/nHfryo+y68ivWv9r4hRGjgqe7jvZBZgJkDSWbBEE/CTa19ZxYOmvO7jg9X9mp+NsrJ78uyS8pph0dAACIhr+OP/nn7dbnIdrcMdw8nrLleMoWftuEUq/sv/hrRpePF3+w1y9msfL+vW2un79w7dSZy4c+7+7eryz3QNuftZ4jHp8f+mjwDBnNMFLXXPN9+5+huUYFT6c7/d10o+Z4yhaLny7fhETR1p1dPn78vdSID7v+K/ad3bVH6iG6fLd7hcx6JGjKYN3jjO8vsfLT2iWBs6hiDf1I0BRi3xg3ii+K/uMSBcmvHWZ/CF0KiF/RT/pRpPgccYI/Sla6Ufxll6noXvV3GhR7pkCi7zfLQmrz8qnPHfq5ihWMFH5qNVWeOHXmrxlmbYbo1dlde/iz3lC/iRNejHR7YpxcEjMA7SEhIVRCZvqVG9XpK9aoVexea9o4pwEO03UTpusm8P9ZXV9bfr2q+NsrV25UY/8EAIBMNd2oKUjeKOld8R3DzbYre6+QWd6zn1bMlf2Pd+91N093x3Cz8sSphwP8CYdEwPXzFw69nmjWGrS68iuH9esO69d5hczqbh3utwWFpfs+JbMM0xYcej0x+vA+igEUJG+kODot9ZevdvfRuHbydNONGofBTE/udGf25uSaL78ikJbooCz3AH966u806Mm18aNmzWA5LQECjV0w/5GgKfeamsVdBgFyN1j3+OzNyd8cO264+q2V69y1niOchj88WDfa5VFPscLr2QDXBy174h3DzczI6PCMNCm+3LSeIyz+SWo9R/yyL7vX6j+1mi4dOPR5YrKk3yHXTp7mUx18QadHgn4j05M42CYkJMyQVXjM0Ny0O07vNMCBdizQOzdnFzdnFz4/we+fqGtqKKuuPHvlUmHZRfSfAABg3I937+Wt1hNeoNc2u6T1HDFyxpNyL2p05fOe8vH5Se8pLCFx91bDJ8tftWYZGv8G0HqO8F36Qoe5xU9fff2O4eaklxYhISGKuvIr189foLVJoulGjYwW/4roy73ZPfztVzm5Mv3S8woJ9goJfvL1VTUXS7KXrCA/iXzHcHPf8vjPE5ORllCAIWPH8F9NumfnEZhSBLlwGOzqMNjVKySY47gf797LWrTMguuN/k6DFu7/iPys8ZCxY1aXF1d/cY6/mDHruddOnj6Zul2Ks0P04X0/3r135fNjwj9lU1Yue2xuCMub0vhUhJWbP8zFZyYO69fx274D419RxsoqUDb0kDCb1lFz8M0NXm7DaAcCVmk1mWpuGbCFwkaghwSAvFh8Ke8VMqv9f1o/4bi2+pKVr0BX2vTZPS89e/nU54pZS1V/+eqHES92d0PLrx3jm0b8sm/ff9+7d/ObCuPt2z33RZ+yctmDj3kZb98u2rqT/0nO3pzMT0ZAm59aTTe/raj+4py5q4m1niNobZLYt/xVy74fhvpNfGi8TuP+kBzfBj/evfeup67nY9ZUlDI7yyPQ3VsN26c9Q3EGub/ToNBtG5lqyNFd4xCLeYXMElg+qG1oi5dC93caZD9ooO/SF2h96K6fv/D+3Ajhx89/f9ugRzza/vPfd1rqr1ztfFjNl1/dbWg0a609yWsSi78kqaByMfNTq+mvwaHmvqupX3f91Gr6/K0kC/aTzX9/m3SrWH68ey8/+b2eoxrqNzFsxxbG59nrL1/dv2KVue+KoX4T7dstdunhulQg6m8zACGQkLBQRtzaef5TaUcBYjI0N33z/fWSym9OXy77R+mXdY0NtCMC0SAhASAjTTdq/vbiMoGX8nzp/0eCftPdUqkf795r/O765SNHLSsWL+uERP3lq9unPdPzMYppbX0u48PD+nVd/tVQv4nBSQk93Jj91GqquVgisMoTEhI9sGAiePGR/eS7ynfX6b0HDE4xW6CHj0kbSaeciLFgonCo38TRv5vdoct9/eWrl48c7aEVTQ+YbRJ+91bD91+V5Se9Z+4/it865jxi+C/721s21XX3VsORN98ROM3doTb6j3fvUZyINGt23oKpwKYbNV/l5PZ6oYKERHdoTb+aezaZnrCGTMeIXvW8hqM7Up+ye0j+DfWbKFHZKLH81GrKXbla+KdmysplI6dN7e80qMt9tz+1mm7X1n5z7B+W9cFGQgJkAQkJy4X5B6GlhIK1GI3X6r4vq648eK4I+Qm5Q0ICQC6Er+IcvzDi8Xmhwu+L+LZy5k7ByDohIXA2QQELort72/R3GvT0+reEz64KSYYhIdGzH+/e2zUnXPinTPg6axH1unOog+kJa3TPzpP7x4TjuPfG+vc6/URx24q4zFrVPtRvYs8Nve/eajjx523mrilmNifBmblK2itkVkD8ClGmt4Rs0+nvNOjZPX8ln6rsgZAEfxuLpwJ/vHvv9I5dPaQl2ExItCXzOv+VkMl6vgJPhwf/fafl8pGjxR/sFThjTnH6NdFtlPCDmZomNnfrD8dx/Z0GLT6yX9LCld1FRWX5gnBNN2p2PTNfyNt1qN/EKa/EuI72EX5RUX/56qm/7DArQcjUOw2gO+ghYbmswmP/KP0S5ZuUyt7OzsttmJfbMH4rDPITAACSEr6wyCtkVkjKu+ZODvZRqx4O8B82eVLxh3t7XSOsAD/evSfw1kXura2vn7/QXRLL3CXtDoNdX8zLtpF3iETu69d35IwnhU/3l+UeEGuiU6Dr5y+YlY3Qeo5gZEGrla6fvyBkrqSu/Ip8W1u3Z9Znf8RTgT0f0G+gZtofX3t8XqhZa4rfnxvBbE6ij1o17Y+vXco73Os/Z6jfRAvOud358d69ng8YvzDiyddXsZb/IzMTel+/vv6xL42cNtWCpevkjV8YMSp4ussoTyt3rmjcH+ryC8d55HD/2Jfu3mr4Ou/Qlc/yrekOJamhfhOFx/aAi4ukwZhlyNgx0xPWmHXBc8dwc8+zL7yYly3dJ3TI2DFeIbM6XMFqPUewnI0QsvuQs6J9iPPI4bM3r5/00iJZfDMACNeHdgDyVtfYoFu+cOuBT1pNJtqxgLT4/MQ8/6m74/Tfvb+vIetI8eZdGXFrw/yDtI5obgkAYJW7txoyI6OFTKDP3pw8e/N6i2+E+qhV4yKfjbtwcqjfRMteQS4ufvyJwCPzk96TNBJJNd2o6W5935SVyyyYB+TfIc/nmF1bGdo8+JiXWccXJG+UJpCuHXrdvM3fv/urJdXeGHR+90cCj/wqJ1fSSIjReo4QeOQjQb8RcpjzyOGxZwrMOn28Pzei6UaN8OMJ0z03r+cD+jsNmmPFObezI2++08NYi4/sn/bH11jLRhDGv8069MRiSn+nQS+f+nzaH18bMnaM1HW0+g3UjIt8NuLDnbM3J0s6kMXszdkuwNp7e1zks+ZeD9eVX8mMjJYoHl5A/IoOj4ycYV6VRZIKN/1FSDbCK2RW7JkCa5L9/DfD+IXmbWoBYBkSEiJYkbZpwspF1fW1tAMBcjrnJwre3bIxOjbMP4h2aAAAMsPXne91fRl/AyxKtZx+AzXhGWnKzkmc/PN2gUfWlV+pv9xFv01Z+NuLy7p8XOs5wi9mscUvO2TsmMVH9lv8dBvXvpurEGW5B4jN2Jq7PWKo30QF7BXgzNkyxXHc8ZQtP97tZRm7LDgNf1j01+yjVoVnpPV3GiT8KV/8NUP0MMQyclovPRHdfZ8QsTzL9fMXunsf8tkIlhdBC2f9Kvg+atXszeunrOz6BEfX+IURVs6rWsYrJPjlU5+b9dEDIZ5au8rcp1w7ebpw01+kCKY7GveHSA4n0E+tpj3PviCkR52Vq6na8Dvb2PxmALAAEhLiKK2qGL4oDFslbJa9nZ2vp/fSWXN2x+n/lXv8xu7cXH2SPjwq0KeXGqkAADZOYBdcfqpCxBtgflJJqTkJgbVZ2ny5N1u6YKTTw+TyyBlPWnnj5zxyOLPrMZWH2CYJc7dHmLX0lWWnd+wy6/grnx+TKBIF6KNWPbvnr8KPP7trD7MJnl/2t+/5gEeCpog1Vs/taqUuTE+SWKvg/WIWs3aVMv/9bRS3sDgMdmV874gc9fol0KXjKVuun78gejDdcR7BYqoyd+VqIdW6ns/ZI27vMf/Yl5CTAGVAQkJM2CoBPKcBDtN1E9aERR5KSGn5JP/qjiwUdwIA6OynVtMny18VMnW+cP9Hok9VKDgnIbw2C4/l+bLu/NRqyl6yoru/HTppgvVDeIUEY2s8GWQ2SZi7PUIxfmo1FX+w16ynFG3tqcMzOI8cbtasqLkJIeXpORsRnpGmmGyEiFi7Snk+Zw/1jlP83hHkJFggXT26zsuPLMuaSKpw01+E7Du0rHxor5CTAGVAQkJk/FaJlTs2txiNtGMBJqhVKjdnl7biTjd25xa8u0UfHuXtbl5BAwAA5cmMjBaytmj++9skKg7QR62as3m9FK9M0d1bDcJrs7SR3YLomoslPaSy7nfRijJKYPwrKBBBhtR9C35qNZm7PUIxev6wdEnWldzICH43QfjBxR/s/anVdrfR95qNYK2wPjv6qFUWVNSRwvSENey0Zw9JeZedPI0t2/XM/Lu3GmhHQUFZbp6QSk1eIbP8Y1+SKAb/2JfwKQC5Q0JCEql5OZ5L5u8tPEo7EGCO0wAHX0/vNWGR5zelt3ySX7x5FzpPAIBtOpfxoZBsxJSVyyRdkddvoEZhlXnO7zFvKTRPdguiL+Ud7uFvf6itE2WU+/r1fXJtvCgvBT2Tum/BpQOHLNgeIWK9GoqOv5dqwbNkWsmNmPv69RU+GXTHcLPmYomk8TAL2QgrOY8cLrw9u0TGL4wYF/ks3Rja41eTYLkAdXcMNz9Z/qqtZVubbtTsW977laHWc0RIyruSRqK8NVVga5CQkEpdY0PkhsSxsVGo4ATdUatUXm7D2jpP8JWdYoLnorITAChe/eWrh/Xrej2sv9Mga1oTC+QVEqyYRUYW1GbhyW5BdM8JiWunzog10KhZMzDrQYZ0ZW1+ajV9nqiovKNwd281CEn9dibHSm6Ejf7dbOEHN9d8L10kzEI2QhQz3lpLN4AnX2dil0Z7/QZqQrdtpB0FcNdOns5duZp2FOT81Gra9cx8IUc+szFJ6u835a2pAluDhIS0+ApOCzYkGJqbaMcCrOMrO6UsWs5XdsrVJyE5AQCK9FOr6cOIF4Uc+fT6t8jMVoTt6H3ntSxYUJulzam/7BA3GOk03ajp+Z95+dDnYo3VR63SPTdPrFeDHki3SeLSgUMWfy7kzrItUzzZVXIjzKw+qw1V30kXCZvu3mrortMPshFmcR3t8/Kpz18+Jdp5zSzjF0aw+ZtyHe1DOwTgOI4ryz1QuOkvkg4hUeFWC3z+VpKQywmvkFnOI0k04vYKCaa+gwrAYkhIkJBVeGzwgpB1WRloLAEC8W2x25ITaDsBAEpyMnW7kKt5recIYu0T7+vXVxnd4bqrzSLkdqUs94BcFkT/+05LzwfUlV8Rsa7xyGlTxXopWzP//W38PJrAHqRSbJKwZnvEvaZmcYMhrLstUwI3/ciukhthZvVZvVVxTbpILHaj+EuJXvnurYbt057p8lyPbIS5+qhVDoNdac3JjgqeTmXcXvVRq8Yv7Hr/DRB2PGXL9fMXxHo1Zq9Fm27UnN21R8iR0974g9TBtKG+gwrAYkhIkJOQma4Jm7a38Girybaq7IGV2redaMg6wicnsHMCAGTq7q0GIY3gOI57ZmOS1MG0NzZC9qvgm27UdFmbZfzCCIG3Kxc//kTsoCRRf6X36lLbpz0jVl1jMsvcFGnQIx78PFpA/Aohx0uxScKa7RE3ii+KGgtplSdOdflvf3JtvJAUbF35FRHnmJTHrAniFltq/WrL2QiByVcZYXkjArPJElmzrEzl+3MjxFoI0tLA6Lfl314UtHRpyspl/QaSm6sZMnYMNkmATCEhQVrkhsRhL4QiLQGWsbez45MTKOsEADJ15M13hBym9RxBeBa430CN3OcRvsrJ7fLxJ16MdB3tI+QO8+Sft4sdFDV3DDczI6PFyknEXTjJr/Qf8WSQKC9oaxwGuwq8YRZ3k0SX2yNePvX52upL/P9EHItB+UnvdX6wv9OgUbNmCEzBnt/9kdhB2SjDlW9oh2AJs8pS8XrIRoxfGKHsbITyMFuvicdyskS+BGasO9s+7RlmNzdY7/r5C3XlV4QcOXHRQqmD6QCbJECmkJCggO93jbQEWKlDWadcfVKYP2ZJAIBpTTdqynIPCDkycNUrUgfTWUD8Cq+QWe3/Rz4Gi/3Uaupy64lXyCyHwa591Cq/l3tvD37HcFNJC6KvnTydGRktyu1xv4EafqX/ff36Wv9qtkngDXPxB3vFSiNxXW2P0HqOYKcataSabtR0OXuie25eH7VKYApWRpXcyDPrjeru+4R0kUjHrLJUXI/ZiCkrl03742ssz25De7M3r19bfWnaH1+jHUhP+qhVfGrZRr7VifGPfWmo30Rzn3XHcDNr0TIRz+BMOfR6opDDvEJmkb9QHDJ2zPiFEe3vX37ZFxerIANq2gHYLj4tEZ+emhwVEzopQK3CxRlYjk9OTNdN2B2nr66vPXiu6O9fnMwvKaYdFwDA/+huCX9nwyZPkjSSLjkMdp29eT35cUVReeJUl4+PXTCf/8Po3845rF/X6+uc3/3RkLFjxIyMqmsnT2/xn7b4yH6S2+ehS/w2nV6rJ90x3Lx04JBXSLD1I/54996+5fEdHqSS7KSiu+/btsWbYxfMF5IhvvjxJ+MinxUzMqW4XVtLOwS2/NRq+mT5q11+xqcnrMG7CEBGwjPS/hocKnBPQJtrJ09//lYS43ksCwjfHtF21U2Y8n7mYAuwQ4Ky9rsl0PIaROHm7LJ01pxDCSktn+SjGzYAsOPHu/cEdo+YsnIZFlGaq8vaLFrPEW3Zhfv69RXS/lEWC6IH6x4XfvAdw80NY/zKcvOUumpPLgRu0+E47vPEZFF+WV1Wf6KS7CSvuy1TU1Yua1u8KbDwtJIquYnrh9o64Qc/EjRFukhY8FOrKTMyuss+Rs/n7EE2AkBe+qhVER/utKCfxNlde85lfChFSBQdfy9VyGH9nQYpaU0PgNSQkGACn5bQhE1bl5WBtASIRa1StXXDRk0nAKCu+otzAo8cOW2qpJEoT3e1WToUyXnixUghryZuEX9G7Fse/9fgUCUVpJKj0b+dI+QwfpOElWN1mQG1nWRndz/Ax+aGtP9PIftFFFbJTUTXTp2hHQIres5GYIYOQI76DdQs3G9JG6HD+nVKOmv8ePdel19unemeE9ScCQB4SEiwJSEzXRM2beWOzdX12AIMYuJrOu2O07d8kl+8eRe2TQAAeWf+miHksP5Ogwi3s1aAL7r62fZ3GtSh36PDYFchRYHFLeIvhQdcXCx4Vl35lffnRux59oWmGzWihwRCCNymw4mxSaLLvFqH6XgFK9q6s/ODfEeZ9o8MmzxJyAJYtLbu0uVDnws/eJifr3SRWOxeU7P1L9JdNqK/06CXT32ObASAfDkMdn0+Z48FT8xesuLurQbR46ECC6oAJIKEBItS83KGLwqboV9ZVF5KOxZQGrVK5eU2DNsmAIAw4cuLRgVPlzoYhfnx7r2zu7q4XXxybXznxeBTXonp9QXvGG7WXCwRJzhp9FGrLG45fu3k6T9PenLf8leRlqBC4DYdKzdJdLk9wnbaWddfvtrllqlJLy3q8EgftUrIis6y3AOKmVoSy91bDcKrqw/1m8hmG5sbxRetf5GTqds7n9+1niMWH9lvI584Ufx4916i2yj+f/uWv0o7HICfDRk7ZsrKZeY+647h5vZpzzC+ukUggQuqOI7DgioAsyAhwa78kuKA1cvGxkbtLTzaalLCVzmwpv22Cb7bhNaRxZslAFCA2kvlAo9EQsJcVz4/1uXjo2bN6PzgkLFjhCyIFlgtlyIr2waW5R5AWoIKh8GuQvoWcBz3eWKyxaN0uT3CdtpZf7k3u/ODWs8RXc6VtPW47tn5PXutDUtZvs4zI2EmJBMsU4Wb/tI5+TfUb+KLedls5mCkZnGzkMbvrosbCYBY/GNfEri7sb07hpuZkdFyz0n81GoSuKDKgh8RgI1DQoJ1pVUVkRsS7ecErsvKQB0nkEhbt4nv3t9XvHnXxuhYFHQCAHEJr7U9cKi7lIEoUJe1WXqolf/k2vheX/PaydOML4gWmFnpGdISVAhMDFjcuqDL7RGcRe2sy3IPWBAAXd1tmerQUabNff36CtlvxH4lN5J+ajUJ7/Wt4Dan3WUjwjPSbKRZi4hO/WUH7RDAFv37TouQw558fZWQmp8dXDt5+mSq0K9KNt38tkLgka6PPyZpJADKg4SEbCRkprfVccKGCZCOl9uwpbPm8AWdMuLWBvroaEcEAErw3dliIYf1dxpkm2sqLdZdbZYeVj2PeFJQpT72F0QLyawIgbQEYQL7FnAcd+j1RAteP2+1vvODttPOusstUz3PiXcu5dQZ+5XcSLp04NAdw02BB4v1TcWacxkfdpn5C05KsJHPmoiabtTIMf0JClB/5aqQw/qoVeEZaRYsBDmesqUsN8/8uFgh8OfDcdxg3eOSRgKgPEhIyAxfx2nYC6HrsjIMzU20wwElcxrgMM9/6qGElIasI7n6JGQmAMAaAvc7u/s+IXUkCtPlmkqvkFn39evb3VPu69dXSDlg9hdEj5o1Q2DxHyH4tMSRP7794917Yr0mdElg3wKO4+rKr5i7SaK7eT3baWfdZamrnufEnUcOF/JRYr+SGxl3bzUIryc2ZeUyr5BgSeOh4vr5C4f162hHoRx5q7pIowIwpY9atfjIfgtyEvuWx1u235EF3xw7LvBItMwBMBcSErJU19iQkJk+eEHIDP3Kw8VnsGECJGVvZzddN6EtMxETPJd2RAAgM8IXnltcfNk2/Xj3XpdzrwHxK3p+opDJ2TuGm5UnTlkWGBl91Krf/bWL9bnWOLtrz7ueunMZHzKejJE7gX0LOPM3SRQkb+z8oO20s75+/kKXK/e77CjTnpA6WuxXciPgx7v3tk97RuD2iCkrl/nHviR1SORdP3/h/bmoli6awk1/EbhoA4CufgM1ods2WvDE7CUrZHr6MFz9VshhFtSzAgAkJOQtv6Q4JGEV32GirLqSdjigcHxmImXRcr4JNjITACC6Aa4P0g5BTi5+/EnnB4f6Tex17tVhsKuQe6f8pPcsjIwUh8Guszdb3vq4O4f16zZNCKi/LHSfPphLYN8CzsxNEt1tj7Cddtbnd3/U+UEh5aoE1tFiv5KbZew1gkoFfltQuMV/msBsxFC/ibaZjfjirxnEgpG7phs1adNnd1n5CoBNQ8aOmf/+NnOfdcdwc8+zL8hxqUeXZVE7s0e9WQDzqWkHAOJIyExPyEz3dvd44angBUEz7e3saEcESsY3wfb19E6Oijl79VLiR7vySwRVhwd2JLqNoh0CSGJt9SXaIXThRvGXAo+830UraSQK02VX1SmvxAh57pRXYnpdkllXfqXpRg3jS8u9QoK/OXZc9Orbdww3t097xitkVvC7CT3UvwKLTXppkcDf2qHXE6MP7xNyZJfbIziL2lnLUXdbpoTsR+HraPU6MXo8ZYtfzGLldQi48vmxUbNmdPfv+vHuveovzuUnvSdwZqq/06An18b3uitFFjo0vBWyN+Lsrj2jgqcrtY+3NX68e6+loeGH2rrmmu8bqr777mwxNkaAHD0c4D9l5TJzE2l15VcyI6MjPtwpUVRSEF7Dc7ButJSBKB+mJhSp13kJJCQUpbSqYkXaphVpmwJ9dLEhv5v86GhkJkBSfGbiUEJKi9F44uuLm3L/hswEAAAZXdZm0XqOEDgNNGTsmP5Og3pd6vtVTi77i3xnb14/0GOoFItMy3IPVBV9EbptIybXRMf3LRAyw1tXfqX+8lXnkcN7Pqy77RG20866yy1T4xdGCMyoTVy0UMiHqPLEqYcD/M0Ojm37lsfvWx7PcZxXyKyBHkM17g9xHHevqflG8UXD1W8F5iF4U1Yum7hooWKymPVX/vvRu3urIXvJCiHPyl6yYvGR/f1scslw23sJQMH8Y1+yIKN27eTpwk1/Yf+qsk1Lg9AyU30dBkgaCYAioWSTMvGlnDRh0xZsSCgqL0WTCZBa+z4TGXFrvd09aEcEALIksG4GcN3UZjGrNE3PfW55x1O2yGKLvX/sS0I6dVvgjuHm+3MjzmV8KMWL2zjhb9cum7d30N32CBtpZ/1Tq6nLLVNPvBgp8BXu69d3/MLeGwOwX8nNGmW5B46nbOHnlA/r15XlHhCejfAKmfXyqc/9Y19STDaivbu3GoQ3z7hjuPnJ8ldlce4AAMuEZ6RZ0DjheMqWbwsKpYgHAGQHOyQULqvwWFbhMY7jYoLnhvoFjB8+Sq2yiTViQIu9nd08/6nz/KcampuyT+a/8/EHdY2ybGAFAFRIMY9zLuPDG8UXRXmpsQvmM7JSvsvaLP2dBplVmmbUrBlCFnLKZUG0f+xLfR0GHNavk+LFD+vXXfksPzwjzUbW2pPB9y0QMsVZlnsgIH5FD9XD6i9f7XJ7hO20s665WNL5J+kVMsusf/4TL0ae3bWn52NkUcmNpPELI0YFT3cd7aPIL4eLf9v3zbHjHMdVFX0hMBvBu3bydPGHe8dFPitZaABAUx+1as7m9cLzlG0+en7J4iP7e9j1ePObCqujI815RC+bOC1w/fyFLtceWWCwbjS+jYFBSEjYitS8nNS8HA6ZCSDFaYDD0llzls6aU1Zdmf5ZHv/2AwAg70bxRbEaDDwSNIWRhMTpHbs6P/jk2nizZsT6qFVCqgDnJ70ni4QEx3HjIp91e2LchxEvmnt7LMS1k6czI6ORkxCRwL4FvILkjbM3r+/ub/evWNXl47bTzvr4e6mdHxy7YL5ZL+Iw2FVIHS1ZVHKT2viFEY/PCx30sIeyvxC6K8ky1G9ir9VaDuvXuT0xrtdia0rV32mQu+8T7R9pudWAphGgJP0Gahbu/+jPk54094kfRrzYQ1U34+3bVodG2i/724v+ms0134vYIA0JCWAQSjbZnNS8nIDVy+znBK7csRnVnIAAL7dhKYuWt3ySn6tPCvTR0Q4HAED2fmo1FX+wt/PjI54MMvelxkbM6/UYfkG0ua9Mi/PI4bFnCoRUnrEAn5NAHRIRCem3zCvLPdDd+/D6+QvdzaHbSDvru11NdArvKNPejLfW9nqMXCq5SepS3uF/3bmj7GxEl/o7DVp8ZH94RprWc0SvB38Y8aLwrrDKMHtz8trqS2urL71yvnD25vXt/xfx4U7+r9ZUlC4+sn96whoLKt4AMMVhsOvzOb3sq+tMeVXd5LirA4A6JCRsV4fMRIvRSDsiUDK1SsU3mbixO1cfHqV1RJl4AOiCkm5OpNNlbZYpK5dZUPCq30CNV8isXg/74q8Z5r4yRX3Uqml/fO3lU58LmS8z17WTp3NXrhb9ZW3Wff36CnkH8rrrEnHo9cQuH7eddtbn93SRoRSSWujMdbRPf6dBvR526cAhC16cWW2TyG3/6/Xbwza7ywz1m8jXWumjVv3ur73vbbpjuJm1SJLuPrLWR61yHjl8XOSzER/ulOhUBUDMkLFjpiesMfdZCruakuOuDgDqkJCAnzMTmrBpM/QrDxefMTQ30Y4IlMxpgMOasMjKndkF727BhgkA6OB2bS3tEGSgy9oswleadyCkqMvZXXtkt8rVYbDri3nZszcnC5ldNUtZ7gG0ZBSR8LJCXW6S6GF7hO20s+68Zaq/0yDX0T4WvFoftcrv5cW9Hla0dacFLy4jQmbbOY47rF+3T1nrfHswZeWy8Iy0tiorDoNdhcxCXjt5uiw3T+LQZMxhsGv04X3ISYCsjYt81oKdqWW5Bwo3/UWKeABAFtBDAv4rv6Q4v6SY4zhvd48Xngqe7DXay20Y7aBAmdQqla+n96GEFENzU9rh3O2H9qP3NWFrqy/RDgEALNFlbRavkFkW9wMfMnaMkKrxVz4/5hUSbNkQtPRRq7xCgkfNmnEydbvARgUCffrq68sKj0jRg90GCXwH8jp3L+hue4TttLOuPHGq85YpczvKtDf6t3N6bQ5fV36l/vJVBbcHcBjsOn5hRK8tvjmO42t899DgRAH6Ow0K3baxcwUw3bPzvvwou9cP777l8YN1j9vI59Eyz2xM2j7tGdpRAFjuyddXGb6pMLdLyvGULUMnTWCkPVtnD7i40A7BVmBqwjZhhwR0obSqYkXaJt3yhb8KmYKCTiCp9hsmvN09aIcDAJIYrHtc4JE/1NaJPvrszevbqnA8n7PH4nWIWs8Ro2bNEDc2C3RZmyUgfoU1r+m79IVej5Hvgug+apV/7Eury4unrBStcsgdw80u+4qDZYS8A3nHU7a036zTw/YI22lnnZ/0XucHrfmyuq9fXyFrXb/cm23xELIQGP+KwP1Vyl7ny5dp6nLGUGDhJo7jdj0z30b2kVjGeeRw4cXrABjUR60Kz0izYE/q+3MjmG1UJjyvX/PlV6KP7hUS3Hb/8vKpz635ihC+FRWAJOyQgF6k5uWk5uVwHBfoo/t/T/hh2wRIgd8wcX5TenV97Z6CIwmZ6bQjAgAx2WuEto1prvle0nVSQ8aOeTEv+/O3koSse22vv9OgiA93Uq9H/1OrqcuV/n+e9KTUQ8t9QfR9/fr6x740cdHC0zt2ibJb4njKFr+YxdTfEsowataMfcvjBR58escufpPET62m7rZHcDbTzrrpRk2XKZl1Ht5SD312157A+FcUvE/ovn59n17/1kfPLxFyMOPrfC021G9ieEZaD190fOGmXrfU3DHczF25Wtn7SKw06aVF/G4bAJnqo1YtPrJ/+7RnOm/a69muZ+YvPrK/rRwcU4b6TRSy7cMgcVNrh8GuszevH7tg/vtzzS6N9XzOHuWdm0AZsEMChMovKW7bNrFgQwK6TYAU3Jxd1oRFNmQdyYhbi8bXAIohfMbqm2PHJY2E47g+apXr44+Z9ZT+ToMYuVOqPHGK4ugKWBDNpyXE2i1B99ehJH3UKuG/kbZNEpcOHOpue4TttLP+KieX4uhXPj9GcXQCHg7wH+o3UeDB78+NqL98VdJ4yLMfqOn1o6R7dp6QrYdluQfQTKIHjg8NoR0CgLX6DdSEbtto7rPuGG5+wmozHqdHBJVwMLdWlWXud9Ga+5T5729DNgKYhYQEWCKr8FhIwqrBC0Ieen72yh2bkZwAcdnb2c3zn4o6TgBKInBOx3D1W6kjsYC77xMsZCO4bmqzECPH1tZdaktLWFkio3Tfp2KFBGMj5gk/+PSOXT+1mj5PTO7uANtpZy1ucxRzybeSm3DBSQnCD/4w4kU259QkJbxw077l8czWZqFOwZuNwKYMGTtm/vvbzH3WtZOnP38rSYp4rCR8DROBLzcLelrYyG5RkCkkJMAqdY0NqXk5fHJibGzUuqyMovLSVpPNXYiDFNrqOBW8uyXMP4h2OABglRFPBQo5rK78CoHZHOcR5tUdGugxVKJIzNJdbRaSLn78Cd0ARHRfv76zN6+3piwvymuIqN9AjfBfxPGULSdTt3dXFMJ22llfOnCIbgB15Veun79ANwapOQx2Fb59547hJptzalLjCzcJOXLXM/OVkdiWwsunPuf/N+2NP9COBcByDwf4W7AP9eyuPecyPpQiHmvQbYPXgQVbP21ktyjIFBISIJrSqoqEzPSA1cvs5wTO0K/ceuCT6vpa2kGBEvh6eu+O01/dkYW0BIB8uT0xTuCRNRdLJI2E47j6K+ZV1dC4PyRRJGahW5uFd/LP22mHIDK+LO/zOXss6MTIEVkQZzvMarrYw84A22lnzcIGhfO7P6IdguT8YhYL/344u2uP8go3CTEu8lkhWyHvGG5mLRKhYp4iOQx25f/HyKZMAIv5x74kvN5dm8P6dawluR0Guwr8/r+Ud1jqYMxl5T5gAKmhqTVIIr+kmO85oXXUhE4KeGrM+MmPjra3s6MdF8iYm7PL7jj9hheXpx3ORddrANkR3gz5Ut5hVDvt7Me79zrPwD6fs8eCerJd+ubYP3rtSspx3B3DzevnLyjvFzRk7JjYMwUWdDv/950WiUKyQUPGjtF6jrB+G5CNFCiov3y1w8+qv9OghftFSw8UJG8UsgeoLPdA8LsJyq4200etEt7dmuO4DyNejD1TIIt1qYN1o0Xc6TVn8/oNY/x6PezaydPnMj4cF/msWOMCAIPCM9IyI6PNba7w/tyI8QvN7tssqVHB04VcHJ7dtWfaH1+TNBIsggGFwQ4JkFZbTSdN2LSxsVFbD3xSVl1JOyiQMacBDnzXa314FO1YAMA8Am8wGFxhxILOzWO9QmYNGTumbUGllf/TPTtP4BIwpS6I7qNWTfvja8/nmJeQMHe3DfTMd+kLVr6C7bSz7txk/sm18WJ9ITgMdg2IXyEwEiVVcuvOwwH+wpeayqhwU1+HASK+Wr+Bmtmbu+3s0h6D66ABQFx91Ko5m9dbsP3U3KUhUhsVPF3gkba5PQ7AYkhIADmlVRUr0jbpli/8VciUlTs2F5WXthiNtIMCWbK3s0NaAkB2Hp8XKuSwO4abuKDvrHNtFuHThUL0Uat0zwnqKlyWe0DBFcCHjB1jbk4CRDTiSWtrM9pIO+sf797rPGUzatYMEYdwGOyq9Rwh5EjlVXLrklll/c/u2mObS1m9QoIF1mnJXrLi7q0GqeMBAIr6DdSIuG+PFtfRPgLTKpePHJU6GAAlQUIC6EjNywlYvUwTNo3vNmFobqIdEcgPn5ZAbwkAuXAeOVzg9BYu6DvoXJtFira9I6dNFXgk9QXR5zI+3Lf81W8LCqV4ceQkKLqvX18LOmG2Ef1zIfAri7zOW6ak2BoisBsHX8lN3KEZ1G+gRmDfZt7fXrTRTglzNq8Xctgdw81Plr/6U6tJ6ngAgCKHwa5yv6YSvmSn+IO9+E4DEA4JCaCMbzUxeEHIQ8/PXpeVUVZd2WrClziYge8tgbQEgCwIrMdyPGWLgtfgW6BzbRbrK9t0JjxjdPLP2+necd0ovliWe+Cj55dI9D4ZMnaMwGnxAa4PShGALbNmi4Po7aydhj8s7guKpfOWKeEJReHcnhgn8EilVnLrQHhpO47j6sqvlOXmSRqP9ZxHCG3vJJzwwk3XTp7OXbla9ADowkkBoIMhY8eYlc1l0MRFC4Ucdsdws/LEKamDAVAMJCSAFXWNDQmZ6brlC+3nBK7csflw8RkUdALh+LREwbtbtI4a2rEAQLdGzZohcMpb0jX4NV9+Jd2Li67L2izD/HylGEtgnuOO4WbNxRIpAjDX6R27JHplv5jFQmYexWoqDm0cBrsKLPnSmY20s75+/kLndtbOI8WfWb6vX1+BvX+UXcmtTR+16tk9fxV+/OeJyYz/WH7Z316KlxVeuKks98C5jA+liIEWnBQAOhsX+SxrrarNcl+/vgJzKvlJ70kdDIBiICEBLGrrg42CTmAWX0/vyp3ZGXFraQcCAF3ro1Y9s1FQq09J1+DfbWiU6JWl0Dk3o/Uc0W+gJMlX4RX8j7+XKkUA5jqeskWiKuR91CrhbQxBXFNeibHkWTbTzrrzdgSB1SQsILD3DydldpApziOHm9XdOj9Z3pNT/TSOlj1RYOEmDg2uzdR0o6bpRg3ab4DsPPn6KotXG7BA4A65uvIr+EIDEAgJCWBaW0GnsbFRfEEn2hEB69Qq1Tz/qQ1ZR1DBCYBNziOHC1kkdcdw89KBQxLFMNBjqESvLIXODWMfny90itBcwhdEXzt5mpEJkU+WvyrRK/eakOjvNEj0Th7AcdyQsWOEF8ZpYzvtrMtyD3R4UIp6TTzhldxsp3Z28LsJwg+We3dr18cfs+yJwgs3cWhwLVjTjZo/T3ryz5OePPLmO7RjATBPH7UqPCPNgpM7I/qoVaHbNgo5kpElOwDsQ0IC5KG0qoIv6IRWEyCEvZ0dKjgBMCsw/hUhNyT7lsdLNEOhcX9IipeVwvXzF+4YbnZ48NHgGdKN+MSLkQKPPL9nr3RhCHft5GmJCrW7jPLs+QDplqXDk2vjzTpeijbvHJPJyy7L2UlRr6mN7Cq5Se2+fn2FT7VzNtzdWnjhpjuGm9unPcN4eSsWFCRv5P/wSNAUqoEAWKKPWrX4yH755iSGjB0jZIecdNelAAqDhATITIdWE0XlpchMQHf4Ck4xwXNpBwIA/+O+fn0FzjZiDWDn2izS1WviOQx2ld2C6H3L4+svXxX9Ze/r17fnH4V0y9JBePUwnujtrHkMJi87b5mSujD3qFlCM6C2syxUeD8kju3u1jeKv5T09YUXbrpjuJm1aBkj5xQ21V++2rY7Sopu5AAE9BuoEbjPgE3T3viDkMPYbyAEwAIkJEDGUvNyAlYvQ2YCeqBWqVIWLacdBQB05BUSLGQSrSz3wLcFhQTiYVOXtVmchj8s9bgz3hLUiYepBdEfRrwoxX6aDq2D25OojbAC/LJvX+tf5L5+faesNGNdue20s+68ZcriojoC9VGrBP4u2KnkJjXh/ZB4zE5O3WtqlvT1zSrcdO3k6dyVqyWNR75+ajV9GPFi23/Kd405sKah6jvCIw4ZO2b++9sIDyqWfgM1z+fs6fWwO4abeav1BOIBkDUkJEAJkJkAAJAdgd3tPnp+iY3McHXW5c0MgRoyrqN9BE52HHo9UepgBOIrfoj7Vun51WS9xE9SYu3gEd4TwuJ21mzOEXfnp1ZTl584uwcekHrosRFCq5MxUsmNAIH9kHj88n9J47HMjeKLPR/wzbHjVg4hvHATx3FluQcKN/3FyhEVKTMyun0yUtKNkoRJnRWDnt2quNbhkZovv5J60IcD/M1ac8CUIWPHTE9Y0+thZbkHzmV8SCAeAPlCQgIUBZkJAAC54LvbCZmnEH2iWRa+LSjsvD2CI1JDpo9a5e77hJAj68qvsNOvVfScxPdflXX3V0P9Jv7/9u48KKoz3eP4GSApt4kGQSzBGOKCOIxRQSMJFy4aB71yC41moiMjool4jYJxSs3NRMtBzYxOaYKjWUwmaq7r1BilghUZo5aMu0A0UrIoLgFSrA1msEmFxfvHSVGM0s17Tp+t4fv5C6HpfqS7T5/z/t73eQeGjdHqgdAu8e5hqntn3b142fkNDBiXEXd8/cZ2l+z4DB2s90P36OstGFJap5ObAQT3Q5JZs6t4VdFNAx5FvHGTJEmnt2xjFO8hWWnv3z5zvvWfIl3srUAw8e0wFYPB7LZaAx4lMmWxeFRpNWMT5ogEKsfWbCjJzjWgHsBNEUigcyKZAADrE8wk6quq98xZoNUgV0tT87kP/qrJXenHXmP7YsXb7f6ot/8AAwoQ3zDzm0PpulaiiLaZxLXDXzj6UezGVE0eolPScDxapHuYX3CQ6t5ZFz7Z7fwGxozLiCjJzr20s/02EZr0yOqQYEhZX1V9659n9S7GIsT3Q5IdTl5pqXC9JDvXSVc6DSlq3CQxitdGS1Nz5tp3Tm/Z1vabBiyU1ET59XyRm905d1HvSuBEu3NfjCE4LcmaIlMWi2QSu2bEa3g0s9QpN+A6Agl0cm2Tiby7t8wuBwDwbwQziYr8wrTx0ZpMxs/Ze8CYIRjV7DW2j2KmPdopXvbz/n4G1CAee5zess1SE6Lrq6o3j4lwfeuRy7v3OrpKn73rwz4B/i7efyf2fXl569f/Kq9w5a5Euoc9/z8L1N15o72h7aTjduWlH7XCCHJJdu6uGQ67AxnTvCUgdJTgLU9ufFfPQqwlJC5WfHdryWIL/kQ2Ib9z7qImnc0UNW6SJGnXjPiuvIOUrNHesC9h4aNJ5ICRIabUo9T1jGMiN6uvqiZ/Mku75293zl005rzOw8vzpa1/dt8NUcQzCU2OZpUFRQ9lk4C7I5BAV7E941BocuJT86ZvOLibZAIArEPOJDpsxl1fVf2XFya5eE5fV1p2bM0Gpb/VL8i4vYudpxGGuX32gviNj69XsLmrMfbPW5SV9r7qK+rKgiJHr5Oo5UuGREe6UFrn92P9/davRUY8nfDw8oxYmuT8NkGTJqq455amZsGe/pl/+KOK+9eQ8zTCGC1NzYX/OCl444r8QrfbBsCV3ly//kTBCJG2C/5cUVda1mEgJ2m6Nauixk2SJO2ftyhz7TtW+FtJkmRwc8KWpua89IxtkTHtPkcGdGlznb3G5mhR16OyP9uvazG6UtT3zGobF7W7oK2+qvr60S+NKaBHX+/EI2787EemLBbZ49r1o9lD29oLcpfVVOiyCCTQtVTU2lL3fRqanBiWMv+Do59X3aszuyIAgOTh5Rmz9i2Rlg775y06nLyisqBIxaPYa2w7p81W9CuBEeG/yz2juiGMUnWlZR2mES5OORdRkp2raBLWpZ17VD8p+jm9ZVva+GgV8y7rSsscXfVFLV8SmbLY5dI6M3uNre1f7/aZ8y4O5I16+SUnPx2XGP9YD8UNi+w1tn0JC0VGYyVJyks/mrn2HbNGkW6eyuowjTBgqDRn7wHBP5fs9JZtWWnvm7vBTKO9QbwVzKWde1QfwfoE+Ivvbi1JUkV+4b6EheaOS7Y0NWesEo0Z8tKP7pg83fU57EobN0mSdGnnnn0JC62wpqT6RrHgLXdOm315915Xnt+bp7LSxkcfTl7p6GTA+kv0WpqaP09eIX77vPSjFtxhRZCiRbe135boV4kKjha0HU5eadiylT4B/iJj+pY1MGzM0rPHO1zncWnnnk9iZ+alZ6iIJVqamh/a1r5DvXx95h3awykrLO5nDx48cPLjA1lfJWxeZ1g1gPEmPBuaEvfr//jFqJ7dupldCwB0dfYa2545C0Su7vyCg6asXy2+sXBJdu7fFy0TP5vv5eszafXKkLhYwdu77vLuvSKrN3r5+iQe2a/feITSP1RbvXx9nn7+uaETo3r7D9B1z+fDySvEux6HxE0NmztbpJ5Ge8PJTe86mtQ5OfX3YxPmKKiy66krLds5bfajL56o5UvCX0tUkRzIMte+4+hJSco8oigvtNfYvvsm74sVb6t4hbv4v1Cqpan5+PqNIlOMAyPCX9r6Z50aN7U0NV8/+uXhZAU7JbTlFxw0aPxY/9EjA0JHGzmKKg/fKApR/IKDXs34u4eXp7qHSxsfrXS0aOaH7+l6nHSkrrTsb68uUdG60C84aMKqNwY9N9aVd8GOydOVPnQvX5+IpUmDnhv75FMDDXsDttVob9gWGaP0oBEYET7+1YQBI0NE3pstTc1lV67ePnsh5/8OOH8gv+CghccOK6rEeFlp76voLfO73DPGNKDT0M1TWfvnLRK//bjE+Ji1b+lXjyLOnya9zzYf4vwcePXd68aUoZr4R7Z8fTFi6hTBj5u60rKMVWsUfZxFLV8S8XqSuo8zwEgEEsBPXo+dMTMi+vngX5pdCAB0afL41/F1m0Qu/v2Cg4ZPmRT4wnj/Uc86OvMuyc798u11ikZAJqf+PnTOLMNO5VVUOH3rJs3DksqCon+s26jomseRkLip0xV251BEUSAhk4e0hk78z3avru01tltnzjkade3l6zNnzyeGLZRxF432hvs2myRJ/yqvuH32QsGXx52/hkPipgaEjurep3dA6Ogn+vcXf3/VlZb95YVJj35ffGCurrTs1Kb37py7qEkztMCI8J59vXV9heelZwgeA2V6jG7LI6RKD02O6HHIcqSyoGhv/KsqnuvAiPDYjanqBuDUddYalxg/YeUbBgyy15WWleZ8LUnSlb8ddv0gHxgRHvXG6+peb5UFRR/FTFP90PKHvsHTfp1koiLkqF4++j30I9udb2uKb9+vsYk/KZYa0X5Uo70h4801qvdJNvJA4Tp7jW3zmAilvzV714emN36UPxNFniYjnxEnAYn1AwlZZUHRkWWrBD80xyXGD46KcBLxij9NrULipkavXGb9RVSAjEAC+Dd+T3onTZm2cHKcb+8+ZtcCAF1Xo73h/Mc7Fc2w8wsO8h02RL7mL/v6G7uttqropqKhNKWzljThaLBVEVcu1VRc8HRI70CiJDv3Xtl3rf/s7T/g5/39enp7l1/PFxkZl4eHWv/pfJw6JG5q7J9STZmWa3F56RmqJ84vPXtc0QVzu7OqxQdKXCnVEf3GR1yv1sU3YEl27ul3t2uSTbYyYFRLHnN3fcA9JG7q0IlRAaGjJYW9cZTOlX70ER/v3l2nGeKWeguoSJS1emilSrJzr2cccyWN0Jxlh+xbmppv/fOsuvVnbQVGhP9q9SqLzwBotDeUX89XvZx0+tZNz0Q8b/ByELnm22cvfHspR+nqsdGzZ456+SW9T4ScrGxzl0BCUjitSiaflPYdHOj99FMqQkpZ1PIlYfGz3G6NEbo4L7MLAKxF3mQidd+nciunF0eN9fJksRsAGO2xHt0jUxaHv5YoHktU5BdW5BeqGObo5esT+ttZw2NetPgFsE5Kc77WNo0wwMCwMe1OzpW/H5myuK607JtD6Y56X9RXVYv8l5loZh1T1q9+dPq5uu2s0aHsz/Zrm0YY4N2wSE2Wv0g/tbP/6figaBRsSHRkUuYRFYszWh9R7yjXImL/lKrVciWdXN69tzTnijWL7BdkxROVPXMWaHXQuH3m/Ecx00zsaeZESXZu9mf7XT9lak0HAyPCg341wZhukIXHT6hLJSvyC4+t2dC9T2+9kzAPL8/f7N6htP2d1Xh4eYbExY6YOkU8lhA8KW3XuMT4EbGTnSwTB6yMQAJo38mrOSev5kiStOY38+OjYwb16292RQDQ5cixRMTrSWVXrmo+Y7d1Uiojzp1PnwD/1leO0vmtckuQkTPieGFYh/+oZ3v5+rS9sFe3nTU6K4sMYPUbPizlwqmcvQdENgTqsh7r0X3mh++p6HBlmNKcK5bN6Z98aqDZJbRD8wizvqr6Xtl3Vgsk7pV9p+0L4/aZ8z37ehu5PZW8mFjRryhdbewKDy/PpMwjH8VMs8ghXTU5lgiJi60sKCrI/KrDvWEUkXOsQc+N9RkymBwCbo1AAuiAvGDil08PXj93IQsmAMB4Hl6eA8PGxO/9a0tTc/XN4srCohsnTiuduii3fQ8IHdX/F8F9A59mUXNXIL9yBoaNiVn7VmVBUWVhkdzLq+2Lp/XivO/gwMAXxvcfEcwwtwV5eHlGLE1qO8g7etZME+uB1UzfusnsEn7i4eU5NmGOPMJYV1omSVL1jeIfvv++7W26PfGEz9DBj/7u4931OvgEhI62zp9IkqSBYWOSMo9UFhaZXUj7wubOHjoxyuwq2mfNTyg9Xl1y2zRLkd9H/YKGPd6rpyZ3WJrzdW//AZrcVYfk8XFjHssVPfp6L8nKbGxoaD1Rlw+kbqrf8GH9hg+TV+5W3yguPn2m6kax0pZZvsOG9B0cOGBkiM/QwcyVQWfCHhKAMiyYAADrsNfYfmxokL9uO+jT9jpW0fa5xmu0NxQeP+Hinbhykdm62amGevsPsNrERsBduP6WdPEN+NAeLZpgLRoAALKWpubvy8vlr3+sv9+azraNu3p6e1szgAS0QiABqDHh2dDVsxPHDRvBggkAAAAAAAAAEEHLJkANeYcJvye9//fl386d+F89u3UzuyIAAAAAAAAAsDQPswsA3FhFrW3ZjjTvV2KWf7z1bmW52eUAAAAAAAAAgHURSAAa2J5xaNhrr8zdnHou/5rZtQAAAAAAAACAFRFIAJo5mHUi+s0lYSnzj+VcaGpuNrscAAAAAAAAALAQAglAY9fuFMelrnpmwcwDWV8RSwAAAAAAAACAjEAC0EVFrS1h87pnFszccHD3/R9+MLscAAAAAAAAADAZgQSgo4paW+q+T71fiSGWAAAAAAAAANDFEUgARiCWAAAAAAAAANDFEUgAxpFjiQ+Ofs7eEgAAAAAAAAC6GgIJwGjLdqSx5TUAAAAAAACAroZAAjBB65bXx3IumF0LAAAAAAAAABiBQAIwTUWtLS51VVjK/Ly7t8yuBQAAAAAAAAD0RSABmOzaneLQ5MS5m1PZ7xoAAAAAAABAJ0YgAVjCwawT3q/EsLEEAAAAAAAAgM6KQAKwkITN68Yvf40OTgAAAAAAAAA6HwIJwFrkDk7LP97KUgkAAAAAAAAAnQmBBGBF2zMOPbNgJkslAAAAAAAAAHQaBBKARVXU2lgqAQAAAAAAAKDTIJAALG17xqH//sOKqnt1ZhcCAAAAAAAAAC4hkACs7uTVnNCUxHP518wuBAAAAAAAAADUI5AA3EBFrS36zSUbDu42uxAAAAAAAAAAUIlAAnAbqfs+ZUsJAAAAAAAAAG6KQAJwJ/KWEmQSAAAAAAAAANwOgQTgZk5ezSGTAAAAAAAAAOB2CCQA93Pyas789zaYXQUAAAAAAAAAKEAgAbilg1knDmR9ZXYVAAAAAAAAACDqZw8ePHDy47uV5ecL8gyrBoAisyJfNLsEAAAAAAAAABDSQSABAAAAAAAAAADgOlo2AQAAAAAAAAAA3RFIAAAAAAAAAAAA3RFIAAAAAAAAAAAA3RFIAAAAAAAAAAAA3RFIAAAAAAAAAAAA3RFIAAAAAAAAAAAA3RFIAAAAAAAAAAAA3RFIAAAAAAAAAAAA3RFIAAAAAAAAAAAA3f0/lGOSpk1bcRYAAAAASUVORK5CYII='

def get_logo_bytes():
    return base64.b64decode(LOGO_BASE64)


ROOT=Path(__file__).resolve().parent
st.set_page_config(page_title='REM | Deal Economics',layout='wide',initial_sidebar_state='expanded')
st.markdown("""<style>
.block-container {max-width:1550px;padding-top:2rem;}
h1,h2,h3 {font-family:Arial,sans-serif;text-align:center;color:#074F42;}
[data-testid="stVerticalBlockBorderWrapper"] {border-color:#b7cbc4 !important;border-radius:12px;}
[data-testid="stMetric"] {background:#f2f6f3;border:1px solid #d0dee5;border-radius:8px;padding:12px 6px;text-align:center;min-height:95px;}
[data-testid="stMetricLabel"] {justify-content:center;font-size:.82rem;}
[data-testid="stMetricValue"] {font-size:1.35rem;font-weight:600;color:#074F42;}
[data-testid="stMetricLabel"] p {font-size:.82rem;}
[data-testid="stNumberInput"] input {text-align:center;}
</style>""",unsafe_allow_html=True)
logo_left,logo_center,logo_right=st.columns([1,3,1])
with logo_center:st.image(get_logo_bytes(),use_container_width=True)
st.caption('Natural gas deal economics • Monthly volume-weighted pricing')
st.caption('Pricing: GitHub input folder | Build TEXAS-INDEX-02')
def workstream_config():
    def setting(name,default=''):
        try:return str(st.secrets.get(name,os.getenv(name,default)))
        except Exception:return os.getenv(name,default)
    return setting('WORKSTREAM_REPO'),setting('WORKSTREAM_GITHUB_TOKEN'),setting('WORKSTREAM_BRANCH','main')

ws_repo,ws_token,ws_branch=workstream_config()
if '_pricing_carry' in st.session_state:
    carried=st.session_state['_pricing_carry']
    for k,v in carried.get('widgets',{}).items():
        if k not in ['wk_pricing_file','wk_asof'] and k not in st.session_state:
            if k in ['wk_start','wk_due']:v=pd.Timestamp(v).date() if v else None
            if k.startswith('wk_component_'):v=tuple(v)
            st.session_state[k]=v
    st.session_state['_restored_usage']=carried.get('usage',[])

if '_pending_workstream' in st.session_state:
    pending=st.session_state.pop('_pending_workstream')
    st.session_state.pop('_current_snapshot',None)
    for k in list(st.session_state):
        if k.startswith(('wk_','offer_mode_','offer_override_','usage_')):del st.session_state[k]
    for k,v in pending.get('widgets',{}).items():
        if k.startswith(('wk_','offer_mode_','offer_override_')):
            if k in ['wk_start','wk_due']:v=pd.Timestamp(v).date() if v else None
            if k.startswith('wk_component_'):v=tuple(v)
            st.session_state[k]=v
    st.session_state['_restored_usage']=pending.get('usage',[])
    st.session_state['_saved_pricing_file']=pending.get('pricing_file')
    st.session_state['_restore_saved_file_once']=True
    st.session_state['_saved_trade_date']=pending.get('trade_date')
    st.session_state['_restore_pricing']=True
    if 'wk_credit_mode' not in pending.get('widgets',{}):st.session_state['wk_credit_mode']='Manual'
    st.session_state['ws_latest']=False
    st.session_state['ws_name']=pending['name']

def save_current_workstream(copy=False):
    try:
        snapshot=st.session_state.get('_current_snapshot')
        if not snapshot:raise WorkstreamError('Load pricing and enter deal assumptions before saving.')
        path,sha,payload=save_workstream(ws_repo,ws_token,st.session_state.get('ws_name',''),snapshot,ws_branch,None if copy else st.session_state.get('_ws_path'),None if copy else st.session_state.get('_ws_sha'))
        st.session_state['_ws_path']=path;st.session_state['_ws_sha']=sha
        st.session_state['_saved_message']='Saved '+payload['name']+' at '+payload['saved_at']
        st.session_state['_last_saved_snapshot']=snapshot
    except Exception as e:st.session_state['_save_error']=str(e)

def open_workstream(path):
    try:
        payload,sha=load_workstream(ws_repo,ws_token,path,ws_branch)
        st.session_state['_pending_workstream']=payload
        st.session_state['_ws_path']=path;st.session_state['_ws_sha']=sha
        st.session_state['_saved_message']='Opened '+payload['name']+' | saved '+payload.get('saved_at','')
    except Exception as e:st.session_state['_save_error']=str(e)

with st.container(border=True):
    st.subheader('Saved workstreams')
    st.caption('Save a named working deal across browser refreshes and days. Saving is explicit; pricing updates do not change the saved assumptions until you save again.')
    st.text_input('Workstream name',key='ws_name',placeholder='Customer - December 2026 evaluation')
    if not ws_repo or not ws_token:
        st.info('One-time setup: configure a private workstream repository and GitHub token in Streamlit Secrets to enable durable Save / Open. See the included setup guide.')
    else:
        if st.checkbox('Show saved workstreams',key='ws_show'):
            try:
                saved=list_workstreams(ws_repo,ws_token,ws_branch)
                if saved:
                    selected=st.selectbox('Saved workstream',saved,format_func=lambda x:x['name'])
                    st.button('Open saved workstream',on_click=open_workstream,args=(selected['path'],))
                else:st.caption('No saved workstreams yet. Save the current deal below.')
            except Exception as e:st.error(str(e))
        one,two=st.columns(2)
        one.button('Save current workstream',on_click=save_current_workstream,disabled='_current_snapshot' not in st.session_state)
        two.button('Save as new workstream',on_click=save_current_workstream,args=(True,),disabled='_current_snapshot' not in st.session_state)
    if '_saved_message' in st.session_state:st.success(st.session_state['_saved_message'])
    if '_save_error' in st.session_state:st.error(st.session_state.pop('_save_error'))
    backup=st.file_uploader('Restore a workstream backup (optional; not pricing)',type=['json'],key='ws_backup')
    if backup is not None and st.button('Open backup workstream'):
        try:
            import json
            payload=json.loads(backup.getvalue())
            if payload.get('schema_version')!=1:raise ValueError('Unsupported backup version')
            st.session_state['_pending_workstream']=payload
            st.session_state.pop('_ws_path',None);st.session_state.pop('_ws_sha',None)
            st.rerun()
        except Exception as e:st.error(str(e))
    st.caption('To reprice a saved deal, select any existing file under Pricing setup and click Load selected pricing file. Customer assumptions, usage and manual offers are retained. File selection is saved with the workstream.')

summary_panel=st.container()

def format_table(df):
    formats={}
    for name in df.columns:
        if not pd.api.types.is_numeric_dtype(df[name]): continue
        key=str(name).lower()
        if '%' in key: formats[name]='{:,.2f}%'
        elif key in {'irr','gross_margin','discount_rate'}: formats[name]='{:.2%}'
        elif any(t in key for t in ['price','profit','revenue','cost','rate','npv','spread','margin','capital','supply','broker','transport','offer','break_even']):
            formats[name]='${:,.4f}' if any(t in key for t in ['price','rate','spread','break_even']) or '/mmbtu' in key else '${:,.0f}'
        else: formats[name]='{:,.0f}'
    return df.style.format(formats,na_rep='N/A')

@st.cache_data(ttl=300)
def load_location(path, modified_ns):
    return read_prices(path)

# Streamlit checks out this GitHub repository before running the app.
# Read its input folder directly: no upload or separate GitHub credentials needed.
input_candidates=[ROOT/'input',ROOT.parent/'input']
INPUT_DIR=next((d for d in input_candidates if d.is_dir()),ROOT/'input')
defaults=dict(customer='New customer', start=str(pd.Timestamp.today().to_period('M').to_timestamp().date()), transport=.35, broker=.30, margin=.18, usage={m:1700. for m in range(1,13)})
prices=None
with st.expander('1. Select ICE pricing file and contract dates', expanded=True):
    st.caption('Pricing source: merchg80/NG-Swing-GDD-Futures • main branch • input folder')
    if st.button('Refresh pricing files'):
        load_location.clear();st.rerun()
    if not INPUT_DIR.is_dir():
        st.error('The repository input folder was not found. Deploy this app from merchg80/NG-Swing-GDD-Futures, branch main, with streamlit_app.py at the repository root.')
        st.code(f'App location: {ROOT}\nExpected data folder: {ROOT / "input"}')
        st.stop()
    available=sorted([p for p in INPUT_DIR.rglob('*') if p.is_file() and p.suffix.lower() in {'.xlsx','.csv','.parquet'}],key=lambda p: ('ice' not in p.name.lower(),p.name.lower()))
    if not available:
        st.warning('The input folder contains no supported pricing files (.xlsx, .csv or .parquet).')
        st.stop()
    names=[str(p.relative_to(INPUT_DIR)) for p in available]
    desired=st.session_state.get('_saved_pricing_file')
    if st.session_state.pop('_restore_saved_file_once',False):
        if desired in names:
            st.session_state['wk_pricing_file']=desired
        else:
            st.warning('The saved pricing file is no longer available. Choose another file below.')
    if st.session_state.get('wk_pricing_file') not in names:st.session_state.pop('wk_pricing_file',None)
    def change_pricing_file():
        # A user selection supersedes the old saved-workstream file reference.
        st.session_state['_saved_pricing_file']=st.session_state.get('wk_pricing_file')
        st.session_state.pop('_restore_saved_file_once',None)
        if '_current_snapshot' in st.session_state:st.session_state['_pricing_carry']=st.session_state['_current_snapshot']
        st.session_state['_restore_pricing']=True
        st.session_state['_saved_trade_date']=None
        st.session_state.pop('wk_asof',None)
    selected_name=st.selectbox('ICE Cleared Gas pricing file',names,key='wk_pricing_file',on_change=change_pricing_file)
    selected_path=INPUT_DIR/selected_name
    st.caption(f'Selected repository file: input/{selected_path.relative_to(INPUT_DIR)}')
    stamp=f'{selected_path}:{selected_path.stat().st_mtime_ns}'
    if st.session_state.pop('_restore_pricing',False):
        st.session_state['loaded_pricing_stamp']=stamp
    if st.button('Load / refresh selected pricing file',type='primary'):
        st.session_state['loaded_pricing_stamp']=stamp
        st.session_state.pop('wk_asof',None)
        st.session_state['_saved_trade_date']=None
    if st.session_state.get('loaded_pricing_stamp')!=stamp:
        st.info('Choose an existing ICE file above, then click Load / refresh selected pricing file. No upload is needed.')
        st.stop()
    try:
        with st.spinner(f'Reading {selected_path.name} — large ICE files may take time…'):
            prices=load_location(str(selected_path),selected_path.stat().st_mtime_ns)
        st.success(f'Loaded {len(prices):,} settlement rows • latest trade date {prices.trade_date.max():%Y-%m-%d}')
    except Exception as e:
        st.error(f'Could not read input/{selected_path.relative_to(INPUT_DIR)}: {e}')
        st.caption('Select an ICE settlement file with trade date, hub, product, strip, contract and settlement price columns. A daily Date-by-hub master is not a monthly forward curve.')
        st.stop()
    dates=sorted(prices.trade_date.unique(),reverse=True)
    saved_date=st.session_state.get('_saved_trade_date')
    if saved_date and pd.Timestamp(saved_date) in dates:st.session_state.setdefault('wk_asof',pd.Timestamp(saved_date))
    if st.session_state.get('wk_asof') not in dates:st.session_state.pop('wk_asof',None)
    asof=st.selectbox('Settlement trade date',dates,format_func=lambda d:pd.Timestamp(d).strftime('%b %d, %Y'),key='wk_asof')
    st.caption('All components use this exact trade date. Missing months block a quote.')
    customer=st.text_input('Customer',value=defaults['customer'],key='wk_customer')
    utility=st.text_input('Utility',value='',placeholder='Serving utility name',key='wk_utility')
    client_location=st.text_input('Client location (address or city)',value='',placeholder='Street address or city, state',key='wk_client_location')
    facility_type=st.selectbox('Facility type',['Housing','Commercial','Industrial','Other'],index=None,placeholder='Select facility type',key='wk_facility_type')
    broker_name=st.text_input('Broker name / firm',value='',key='wk_broker_name')
    offer_due=st.date_input('Offer due date',value=None,key='wk_due')
    start_entered=st.date_input('Contract start date',value=(pd.Timestamp(asof).to_period('M')+1).to_timestamp().date(),key='wk_start')
    start=start_entered.replace(day=1)
    if start_entered.day!=1: st.caption('Monthly pricing uses the start month. Enter first-month usage for the remaining delivery days; usage is not automatically prorated.')
    term=st.slider('Term (months)',min_value=1,max_value=48,value=12,step=1,key='wk_term')

st.subheader('Market components')
st.caption('Choose products by name. Additive mode follows the workbook: Henry Hub + HSC basis + HSC index adjustment. Outright mode uses one full-price product.')
kind=st.radio('Curve construction',['Additive components','Outright fixed price'],horizontal=True,key='wk_kind')
options=list(prices[['hub','product','contract']].drop_duplicates().itertuples(index=False,name=None))
selections=[]
cols=st.columns(3 if kind=='Additive components' else 1)
for i,c in enumerate(cols):
    preferred=['H','HXS','SHS'][i] if kind=='Additive components' else 'H'
    idx=next((j for j,x in enumerate(options) if x[2]==preferred),0)
    if st.session_state.get(f'wk_component_{i}') not in options:
        if st.session_state.get(f'wk_component_{i}') is not None:st.warning(f'Saved component {i+1} is not in this pricing file. Review the replacement selection before quoting.')
        st.session_state.pop(f'wk_component_{i}',None)
    selections.append(c.selectbox((['Henry Hub','Basis','Index adjustment (SHS)'][i] if kind=='Additive components' else 'Outright market curve'),options,index=idx,format_func=lambda x:f'{x[0]} | {x[1]} ({x[2]})',key=f'wk_component_{i}'))
if len(set(selections))!=len(selections): st.error('Select distinct components.');st.stop()
st.subheader('Customer pricing structure')
def reset_offer_structure():
    for key in list(st.session_state):
        if key.startswith(('offer_mode_','offer_override_')) or key in ('wk_custom_override','wk_custom_price','wk_transport_pass_through'):
            st.session_state.pop(key,None)
pricing_mode=st.radio('Contract pricing',['Fixed price','Index + adder'],horizontal=True,key='wk_pricing_mode',on_change=reset_offer_structure)
is_indexed=pricing_mode=='Index + adder'
customer_adder=st.number_input('Customer adder above monthly index ($/MMBtu)',value=.38,step=.01,format='%.4f',key='wk_customer_adder') if is_indexed else .38
supplier_premium=st.number_input('Supplier premium / discount to the same index ($/MMBtu)',value=0.,step=.01,format='%.4f',key='wk_supplier_premium')
transport_pass_through=st.checkbox('Transportation is passed through in full to the customer',value=is_indexed,key='wk_transport_pass_through')
deal_options=dict(pricing_mode=pricing_mode,customer_adder=customer_adder,supplier_premium=supplier_premium,transport_pass_through=transport_pass_through)
quote_label='Index adder ($/MMBtu)' if is_indexed else 'Offer excluding transport pass-through ($/MMBtu)' if transport_pass_through else 'All-in offer ($/MMBtu)'
break_even_label='Break-even adder ($/MMBtu)' if is_indexed else 'Break-even offer ($/MMBtu)'
st.caption('Monthly index = Henry Hub + signed basis + signed index adjustment. Negative components remain negative. Index + adder bills that monthly index plus the customer adder. Transportation pass-through is billed separately at its cost; financing that expense still incurs the capital charge.')
st.subheader('Deal assumptions')
a,b,c,d=st.columns(4)
transport=a.number_input('Upstream transport ($/delivered MMBtu)',min_value=0.,value=defaults['transport'],step=.01,format='%.4f',key='wk_transport')
broker=b.number_input('Brokerage ($/delivered MMBtu)',min_value=0.,value=defaults['broker'],step=.01,format='%.4f',key='wk_broker')
margin=c.number_input('Target gross margin ($/MMBtu)',value=defaults['margin'],step=.01,format='%.4f',key='wk_margin')
fuel=d.number_input('Supply fuel retention (%)',min_value=0.,max_value=99.,value=0.,step=.1,key='wk_fuel')/100
st.caption('Margin is after supply, transport and brokerage, after the capital charge, before overhead, other credit fees and taxes. Transport and brokerage apply to delivered usage. Fuel gross-up applies to gas supply only.')
st.subheader('Cost of capital')
cap1,cap2=st.columns(2)
capital_hurdle=cap1.number_input('Annual capital hurdle (%)',min_value=0.,value=8.,step=.5,key='wk_capital')/100
financing_days=cap2.number_input('Days of costs financed',min_value=0.,value=30.,step=1.,key='wk_days')
st.caption('Capital charge per delivered MMBtu = positive supply + transport + brokerage cost per MMBtu × annual hurdle × financing days ÷ 365. The charge is included in the target offer and break-even price, and deducted from modeled profit. Financing days represent the cash collection gap. Negative total costs have a zero capital base.')
months=pd.date_range(start,periods=term,freq='MS')
base=pd.DataFrame({'month':months,'usage_mmbtu':[defaults['usage'].get(m.month,0.) for m in months]})
if st.session_state.get('_restored_usage'):
    saved_usage={pd.Timestamp(row['month']):float(row['usage_mmbtu']) for row in st.session_state['_restored_usage']}
    base['usage_mmbtu']=[saved_usage.get(m,v) for m,v in zip(base.month,base.usage_mmbtu)]
st.subheader('Customer monthly usage')
st.caption('Enter total delivered MMBtu for each month. The attached workbook starts at 1,700 MMBtu/month. Values are estimates, not daily quantities.')
usage_file=st.file_uploader('Load a saved usage schedule (optional)',type=['csv'])
if usage_file is not None:
    try:
        loaded=pd.read_csv(usage_file)
        loaded['month']=pd.to_datetime(loaded['month']).dt.to_period('M').dt.to_timestamp()
        if loaded.month.duplicated().any(): raise ValueError('Duplicate usage months')
        base['usage_mmbtu']=base.month.map(loaded.set_index('month')['usage_mmbtu'])
        if base.usage_mmbtu.isna().any(): raise ValueError('Schedule must cover every selected month')
    except Exception as e: st.error(f'Invalid usage schedule: {e}');st.stop()
usage=st.data_editor(base,hide_index=True,disabled=['month'],column_config={'usage_mmbtu':st.column_config.NumberColumn('Usage (MMBtu/month)',min_value=0.,required=True,format='localized')},key=f'usage_{start}_{term}',use_container_width=True)
with st.expander('Upfront credit and investment returns',expanded=True):
    a,b,c=st.columns(3)
    discount=a.number_input('Required annual discount rate (%)',key='wk_discount',min_value=0.,value=10.,step=.5,format='%.2f')/100
    credit_mode=b.radio('Upfront credit sizing',['Automatic','Manual'],key='wk_credit_mode')
    automatic_credit=credit_mode=='Automatic'
    credit_pricing=curve(prices,asof,selections,months)
    try:
        credit_result,_=calculate(credit_pricing,usage.usage_mmbtu,transport,broker,margin,fuel,capital_hurdle=capital_hurdle,financing_days=financing_days,**deal_options)
        credit_parts=required_credit(credit_result,transport)
    except ValueError as e:
        st.error('Unable to size credit for the selected term: '+str(e));st.stop()
    auto_amount=float(credit_parts['total_credit'])
    investment=auto_amount if automatic_credit else b.number_input('Manual upfront cash posted for credit ($)',key='wk_investment',min_value=0.,value=auto_amount,step=100.,format='%.2f')
    b.metric('Upfront credit cash',f'${investment:,.2f}')
    full_recovery=c.checkbox('Return all posted credit cash at contract end',value=True,key='wk_full_recovery')
    recovery=investment if full_recovery else c.number_input('Credit cash returned at contract end ($)',key='wk_recovery',min_value=0.,value=0.,step=100.,format='%.2f')
    c.caption(f'Credit cash recovered: ${recovery:,.2f} in month {term}')
    st.caption(f"Automatic FT credit: ${transport:.4f} x {credit_parts['top_two_average_volume']:,.0f} MMBtu (average of the two highest-use months) = ${credit_parts['ft_credit']:,.2f}. Supply credit: {credit_parts['highest_volume']:,.0f} MMBtu x ${credit_parts['highest_supply_price']:.4f} = ${credit_parts['supply_credit']:,.2f}. Total: ${auto_amount:,.2f}.")
    st.caption('Supply price is the selected monthly gas-supply curve, fuel-adjusted, before transportation, brokerage and marketer margin. The highest price and highest volume need not occur in the same month. One-month terms use that single month for the FT average. Negative peak supply prices are floored at zero for collateral sizing. Automatic sizing is calculated separately over each standard term; manual sizing uses your same input across all terms. Collateral is returned only at the end of that term when enabled.')
offers=standard_offers(prices,asof,selections,start,usage,defaults['usage'],transport,broker,margin,fuel,capital_hurdle,financing_days,discount,investment,recovery,automatic_credit=automatic_credit,full_credit_recovery=full_recovery,**deal_options)
with summary_panel:
    with st.container(border=True):
        st.subheader('Projected offers | standard contract terms')
        st.caption(f'Start: {start_entered:%b %d, %Y} • Target margin after capital: ${margin:.4f}/MMBtu • Settlements: {pd.Timestamp(asof):%Y-%m-%d}')
        st.caption(f'Shared assumptions: transport ${transport:.4f}/MMBtu | brokerage ${broker:.4f}/MMBtu | capital hurdle {capital_hurdle:.2%} | financing {financing_days:g} days | target margin ${margin:.4f}/MMBtu')
        offer_cols=st.columns(5)
        overrides={}
        output_panels=[]
        for col,row in zip(offer_cols,offers.itertuples(index=False)):
            with col:
                with st.container(border=True):
                    st.markdown(f'### {row.term_months} months')
                    if row.status!='Available':
                        st.warning(row.status)
                        output_panels.append(None)
                        continue
                    st.caption(f'Model required {"adder" if is_indexed else "offer"}: ${row.target_price:,.4f}/MMBtu')
                    setting=st.selectbox('Offer setting',(['Entered adder','Model target','Target less $0.10','Target less $0.20','Manual price'] if is_indexed else ['Model target','Target less $0.10','Target less $0.20','Manual price']),key=f'offer_mode_{row.term_months}')
                    if setting=='Manual price':
                        overrides[row.term_months]=st.number_input(quote_label,value=float(row.offer_price),step=.01,format='%.4f',key=f'offer_override_{row.term_months}')
                    elif setting=='Model target':
                        overrides[row.term_months]=row.target_price
                    elif setting.startswith('Target less'):
                        overrides[row.term_months]=row.target_price-(.1 if setting=='Target less $0.10' else .2)
                    output_panels.append(st.container())
        offers=standard_offers(prices,asof,selections,start,usage,defaults['usage'],transport,broker,margin,fuel,capital_hurdle,financing_days,discount,investment,recovery,overrides,automatic_credit=automatic_credit,full_credit_recovery=full_recovery,**deal_options)
        for panel,row in zip(output_panels,offers.itertuples(index=False)):
            if panel is None: continue
            with panel:
                st.metric(quote_label,f'${row.offer_price:,.4f}')
                if is_indexed: st.caption(f'Projected all-in: ${row.projected_all_in_price:.4f}/MMBtu')
                st.metric(f'NPV @ {discount:.2%}',f'${row.npv:,.0f}' if pd.notna(row.npv) else 'N/A')
                st.metric('Annualized IRR',f'{row.irr:.2%}' if pd.notna(row.irr) else 'N/A')
                st.caption(f'Credit posted: ${row.upfront_credit:,.0f}')
                st.caption(f'Collateral returned: ${row.credit_recovery:,.0f} in month {row.term_months}' if row.credit_recovery>0 else 'No collateral recovery included')
                st.metric('Contract profit',f'${row.contract_profit:,.0f}')
                st.metric('Spread ($/MMBtu)',f'${row.spread:,.4f}')
                st.metric('Gross margin',f'{row.gross_margin:.2%}' if pd.notna(row.gross_margin) else 'N/A')
        st.caption('Entered adder uses the shared customer adder for all indexed terms. Model target calculates the adder needed for your target margin. Manual entries remain as entered; target reductions use the model required adder / offer.')
        st.caption(f'Credit sizing: {credit_mode}. Each card shows its own posted and recovered credit cash. Discount rate: {discount:.2%} effective annual. NPV is valued at contract start; IRR is annualized.')
        st.caption('Each offer follows the selected pricing structure. Indexed offers show the adder above the monthly index; projected revenue uses the forward index. Full transportation pass-through is added to the bill and deducted as an equal cost. Each model target includes supply, fuel, upstream transport, brokerage, capital charge and target margin. Entered months use your exact volumes. Beyond the entered schedule, the first 12 months repeat by calendar month; missing calendar months use the displayed default.')
        with st.expander('Standard offer detail and volume assumptions'):
            st.dataframe(format_table(offers),hide_index=True,use_container_width=True)
        st.download_button('Download standard offers CSV',offers.to_csv(index=False).encode(),'standard_contract_offers.csv','text/csv')
with summary_panel:
    with st.container(border=True):
        st.subheader('Detailed summary selection')
        st.caption(f'The summary, monthly economics, offer comparisons and PDF below use the selected {term}-month contract. Change the contract term in the pricing setup to review another term.')
        if term in [12,18,24,36,48]:
            quote=overrides.get(term)
        else:
            custom_override=st.checkbox(f'Manual offer for selected {term}-month contract',key='wk_custom_override')
            quote=st.number_input(quote_label,value=(customer_adder if is_indexed else 4.),step=.01,format='%.4f',key='wk_custom_price') if custom_override else None
import json
widget_snapshot={}
for k,v in st.session_state.items():
    if k.startswith(('wk_','offer_mode_','offer_override_')) and k not in ['wk_pricing_file','wk_asof']:
        if hasattr(v,'isoformat'):v=v.isoformat()
        elif isinstance(v,tuple):v=list(v)
        widget_snapshot[k]=v
current_usage=[dict(month=str(pd.Timestamp(row.month).date()),usage_mmbtu=float(row.usage_mmbtu)) for row in usage.itertuples(index=False)]
snapshot=dict(pricing_structure=deal_options,credit_sizing_mode=credit_mode,calculated_credit=credit_parts,credit_posted=investment,schema_version=1,widgets=widget_snapshot,usage=current_usage,pricing_file=str(selected_path.relative_to(INPUT_DIR)),trade_date=str(pd.Timestamp(asof).date()),latest_pricing=False,name=st.session_state.get('ws_name',''))
st.session_state['_current_snapshot']=snapshot
st.session_state.pop('_pricing_carry',None)
with summary_panel:
    st.download_button('Download workstream backup (JSON)',json.dumps(snapshot,indent=2,allow_nan=False).encode(),'REM_Workstream.json','application/json')
try:
    p=curve(prices,asof,selections,months)
    if p.isna().any().any():
        st.error('Pricing is missing for these delivery months/components. Select another trade date, product or term.')
        st.dataframe(p[p.isna().any(axis=1)].style.format('${:,.4f}',na_rep='Missing'));st.stop()
    result,target=calculate(p,usage.usage_mmbtu,transport,broker,margin,fuel,quote,capital_hurdle,financing_days,**deal_options)
except ValueError as e: st.error(str(e));st.stop()
summary=deal_summary(result,discount,investment,recovery)
with summary_panel:
    with st.expander('Verify collateral timing and investment cash flows'):
        st.caption(f'For this {term}-month contract, upfront credit is posted in month 0. '+(f'Collateral is returned once, in month {term}, alongside the final operating profit.' if recovery>0 else 'No collateral repayment is included.'))
        cash_audit=investment_cashflows(result,investment,recovery)
        monthly_discount=(1+discount)**(1/12)-1
        cash_audit['discount_factor']=1/(1+monthly_discount)**cash_audit.month_number
        cash_audit['present_value']=cash_audit.net_cashflow*cash_audit.discount_factor
        st.dataframe(cash_audit.style.format({'month_number':'{:,.0f}','operating_profit':'${:,.2f}','collateral_cashflow':'${:,.2f}','net_cashflow':'${:,.2f}','discount_factor':'{:.6f}','present_value':'${:,.2f}'}),hide_index=True,use_container_width=True)
        st.caption(f'Sum of present values = NPV ${cash_audit.present_value.sum():,.2f}. IRR uses the same net cash-flow column, including month 0.')
        st.download_button('Download investment cash-flow audit CSV',cash_audit.to_csv(index=False).encode(),'investment_cashflow_audit.csv','text/csv')

with summary_panel:
    pdf_panel=st.container()
with summary_panel.container(border=True):
    st.subheader('Deal summary')
    st.caption(f"Utility: {utility or 'Not specified'} • Facility: {facility_type or 'Not specified'}")
    st.caption(f"Client location: {client_location or 'Not specified'}")
    st.caption(f'Broker: {broker_name or "Not specified"} • Offer due: {offer_due or "Not specified"}')
    st.caption(f'{customer} • Start {start_entered:%b %d, %Y} • {term} month term • Settlements {pd.Timestamp(asof):%b %d, %Y}')
    top=st.columns(4)
    top[0].metric('Annual revenue (annualized)',f"${summary['annual_revenue']:,.0f}")
    top[1].metric('Volume sold (contract MMBtu)',f"{summary['volume']:,.0f}")
    top[2].metric('Projected weighted all-in price ($/MMBtu)',f"${summary['contract_price']:,.4f}")
    top[3].metric(f'NPV @ {discount:.2%} annual',f"${summary['npv']:,.0f}")
    bottom_break_even=target-margin
    bottom=st.columns(4)
    bottom[3].metric(break_even_label,f'${bottom_break_even:,.4f}')
    bottom[0].metric('IRR (annualized)',f"{summary['irr']:.2%}" if summary['irr'] is not None else 'N/A')
    bottom[1].metric('MMBtu spread ($/MMBtu)',f"${summary['spread']:,.4f}")
    bottom[2].metric('Gross margin (%)',f"{summary['gross_margin']:.2%}" if summary['gross_margin'] is not None else 'N/A')
    st.markdown('**All-in price breakdown ($/delivered MMBtu)**')
    weighted_supply=float(result.supply_cost.sum()/result.usage_mmbtu.sum())
    weighted_capital=float(result.capital_cost.sum()/result.usage_mmbtu.sum())
    bridge=pd.DataFrame({'Component':['Gas supply (volume-weighted)','Transportation pass-through' if transport_pass_through else 'Transportation included','Brokerage included','Capital charge included','Marketer margin after capital','Projected weighted all-in customer price'],'Rate ($/MMBtu)':[weighted_supply,transport,broker,weighted_capital,summary['spread'],summary['contract_price']]})
    st.dataframe(bridge.style.format({'Rate ($/MMBtu)':'${:,.4f}'}),hide_index=True,use_container_width=True)
    st.caption('The projected all-in customer rate includes any separate transportation pass-through once. Transportation recovery and expense are equal when enabled. Brokerage, supplier premium, fuel and the financing charge reduce the remaining spread. The financing base includes transportation even when reimbursed.')
    if is_indexed:
        st.metric('Customer monthly index adder',f'${result.offer_rate.iloc[0]:.4f}/MMBtu')
        st.caption('Customer settlement: monthly Henry Hub + signed basis + signed SHS adjustment + customer adder + transportation pass-through when enabled. The projected all-in price is volume-weighted, not a fixed quote.')
    st.caption(f'PV basis: {discount:.2%} effective annual discount rate; equivalent monthly rate {(1+discount)**(1/12)-1:.4%}. Valuation date: contract start. Monthly cash flows occur at month-end. IRR is solved independently of this rate.')
    st.caption(f"Capital hurdle: {capital_hurdle:.1%} • Financing days: {financing_days:g} • Weighted capital charge: ${(result.capital_cost.sum()/result.usage_mmbtu.sum()):.4f}/MMBtu • Contract capital charge: ${result.capital_cost.sum():,.0f}")
    st.caption('Annual revenue = total contract revenue × 12 ÷ term months. Spread = gross profit ÷ delivered volume. Gross margin = gross profit ÷ revenue. '+summary['irr_status']+'.')
    st.download_button('Download deal summary CSV',pd.DataFrame([dict(**deal_options,customer_offer_rate=float(result.offer_rate.iloc[0]),customer=customer,utility=utility,client_location=client_location,facility_type=facility_type,broker_name=broker_name,offer_due_date=str(offer_due) if offer_due else '',start_date=str(start_entered),term_months=term,trade_date=str(pd.Timestamp(asof).date()),capital_hurdle=capital_hurdle,financing_days=financing_days,capital_charge_per_mmbtu=result.capital_cost.sum()/result.usage_mmbtu.sum(),annual_discount_rate=discount,credit_sizing_mode=credit_mode,automatic_ft_credit=credit_parts['ft_credit'],automatic_supply_credit=credit_parts['supply_credit'],upfront_credit_cash=investment,credit_cash_recovery=recovery,**summary)]).to_csv(index=False).encode(),'deal_summary.csv','text/csv')
vol=result.usage_mmbtu.sum();gp=result.gross_profit.sum()
a,b,c,d=st.columns(4)
a.metric('Required index adder' if is_indexed else 'Required offer excluding pass-through',f'${target:,.4f}/MMBtu')
b.metric('Total contract volume',f'{vol:,.0f} MMBtu')
c.metric('Contract gross profit',f'${gp:,.0f}')
d.metric('Realized gross margin',f'${gp/vol:,.4f}/MMBtu')
st.caption(f'Customer: {customer} • Settlements: {pd.Timestamp(asof):%Y-%m-%d} • Annualized gross profit: ${gp*12/term:,.0f}')
st.subheader('Offer comparison')
if is_indexed: st.caption('Offer and break-even columns below are adders above the monthly index. Projected all-in prices include the signed market curve and any transportation pass-through.')
st.caption('Offer uses the proposed customer price, or target price when no proposed price is entered. Costs and volume remain constant. Break-even includes the capital charge and excludes overhead and taxes.')
scenario_rows=[]
scenario_monthly=pd.DataFrame(index=result.index)
for label,reduction in [('Offer',0.),('Offer less $0.10',.10),('Offer less $0.20',.20)]:
    offer_price=float(result.offer_rate.iloc[0])-reduction
    sr,_=calculate(p,usage.usage_mmbtu,transport,broker,margin,fuel,offer_price,capital_hurdle,financing_days,**deal_options)
    ss=deal_summary(sr,discount,investment,recovery)
    scenario_rows.append({'Scenario':label,'Offer $/MMBtu':offer_price,'Pricing structure':pricing_mode,'Projected all-in $/MMBtu':ss['contract_price'],'Break-even $/MMBtu':target-margin,'Spread $/MMBtu':ss['spread'],'Gross margin %':ss['gross_margin']*100 if ss['gross_margin'] is not None else None,'Average monthly GP $':sr.gross_profit.sum()/term,'Annualized GP $':sr.gross_profit.sum()*12/term,'Contract GP $':sr.gross_profit.sum(),f'NPV @ {discount:.2%} $':ss['npv'],'Annual discount rate %':discount*100,'Annualized IRR %':ss['irr']*100 if ss['irr'] is not None else None})
    scenario_monthly[label]=sr.gross_profit
scenario_table=pd.DataFrame(scenario_rows)
with pdf_panel:
    try:
        pdf_bytes=build_summary_pdf(dict(**deal_options,customer_offer_rate=float(result.offer_rate.iloc[0]),credit_mode=credit_mode,ft_credit=credit_parts['ft_credit'],supply_credit=credit_parts['supply_credit'],customer=customer,utility=utility,client_location=client_location,facility_type=facility_type,broker_name=broker_name or 'Not specified',due_date=str(offer_due) if offer_due else 'Not specified',start=str(start_entered),term=term,asof=str(pd.Timestamp(asof).date()),discount=discount,break_even=target-margin,transport=transport,broker=broker,margin=margin,fuel=fuel,capital_hurdle=capital_hurdle,days=financing_days,investment=investment,recovery=recovery,source='input/'+str(selected_path.relative_to(INPUT_DIR)),components=' + '.join(' | '.join(x) for x in selections)),summary,result,scenario_rows)
        st.download_button('Download professional deal summary (PDF)',pdf_bytes,'REM_Deal_Summary.pdf','application/pdf')
    except ImportError:
        st.warning('PDF export requires reportlab in requirements.txt. CSV downloads remain available.')

st.dataframe(format_table(scenario_table),hide_index=True,use_container_width=True)
st.caption('Monthly and annualized averages use the selected contract term. The schedule below shows actual monthly profits. NPV includes initial investment and recovery; break-even price is the volume-weighted cost including the capital charge.')
with st.expander('Monthly offer economics',expanded=True):
    st.line_chart(scenario_monthly)
    st.dataframe(format_table(scenario_monthly.reset_index().rename(columns={c:c+" profit ($)" for c in scenario_monthly.columns})), hide_index=True,use_container_width=True)
st.download_button('Download offer comparison CSV',scenario_table.to_csv(index=False).encode(),'offer_comparison.csv','text/csv')
st.download_button('Download monthly offer economics CSV',scenario_monthly.reset_index().to_csv(index=False).encode(),'monthly_offer_economics.csv','text/csv')
t1,t2,t3=st.tabs(['Monthly economics','Term comparisons','Method & source'])
with t1:
    st.line_chart(result[['market_rate','landed_rate','all_in_cost_rate','customer_rate']])
    st.bar_chart(result[['gross_profit']])
    st.dataframe(format_table(result.reset_index()),hide_index=True,use_container_width=True)
    totals=result[['usage_mmbtu','revenue','supply_cost','transport_cost','broker_cost','capital_cost','gross_profit']].sum()
    st.write('Contract totals')
    st.dataframe(format_table(totals.to_frame('Total').T),hide_index=True,use_container_width=True)
    st.write('Selected-strip averages',pd.DataFrame({'Simple average':result[['market_rate','landed_rate','all_in_cost_rate','customer_rate']].mean(),'Volume-weighted average':result[['market_rate','landed_rate','all_in_cost_rate','customer_rate']].mul(result.usage_mmbtu,axis=0).sum()/vol}))
with t2:
    st.caption('Compares standard terms and your selected term within the entered usage schedule. Longer terms require extending the schedule above. Each term is priced separately to the same target margin.')
    rows=[]
    for n in sorted(set([1,3,6,12,18,24,36,48,term])):
        if n>term or usage.usage_mmbtu.iloc[:n].sum()<=0: continue
        r,q=calculate(p.iloc[:n],usage.usage_mmbtu.iloc[:n],transport,broker,margin,fuel,capital_hurdle=capital_hurdle,financing_days=financing_days,**deal_options)
        rows.append({'Months':n,'Volume MMBtu':r.usage_mmbtu.sum(),'Required adder / offer':q,'Gross profit':r.gross_profit.sum(),'Annualized GP':r.gross_profit.sum()*12/n})
    st.dataframe(format_table(pd.DataFrame(rows)),hide_index=True)
with t3:
    st.write('Indexed customer rate = signed market curve + customer adder + transport pass-through. Supply = (signed market curve + supplier premium) / (1 - fuel). Profit = revenue - supply - transport - brokerage - capital charge. Required adder = volume-weighted (all-in cost - market curve - transportation recovery) + target margin.')
    st.write('Uses actual entered monthly volumes, calendar delivery months and a single settlement snapshot. No interpolation, silent zero filling or fallback to older settlements. Daily swing observations are not converted to forward prices.')
    st.write('The capital charge is a modeled financing expense, not a full working-capital schedule. NPV separately discounts profit after that charge; initial investment and terminal recovery are additional cash flows and should not duplicate the financed operating costs. The workbook’s NPV/carry formulas and daily-volume approximation are not reproduced. This model supports fixed offers and a level adder to the same monthly purchase and sale index. It excludes utility charges, imbalance exposure, credit/collateral financing and overhead unless included in your rates.')
    st.dataframe(pd.DataFrame(selections,columns=['Hub','Product','Contract']),hide_index=True)
export=result.reset_index(); export['utility']=utility; export['client_location']=client_location; export['facility_type']=facility_type; export.insert(0,'customer',customer);export.insert(1,'trade_date',pd.Timestamp(asof));export.insert(2,'contract_start_date',str(start_entered));export.insert(3,'term_months',term);export.insert(2,'components',' + '.join(x[2] for x in selections))
st.download_button('Download monthly economics CSV',export.to_csv(index=False).encode(),'rem_deal_economics.csv','text/csv')
st.download_button('Download usage schedule CSV',usage.to_csv(index=False).encode(),'customer_usage.csv','text/csv')
